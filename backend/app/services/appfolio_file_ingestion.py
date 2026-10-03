"""CSV/XLSX staging foundation for Phase 4.13 AppFolio migrations.

Files are parsed as data only. Raw bytes are never persisted, spreadsheet
formulas/macros are never executed, and staging never creates customer business
records. All persisted rows remain tied to the existing PlatformMigrationRun.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.models.gl_account import GLAccount
from app.models.platform_migration import (
    PlatformMigrationItem,
    PlatformMigrationRun,
    PlatformMigrationStagedRow,
    PlatformMigrationUpload,
)
from app.models.property import Property, Unit
from app.models.user import User, UserRole
from app.models.vendor import Vendor

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 10_000
MAX_COLUMNS = 100
MAX_SHEETS = 10
MAX_XLSX_ENTRIES = 2_000
MAX_XLSX_UNCOMPRESSED_BYTES = 64 * 1024 * 1024

SECRET_HEADER_KEYS = {
    "apikey",
    "clientsecret",
    "accesstoken",
    "refreshtoken",
    "authorization",
    "cookie",
    "password",
    "sessioncookie",
}

PROPERTY_ALIASES: dict[str, tuple[str, ...]] = {
    "source_id": ("Id", "Property Id", "PropertyId"),
    "name": ("Name", "Property Name"),
    "address_line1": ("Address1", "Address 1", "Street", "Street Address"),
    "address_line2": ("Address2", "Address 2", "Street 2"),
    "city": ("City",),
    "state": ("State", "Province"),
    "zip_code": ("Zip", "ZipCode", "Zip Code", "PostalCode", "Postal Code"),
    "property_type": ("PropertyType", "Property Type", "Type"),
    "hidden_at": ("HiddenAt", "Hidden At"),
}
PROPERTY_REQUIRED = (
    "source_id",
    "name",
    "address_line1",
    "city",
    "state",
    "zip_code",
)

# Verified Unit Directory export columns. These aliases intentionally stay
# narrower than the target Unit model: staging must not invent AppFolio fields.
UNIT_ALIASES: dict[str, tuple[str, ...]] = {
    "source_id": ("Unit ID", "Unit Id", "UnitID"),
    "source_property_id": ("Property ID", "Property Id", "PropertyID"),
    "unit_name": ("Unit Name", "Unit"),
    "property_name": ("Property Name", "Property"),
    "unit_address": ("Unit Address",),
    "address_line1": (
        "Unit Street Address 1",
        "Unit Address 1",
        "Unit Street 1",
    ),
    "address_line2": (
        "Unit Street Address 2",
        "Unit Address 2",
        "Unit Street 2",
    ),
    "city": ("Unit City",),
    "state": ("Unit State",),
    "zip_code": ("Unit Zip", "Unit Zip Code"),
}
# Unit Name is the only universally required display field in the verified
# export contract. Stable Unit/Property IDs are optional in AppFolio exports,
# but missing IDs remain REVIEW blockers for safe later commit.
UNIT_REQUIRED = ("unit_name",)

# Verified Tenant Directory export contract from current AppFolio export
# instructions used by migration/integration vendors. Keep this narrower than
# the target User/Lease models: staging preserves source facts but does not
# infer occupancy, lease liability, identity or billing semantics.
TENANT_ALIASES: dict[str, tuple[str, ...]] = {
    "source_id": ("Tenant ID", "Tenant Id", "TenantID"),
    "tenant_name": ("Tenant", "Tenant Name"),
    "phone_numbers": ("Phone Numbers", "Phone Number", "Phone"),
    "emails": ("Emails", "Email", "Tenant Email"),
    "tenant_address_line1": ("Tenant Street Address 1", "Tenant Address 1"),
    "tenant_address_line2": ("Tenant Street Address 2", "Tenant Address 2"),
    "tenant_city": ("Tenant City",),
    "tenant_state": ("Tenant State",),
    "tenant_zip": ("Tenant Zip", "Tenant Zip Code"),
    "source_property_id": ("Property ID", "Property Id", "PropertyID"),
    "property_name": ("Property Name", "Property"),
    "property_address": ("Property Address",),
    "source_unit_id": ("Unit ID", "Unit Id", "UnitID"),
    "unit_name": ("Unit", "Unit Name"),
    "move_in": ("Move-in", "Move In", "Move-in Date"),
    "move_out": ("Move-out", "Move Out", "Move-out Date"),
    "lease_from": ("Lease From", "Lease Start"),
    "lease_to": ("Lease To", "Lease End"),
}
TENANT_REQUIRED = ("tenant_name",)

# Lease/occupancy staging is explicit-only. Stable relationship identities come
# from the already verified Tenant Directory export contract. Rent Roll fields
# may supplement those identities as source evidence, but names/addresses are
# never used as relationship keys.
LEASE_OCCUPANCY_ALIASES: dict[str, tuple[str, ...]] = {
    "source_tenant_id": ("Tenant ID", "Tenant Id", "TenantID"),
    "tenant_name": ("Tenant", "Tenant Name"),
    "source_property_id": ("Property ID", "Property Id", "PropertyID"),
    "property_name": ("Property Name", "Property"),
    "source_unit_id": ("Unit ID", "Unit Id", "UnitID"),
    "unit_name": ("Unit", "Unit Name"),
    "status": ("Status",),
    "rent": ("Rent",),
    "deposit": ("Deposit", "Security Deposit"),
    "move_in": ("Move-in", "Move In", "Move-in Date"),
    "move_out": ("Move-out", "Move Out", "Move-out Date"),
    "lease_from": ("Lease From", "Lease Start"),
    "lease_to": ("Lease To", "Lease End"),
}
LEASE_OCCUPANCY_REQUIRED = (
    "source_tenant_id",
    "tenant_name",
    "source_property_id",
    "source_unit_id",
)


# Verified Owner Directory export contract. Name, Phone Numbers, Email,
# Properties Owned and Properties Owned IDs are documented report columns.
# A stable Owner ID is not guaranteed by the export, but preserve it when
# supplied rather than synthesizing source identity from contact data.
OWNER_ALIASES: dict[str, tuple[str, ...]] = {
    "source_id": ("Owner ID", "Owner Id", "OwnerID", "Id"),
    "name": ("Name", "Owner", "Owner Name"),
    "phone_numbers": ("Phone Numbers", "Phone Number", "Phone"),
    "email": ("Email", "Email Address", "Owner Email"),
    "properties_owned": ("Properties Owned",),
    "properties_owned_ids": (
        "Properties Owned IDs",
        "Properties Owned Ids",
        "Properties Owned ID",
        "Property IDs",
        "Property Ids",
    ),
}
OWNER_REQUIRED = ("name",)

# Verified Vendor Directory export contract. These fields remain source/staging
# facts only; 1099/compliance-expiration values are not tax/legal certification.
VENDOR_ALIASES: dict[str, tuple[str, ...]] = {
    "source_id": ("Vendor ID", "Vendor Id", "VendorID", "Id"),
    "company_name": ("Company Name", "Vendor", "Vendor Name"),
    "address": ("Address", "Vendor Address"),
    "phone_numbers": ("Phone Numbers", "Phone Number", "Phone"),
    "email": ("Email", "Email Address", "Vendor Email"),
    "send_1099": ("Send 1099?", "Send 1099", "Send1099"),
    "liability_insurance_expiration": (
        "Liability Insurance Expiration",
        "Liability Insurance Expiration Date",
    ),
    "workers_comp_expiration": (
        "Workers Comp Expiration",
        "Workers Compensation Expiration",
    ),
    "epa_certification_expiration": (
        "EPA Certification Expiration",
        "EPA Certification Expiration Date",
    ),
    "state_license_expiration": (
        "State License Expiration",
        "State License Expiration Date",
    ),
    "contract_expiration": ("Contract Expiration", "Contract Expiration Date"),
    "contact_name": ("Contact Name",),
    "contact_phone_numbers": ("Contact Phone Numbers", "Contact Phone"),
    "contact_email": ("Contact Email",),
}
VENDOR_REQUIRED = ("company_name",)

# Verified AppFolio General Ledger Accounts source contract. The current
# published AppFolio Stack resource exposes Number, Name, Type, FundAccount,
# IsCorporateAccount, OffsetAccountId, ParentGlAccountId, PropertyIds and
# LastUpdatedAt. A report export may also expose a stable GL Account ID; preserve
# it when supplied, but never synthesize source identity from account name.
GL_ACCOUNT_ALIASES: dict[str, tuple[str, ...]] = {
    "source_id": ("GL Account ID", "GL Account Id", "GlAccountId", "gl_account_id"),
    "account_number": ("Number", "Account Number", "GL Number", "GL Account Number"),
    "account_name": ("Name", "Account Name", "GL Account Name"),
    "account_type": ("Type", "Account Type"),
    "fund_account": ("FundAccount", "Fund Account"),
    "is_corporate_account": ("IsCorporateAccount", "Is Corporate Account"),
    "offset_account_id": ("OffsetAccountId", "Offset Account ID", "Offset Account Id"),
    "parent_gl_account_id": (
        "ParentGlAccountId",
        "Parent GL Account ID",
        "Parent GL Account Id",
    ),
    "property_ids": ("PropertyIds", "Property IDs", "Property Ids"),
    "last_updated_at": ("LastUpdatedAt", "Last Updated At"),
}
GL_ACCOUNT_REQUIRED = ("account_number", "account_name", "account_type")

# Verified AppFolio General Ledger Details source contract. Keep this source
# schema independent from our target accounting models. These fields are
# preserved as evidence only until later reconciliation/controlled-commit
# batches prove the accounting contract.
GENERAL_LEDGER_ALIASES: dict[str, tuple[str, ...]] = {
    "source_id": ("LineItemId", "Line Item ID", "Line Item Id"),
    "transaction_id": ("TransactionId", "Transaction ID", "Transaction Id"),
    "source_gl_account_id": ("GlAccountId", "GL Account ID", "GL Account Id"),
    "source_property_id": ("PropertyId", "Property ID", "Property Id"),
    "source_unit_id": ("UnitId", "Unit ID", "Unit Id"),
    "posted_date": ("Date", "Posted Date", "Post Date"),
    "debit": ("Debit",),
    "credit": ("Credit",),
    "description": ("Description",),
    "reference": ("Reference",),
    "remarks": ("Remarks",),
    "transaction_type": ("TransactionType", "Transaction Type"),
}
GENERAL_LEDGER_REQUIRED = ("source_gl_account_id", "posted_date", "debit", "credit")

# Verified AppFolio Bills top-level source contract. Structured LineItems are
# intentionally not flattened or inferred in this first staging batch.
BILL_ALIASES: dict[str, tuple[str, ...]] = {
    "source_id": ("Bill ID", "Bill Id", "BillId", "Id"),
    "source_vendor_id": ("VendorId", "Vendor ID", "Vendor Id"),
    "source_property_id": ("PropertyId", "Property ID", "Property Id"),
    "due_date": ("DueDate", "Due Date"),
    "invoice_date": ("InvoiceDate", "Invoice Date"),
    "posting_date": ("PostingDate", "Posting Date"),
    "reference": ("Reference",),
    "remarks": ("Remarks",),
    "total_amount": ("TotalAmount", "Total Amount", "Amount"),
    "approval_status": ("ApprovalStatus", "Approval Status"),
    "check_memo": ("CheckMemo", "Check Memo"),
    "account_number": ("AccountNumber", "Account Number"),
    "management_company_as_payee": (
        "ManagementCompanyAsPayee",
        "Management Company As Payee",
    ),
    "source_work_order_id": ("WorkOrderId", "Work Order ID", "Work Order Id"),
    "last_updated_at": ("LastUpdatedAt", "Last Updated At"),
}
# Keep source Bill ID out of the explicit-resource required set so a file that
# omits it can still be staged for REVIEW/SKIP. Auto-detection separately
# requires a stable Bill ID to avoid false positives.
BILL_REQUIRED = ("source_vendor_id", "due_date", "total_amount")

# Verified AppFolio Charges source contract from the public Stack API field
# inventory. AmountDue is preserved as the source's current outstanding amount;
# it is not promoted into target Charge.amount or payment state.
CHARGE_ALIASES: dict[str, tuple[str, ...]] = {
    "source_id": ("Charge ID", "Charge Id", "ChargeId", "Id"),
    "amount_due": ("AmountDue", "Amount Due"),
    "charged_on": ("ChargedOn", "Charged On", "Charge Date"),
    "description": ("Description",),
    "source_gl_account_id": ("GlAccountId", "GL Account ID", "GL Account Id"),
    "source_occupancy_id": ("OccupancyId", "Occupancy ID", "Occupancy Id"),
}
# Stable source Charge ID is deliberately excluded from explicit-resource
# required fields so an incomplete export may still be staged for REVIEW/SKIP.
# Auto-detection separately requires source_id to avoid false positives.
CHARGE_REQUIRED = (
    "amount_due",
    "charged_on",
    "source_gl_account_id",
    "source_occupancy_id",
)

# Verified AppFolio Stack Work Orders source contract. These fields are staged
# as source evidence only; this batch does not translate workflow state or
# create/update target WorkOrder records.
WORK_ORDER_ALIASES: dict[str, tuple[str, ...]] = {
    "source_id": ("Work Order ID", "Work Order Id", "WorkOrderId", "Id"),
    "source_property_id": ("PropertyId", "Property ID", "Property Id"),
    "source_unit_id": ("UnitId", "Unit ID", "Unit Id"),
    "status": ("Status", "Statuses"),
    "job_description": ("JobDescription", "Job Description"),
    "assigned_users": ("AssignedUsers", "Assigned Users"),
    "canceled_on": ("CanceledOn", "Canceled On"),
    "completed_on": ("CompletedOn", "Completed On"),
    "permission_to_enter": ("PermissionToEnter", "Permission To Enter"),
    "priority": ("Priority",),
    "scheduled_start": ("ScheduledStart", "Scheduled Start"),
    "scheduled_end": ("ScheduledEnd", "Scheduled End"),
    "source_vendor_id": ("VendorId", "Vendor ID", "Vendor Id"),
    "vendor_trade": ("VendorTrade", "Vendor Trade"),
}
# Stable Work Order ID is deliberately excluded so an explicitly selected,
# incomplete source file can still be staged for REVIEW/SKIP. Auto-detection
# separately requires source_id.
WORK_ORDER_REQUIRED = ("source_property_id", "status", "job_description")


class AppFolioFileIngestionError(ValueError):
    pass


@dataclass(frozen=True)
class ParsedTable:
    file_format: str
    sheet_name: str
    headers: list[str]
    rows: list[tuple[int, dict[str, Any]]]


@dataclass(frozen=True)
class StagedUploadResult:
    upload: PlatformMigrationUpload
    replayed: bool


def _normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def _clean_cell(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    text = str(value).strip()
    return text or None


def _safe_filename(filename: str) -> str:
    name = Path(filename).name.strip()
    if not name or len(name) > 255:
        raise AppFolioFileIngestionError("Source filename is missing or too long.")
    return name


def _validate_headers(headers: list[str]) -> None:
    if not headers:
        raise AppFolioFileIngestionError("Source file has no header row.")
    if len(headers) > MAX_COLUMNS:
        raise AppFolioFileIngestionError(
            f"Source file exceeds the {MAX_COLUMNS}-column limit."
        )
    normalized = [_normalize_header(header) for header in headers]
    if any(not key for key in normalized):
        raise AppFolioFileIngestionError("Blank source headers are not supported.")
    if len(set(normalized)) != len(normalized):
        raise AppFolioFileIngestionError(
            "Duplicate or normalization-colliding source headers require cleanup before staging."
        )
    forbidden = sorted(set(normalized) & SECRET_HEADER_KEYS)
    if forbidden:
        raise AppFolioFileIngestionError(
            "Credential/session fields are not accepted in migration source files."
        )


def _parse_csv(content: bytes) -> ParsedTable:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise AppFolioFileIngestionError("CSV files must be UTF-8 encoded.") from exc
    reader = csv.reader(io.StringIO(text, newline=""))
    try:
        raw_headers = next(reader)
    except StopIteration as exc:
        raise AppFolioFileIngestionError("CSV file is empty.") from exc
    headers = [str(value).strip() for value in raw_headers]
    _validate_headers(headers)
    rows: list[tuple[int, dict[str, Any]]] = []
    for row_number, values in enumerate(reader, start=2):
        if len(values) > MAX_COLUMNS:
            raise AppFolioFileIngestionError(
                f"CSV row {row_number} exceeds the {MAX_COLUMNS}-column limit."
            )
        padded = list(values) + [""] * (len(headers) - len(values))
        if len(padded) > len(headers):
            raise AppFolioFileIngestionError(
                f"CSV row {row_number} has more values than the header row."
            )
        if not any(str(value).strip() for value in padded):
            continue
        rows.append(
            (
                row_number,
                {
                    headers[index]: _clean_cell(value)
                    for index, value in enumerate(padded)
                },
            )
        )
        if len(rows) > MAX_ROWS:
            raise AppFolioFileIngestionError(
                f"Source file exceeds the {MAX_ROWS}-row staging limit."
            )
    return ParsedTable(file_format="CSV", sheet_name="CSV", headers=headers, rows=rows)


def _xlsx_sheet_headers(ws) -> list[str]:
    first = next(ws.iter_rows(min_row=1, max_row=1), ())
    if any(getattr(cell, "data_type", None) == "f" for cell in first):
        raise AppFolioFileIngestionError("Spreadsheet formula cells are not accepted.")
    return [str(cell.value or "").strip() for cell in first]


def _looks_like_properties(headers: list[str]) -> bool:
    normalized = {_normalize_header(header) for header in headers}
    matched = 0
    for field in PROPERTY_REQUIRED:
        aliases = {_normalize_header(alias) for alias in PROPERTY_ALIASES[field]}
        if normalized & aliases:
            matched += 1
    return matched == len(PROPERTY_REQUIRED)


def _looks_like_units(headers: list[str]) -> bool:
    normalized = {_normalize_header(header) for header in headers}
    unit_name = {
        _normalize_header(alias) for alias in UNIT_ALIASES["unit_name"]
    }
    property_ref = {
        _normalize_header(alias)
        for field in ("source_property_id", "property_name")
        for alias in UNIT_ALIASES[field]
    }
    address_ref = {
        _normalize_header(alias)
        for field in ("unit_address", "address_line1")
        for alias in UNIT_ALIASES[field]
    }
    return bool(normalized & unit_name) and bool(normalized & property_ref) and bool(
        normalized & address_ref
    )


def _looks_like_tenants(headers: list[str]) -> bool:
    normalized = {_normalize_header(header) for header in headers}
    tenant_aliases = {
        _normalize_header(alias) for alias in TENANT_ALIASES["tenant_name"]
    }
    tenant_id_aliases = {
        _normalize_header(alias) for alias in TENANT_ALIASES["source_id"]
    }
    unit_id_aliases = {
        _normalize_header(alias) for alias in TENANT_ALIASES["source_unit_id"]
    }
    property_id_aliases = {
        _normalize_header(alias) for alias in TENANT_ALIASES["source_property_id"]
    }
    return (
        bool(normalized & tenant_aliases)
        and bool(normalized & tenant_id_aliases)
        and bool(normalized & unit_id_aliases)
        and bool(normalized & property_id_aliases)
    )


def _single_source_email(value: Any) -> str | None:
    text = str(value).strip() if value is not None else ""
    if not text or not re.fullmatch(r"[^@\s,;]+@[^@\s,;]+\.[^@\s,;]+", text):
        return None
    return text.lower()


def _looks_like_owners(headers: list[str]) -> bool:
    normalized = {_normalize_header(header) for header in headers}
    name_aliases = {_normalize_header(alias) for alias in OWNER_ALIASES["name"]}
    email_aliases = {_normalize_header(alias) for alias in OWNER_ALIASES["email"]}
    property_id_aliases = {
        _normalize_header(alias) for alias in OWNER_ALIASES["properties_owned_ids"]
    }
    return (
        bool(normalized & name_aliases)
        and bool(normalized & email_aliases)
        and bool(normalized & property_id_aliases)
    )


def _looks_like_vendors(headers: list[str]) -> bool:
    normalized = {_normalize_header(header) for header in headers}
    company_aliases = {
        _normalize_header(alias) for alias in VENDOR_ALIASES["company_name"]
    }
    email_aliases = {_normalize_header(alias) for alias in VENDOR_ALIASES["email"]}
    send_1099_aliases = {
        _normalize_header(alias) for alias in VENDOR_ALIASES["send_1099"]
    }
    return (
        bool(normalized & company_aliases)
        and bool(normalized & email_aliases)
        and bool(normalized & send_1099_aliases)
    )


def _looks_like_gl_accounts(headers: list[str]) -> bool:
    normalized = {_normalize_header(header) for header in headers}
    return all(
        normalized
        & {
            _normalize_header(alias)
            for alias in GL_ACCOUNT_ALIASES[field]
        }
        for field in GL_ACCOUNT_REQUIRED
    )


def _looks_like_general_ledger(headers: list[str]) -> bool:
    normalized = {_normalize_header(header) for header in headers}
    return all(
        normalized
        & {
            _normalize_header(alias)
            for alias in GENERAL_LEDGER_ALIASES[field]
        }
        for field in GENERAL_LEDGER_REQUIRED
    )


def _looks_like_bills(headers: list[str]) -> bool:
    normalized = {_normalize_header(header) for header in headers}
    required_fields = ("source_id",) + BILL_REQUIRED
    return all(
        normalized
        & {
            _normalize_header(alias)
            for alias in BILL_ALIASES[field]
        }
        for field in required_fields
    )


def _looks_like_charges(headers: list[str]) -> bool:
    normalized = {_normalize_header(header) for header in headers}
    required_fields = ("source_id",) + CHARGE_REQUIRED
    return all(
        normalized
        & {
            _normalize_header(alias)
            for alias in CHARGE_ALIASES[field]
        }
        for field in required_fields
    )


def _looks_like_work_orders(headers: list[str]) -> bool:
    normalized = {_normalize_header(header) for header in headers}
    required_fields = ("source_id",) + WORK_ORDER_REQUIRED
    return all(
        normalized
        & {
            _normalize_header(alias)
            for alias in WORK_ORDER_ALIASES[field]
        }
        for field in required_fields
    )


def _parse_xlsx(content: bytes, requested_sheet: str | None) -> ParsedTable:
    stream = io.BytesIO(content)
    if not zipfile.is_zipfile(stream):
        raise AppFolioFileIngestionError("XLSX source is not a valid workbook archive.")
    stream.seek(0)
    with zipfile.ZipFile(stream) as archive:
        infos = archive.infolist()
        if len(infos) > MAX_XLSX_ENTRIES:
            raise AppFolioFileIngestionError("XLSX archive contains too many entries.")
        if sum(info.file_size for info in infos) > MAX_XLSX_UNCOMPRESSED_BYTES:
            raise AppFolioFileIngestionError("XLSX expanded content exceeds the safety limit.")
        lowered = {info.filename.lower() for info in infos}
        if any(name.endswith("vbaproject.bin") for name in lowered):
            raise AppFolioFileIngestionError("Macro-enabled spreadsheet content is not accepted.")
        try:
            content_types = archive.read("[Content_Types].xml").lower()
        except KeyError:
            content_types = b""
        if b"macroenabled" in content_types:
            raise AppFolioFileIngestionError("Macro-enabled spreadsheet content is not accepted.")

    stream.seek(0)
    try:
        workbook = load_workbook(
            stream,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except Exception as exc:
        raise AppFolioFileIngestionError("XLSX workbook could not be parsed safely.") from exc
    try:
        if len(workbook.sheetnames) > MAX_SHEETS:
            raise AppFolioFileIngestionError(
                f"Workbook exceeds the {MAX_SHEETS}-sheet limit."
            )
        if requested_sheet:
            if requested_sheet not in workbook.sheetnames:
                raise AppFolioFileIngestionError("Requested workbook sheet was not found.")
            selected = requested_sheet
        elif len(workbook.sheetnames) == 1:
            selected = workbook.sheetnames[0]
        else:
            candidates = []
            for name in workbook.sheetnames:
                headers = _xlsx_sheet_headers(workbook[name])
                if headers and (
                    _looks_like_properties(headers)
                    or _looks_like_units(headers)
                    or _looks_like_tenants(headers)
                    or _looks_like_owners(headers)
                    or _looks_like_vendors(headers)
                    or _looks_like_gl_accounts(headers)
                    or _looks_like_general_ledger(headers)
                    or _looks_like_bills(headers)
                    or _looks_like_charges(headers)
                    or _looks_like_work_orders(headers)
                ):
                    candidates.append(name)
            if len(candidates) != 1:
                raise AppFolioFileIngestionError(
                    "Workbook contains multiple sheets; specify sheet_name unless exactly one supported report can be detected."
                )
            selected = candidates[0]

        ws = workbook[selected]
        headers = _xlsx_sheet_headers(ws)
        _validate_headers(headers)
        rows: list[tuple[int, dict[str, Any]]] = []
        for row_number, cells in enumerate(ws.iter_rows(min_row=2), start=2):
            if len(cells) > MAX_COLUMNS:
                raise AppFolioFileIngestionError(
                    f"Workbook row {row_number} exceeds the {MAX_COLUMNS}-column limit."
                )
            if any(getattr(cell, "data_type", None) == "f" for cell in cells):
                raise AppFolioFileIngestionError(
                    f"Spreadsheet formula cell found on row {row_number}; formulas are never evaluated."
                )
            values = [_clean_cell(cell.value) for cell in cells[: len(headers)]]
            values += [None] * (len(headers) - len(values))
            if not any(value not in (None, "") for value in values):
                continue
            rows.append(
                (
                    row_number,
                    {
                        headers[index]: values[index]
                        for index in range(len(headers))
                    },
                )
            )
            if len(rows) > MAX_ROWS:
                raise AppFolioFileIngestionError(
                    f"Source file exceeds the {MAX_ROWS}-row staging limit."
                )
        return ParsedTable(
            file_format="XLSX",
            sheet_name=selected,
            headers=headers,
            rows=rows,
        )
    finally:
        workbook.close()


def _parse_file(filename: str, content: bytes, sheet_name: str | None) -> ParsedTable:
    if len(content) > MAX_FILE_BYTES:
        raise AppFolioFileIngestionError(
            f"Source file exceeds the {MAX_FILE_BYTES // (1024 * 1024)} MB upload limit."
        )
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        if sheet_name:
            raise AppFolioFileIngestionError("sheet_name is only valid for XLSX uploads.")
        return _parse_csv(content)
    if suffix == ".xlsx":
        return _parse_xlsx(content, sheet_name)
    if suffix in {".xlsm", ".xls", ".xltm"}:
        raise AppFolioFileIngestionError(
            "Macro-enabled or legacy Excel formats are not accepted; export a data-only XLSX or CSV."
        )
    raise AppFolioFileIngestionError("Only CSV and XLSX source files are supported.")


def _auto_mapping(
    headers: list[str],
    aliases_by_field: dict[str, tuple[str, ...]],
) -> tuple[dict[str, str], list[str]]:
    by_normalized = {_normalize_header(header): header for header in headers}
    mapping: dict[str, str] = {}
    ambiguous: list[str] = []
    for field, aliases in aliases_by_field.items():
        matches = []
        for alias in aliases:
            header = by_normalized.get(_normalize_header(alias))
            if header is not None and header not in matches:
                matches.append(header)
        if len(matches) == 1:
            mapping[field] = matches[0]
        elif len(matches) > 1:
            ambiguous.append(field)
    return mapping, ambiguous


def _resolve_mapping(
    headers: list[str],
    *,
    resource_override: str | None,
    explicit_mapping: dict[str, str] | None,
) -> tuple[str, dict[str, str], list[str], list[str]]:
    requested = (resource_override or "").strip().upper()
    if requested not in {"", "PROPERTIES", "UNITS", "TENANTS", "LEASE_OCCUPANCY", "OWNERS", "VENDORS", "GL_ACCOUNTS", "GENERAL_LEDGER", "BILLS", "CHARGES", "WORK_ORDERS"}:
        raise AppFolioFileIngestionError(
            "This migration stage currently supports only verified PROPERTIES, UNITS, TENANTS, LEASE_OCCUPANCY, OWNERS, VENDORS, GL_ACCOUNTS, GENERAL_LEDGER, BILLS, CHARGES and WORK_ORDERS resource mappings."
        )

    if requested == "PROPERTIES":
        resource = "PROPERTIES"
    elif requested == "UNITS":
        resource = "UNITS"
    elif requested == "TENANTS":
        resource = "TENANTS"
    elif requested == "LEASE_OCCUPANCY":
        resource = "LEASE_OCCUPANCY"
    elif requested == "OWNERS":
        resource = "OWNERS"
    elif requested == "VENDORS":
        resource = "VENDORS"
    elif requested == "GL_ACCOUNTS":
        resource = "GL_ACCOUNTS"
    elif requested == "GENERAL_LEDGER":
        resource = "GENERAL_LEDGER"
    elif requested == "BILLS":
        resource = "BILLS"
    elif requested == "CHARGES":
        resource = "CHARGES"
    elif requested == "WORK_ORDERS":
        resource = "WORK_ORDERS"
    elif _looks_like_properties(headers):
        resource = "PROPERTIES"
    elif _looks_like_units(headers):
        resource = "UNITS"
    elif _looks_like_tenants(headers):
        resource = "TENANTS"
    elif _looks_like_owners(headers):
        resource = "OWNERS"
    elif _looks_like_vendors(headers):
        resource = "VENDORS"
    elif _looks_like_gl_accounts(headers):
        resource = "GL_ACCOUNTS"
    elif _looks_like_general_ledger(headers):
        resource = "GENERAL_LEDGER"
    elif _looks_like_bills(headers):
        resource = "BILLS"
    elif _looks_like_charges(headers):
        resource = "CHARGES"
    elif _looks_like_work_orders(headers):
        resource = "WORK_ORDERS"
    else:
        resource = "UNKNOWN"

    aliases = (
        PROPERTY_ALIASES
        if resource == "PROPERTIES"
        else UNIT_ALIASES
        if resource == "UNITS"
        else TENANT_ALIASES
        if resource == "TENANTS"
        else LEASE_OCCUPANCY_ALIASES
        if resource == "LEASE_OCCUPANCY"
        else OWNER_ALIASES
        if resource == "OWNERS"
        else VENDOR_ALIASES
        if resource == "VENDORS"
        else GL_ACCOUNT_ALIASES
        if resource == "GL_ACCOUNTS"
        else GENERAL_LEDGER_ALIASES
        if resource == "GENERAL_LEDGER"
        else BILL_ALIASES
        if resource == "BILLS"
        else CHARGE_ALIASES
        if resource == "CHARGES"
        else WORK_ORDER_ALIASES
        if resource == "WORK_ORDERS"
        else {}
    )
    required = (
        PROPERTY_REQUIRED
        if resource == "PROPERTIES"
        else UNIT_REQUIRED
        if resource == "UNITS"
        else TENANT_REQUIRED
        if resource == "TENANTS"
        else LEASE_OCCUPANCY_REQUIRED
        if resource == "LEASE_OCCUPANCY"
        else OWNER_REQUIRED
        if resource == "OWNERS"
        else VENDOR_REQUIRED
        if resource == "VENDORS"
        else GL_ACCOUNT_REQUIRED
        if resource == "GL_ACCOUNTS"
        else GENERAL_LEDGER_REQUIRED
        if resource == "GENERAL_LEDGER"
        else BILL_REQUIRED
        if resource == "BILLS"
        else CHARGE_REQUIRED
        if resource == "CHARGES"
        else WORK_ORDER_REQUIRED
        if resource == "WORK_ORDERS"
        else ()
    )
    auto_mapping, ambiguous = _auto_mapping(headers, aliases) if aliases else ({}, [])
    mapping = dict(auto_mapping)

    if explicit_mapping:
        allowed_fields = set(aliases)
        if not allowed_fields:
            # A caller must choose a supported resource before using explicit
            # mapping when automatic report detection cannot determine one.
            raise AppFolioFileIngestionError(
                "Choose resource PROPERTIES, UNITS, TENANTS, LEASE_OCCUPANCY, OWNERS, VENDORS, GL_ACCOUNTS, GENERAL_LEDGER, BILLS, CHARGES or WORK_ORDERS before supplying explicit column mapping."
            )
        unknown_fields = sorted(set(explicit_mapping) - allowed_fields)
        if unknown_fields:
            raise AppFolioFileIngestionError(
                "Unknown explicit mapping fields: " + ", ".join(unknown_fields)
            )
        header_set = set(headers)
        missing_headers = sorted(
            {header for header in explicit_mapping.values() if header not in header_set}
        )
        if missing_headers:
            raise AppFolioFileIngestionError(
                "Explicit mapping references unknown source headers: "
                + ", ".join(missing_headers)
            )
        if len(set(explicit_mapping.values())) != len(explicit_mapping.values()):
            raise AppFolioFileIngestionError(
                "A source column cannot be explicitly mapped to multiple target fields."
            )
        mapping.update(explicit_mapping)
        ambiguous = [field for field in ambiguous if field not in explicit_mapping]

    missing_required = [field for field in required if field not in mapping]
    return resource, mapping, ambiguous, missing_required


def _normalized_source_row(row: dict[str, Any]) -> dict[str, Any]:
    return {
        _normalize_header(key): _clean_cell(value)
        for key, value in row.items()
    }


def _fingerprint(
    *,
    run: PlatformMigrationRun,
    parsed: ParsedTable,
    resource: str,
    mapping: dict[str, str],
) -> str:
    canonical = {
        "provider": "APPFOLIO",
        "organization_id": run.organization_id,
        "source_account_ref": run.source_account_ref,
        "resource": resource,
        "sheet_name": parsed.sheet_name,
        "mapping": {
            key: _normalize_header(value)
            for key, value in sorted(mapping.items())
        },
        "rows": [
            {
                "row_number": row_number,
                "data": _normalized_source_row(row),
            }
            for row_number, row in parsed.rows
        ],
    }
    return hashlib.sha256(
        json.dumps(canonical, sort_keys=True, separators=(",", ":"), default=str).encode("utf-8")
    ).hexdigest()


def _row_fingerprint(row: dict[str, Any]) -> str:
    return hashlib.sha256(
        json.dumps(
            _normalized_source_row(row),
            sort_keys=True,
            separators=(",", ":"),
            default=str,
        ).encode("utf-8")
    ).hexdigest()


def _mapped_value(row: dict[str, Any], mapping: dict[str, str], field: str) -> Any:
    source = mapping.get(field)
    return _clean_cell(row.get(source)) if source else None


def stage_appfolio_file(
    db: Session,
    *,
    run: PlatformMigrationRun,
    filename: str,
    content: bytes,
    resource_override: str | None,
    sheet_name: str | None,
    explicit_mapping: dict[str, str] | None,
    platform_user_id: int,
) -> StagedUploadResult:
    if run.provider != "APPFOLIO":
        raise AppFolioFileIngestionError("Migration run is not an AppFolio run.")
    safe_filename = _safe_filename(filename)
    parsed = _parse_file(safe_filename, content, sheet_name)
    resource, mapping, ambiguous, missing_required = _resolve_mapping(
        parsed.headers,
        resource_override=resource_override,
        explicit_mapping=explicit_mapping,
    )
    fingerprint = _fingerprint(
        run=run,
        parsed=parsed,
        resource=resource,
        mapping=mapping,
    )
    replay = (
        db.query(PlatformMigrationUpload)
        .filter(
            PlatformMigrationUpload.run_id == run.id,
            PlatformMigrationUpload.organization_id == run.organization_id,
            PlatformMigrationUpload.provider == "APPFOLIO",
            PlatformMigrationUpload.normalized_fingerprint == fingerprint,
        )
        .first()
    )
    if replay is not None:
        return StagedUploadResult(upload=replay, replayed=True)

    file_sha256 = hashlib.sha256(content).hexdigest()
    upload = PlatformMigrationUpload(
        run_id=run.id,
        organization_id=run.organization_id,
        provider="APPFOLIO",
        filename=safe_filename,
        file_format=parsed.file_format,
        file_sha256=file_sha256,
        normalized_fingerprint=fingerprint,
        detected_resource=resource,
        sheet_name=parsed.sheet_name,
        headers=parsed.headers,
        column_mapping=mapping,
        validation_summary={},
        status="STAGING",
        row_count=len(parsed.rows),
        created_by_platform_user_id=platform_user_id,
    )
    db.add(upload)
    db.flush()

    valid = 0
    warning_rows = 0
    invalid = 0
    duplicates = 0
    possible_matches = 0
    seen_source_ids: set[str] = set()
    seen_lease_occupancy_keys: set[tuple[str, str, str, str, str, str, str]] = set()
    seen_gl_account_keys: set[tuple[str, str, str]] = set()
    seen_general_ledger_line_ids: set[str] = set()

    for row_number, source_row in parsed.rows:
        errors: list[str] = []
        warnings: list[str] = []
        source_id = None
        disposition = "REVIEW"
        normalized_data: dict[str, Any]

        if resource == "PROPERTIES":
            normalized_data = {
                field: _mapped_value(source_row, mapping, field)
                for field in PROPERTY_ALIASES
                if field in mapping
            }
            if ambiguous:
                errors.append(
                    "Ambiguous automatic mapping requires explicit mapping for: "
                    + ", ".join(sorted(ambiguous))
                )
            if missing_required:
                errors.append(
                    "Missing required source columns: " + ", ".join(missing_required)
                )

            source_id_value = normalized_data.get("source_id")
            source_id = str(source_id_value).strip() if source_id_value is not None else None
            if not source_id:
                errors.append("source_id is required.")
            elif source_id in seen_source_ids:
                errors.append("Duplicate AppFolio source ID in this staged upload.")
                duplicates += 1
            else:
                seen_source_ids.add(source_id)

            for field in ("name", "address_line1", "city", "state", "zip_code"):
                value = normalized_data.get(field)
                if value is None or not str(value).strip():
                    errors.append(f"{field} is required.")

            if errors:
                disposition = "INVALID"
                invalid += 1
            else:
                valid += 1
                mapped_item = (
                    db.query(PlatformMigrationItem)
                    .filter(
                        PlatformMigrationItem.run_id == run.id,
                        PlatformMigrationItem.organization_id == run.organization_id,
                        PlatformMigrationItem.provider == "APPFOLIO",
                        PlatformMigrationItem.resource == "PROPERTIES",
                        PlatformMigrationItem.source_id == source_id,
                    )
                    .first()
                )
                if mapped_item is not None:
                    disposition = "ALREADY_MAPPED"
                    warnings.append(
                        f"Source ID is already mapped to {mapped_item.target_entity} #{mapped_item.target_id}."
                    )
                else:
                    existing = (
                        db.query(Property)
                        .filter(
                            Property.organization_id == run.organization_id,
                            Property.name == normalized_data.get("name"),
                            Property.address_line1 == normalized_data.get("address_line1"),
                            Property.city == normalized_data.get("city"),
                            Property.state == normalized_data.get("state"),
                            Property.zip_code == normalized_data.get("zip_code"),
                        )
                        .first()
                    )
                    if existing is not None:
                        disposition = "POSSIBLE_MATCH"
                        possible_matches += 1
                        warnings.append(
                            f"Possible existing target property match: local property #{existing.id}."
                        )
                    elif normalized_data.get("hidden_at") not in (None, "", False, 0):
                        disposition = "REVIEW"
                        warnings.append(
                            "Source property is marked hidden; review inclusion before dry run."
                        )
                    else:
                        disposition = "NEW"
        elif resource == "UNITS":
            normalized_data = {
                field: _mapped_value(source_row, mapping, field)
                for field in UNIT_ALIASES
                if field in mapping
            }
            if ambiguous:
                errors.append(
                    "Ambiguous automatic mapping requires explicit mapping for: "
                    + ", ".join(sorted(ambiguous))
                )
            if missing_required:
                errors.append(
                    "Missing required source columns: " + ", ".join(missing_required)
                )

            unit_name = normalized_data.get("unit_name")
            if unit_name is None or not str(unit_name).strip():
                errors.append("unit_name is required.")

            source_id_value = normalized_data.get("source_id")
            source_id = str(source_id_value).strip() if source_id_value is not None else None
            if source_id:
                if source_id in seen_source_ids:
                    errors.append("Duplicate AppFolio Unit ID in this staged upload.")
                    duplicates += 1
                else:
                    seen_source_ids.add(source_id)

            source_property_value = normalized_data.get("source_property_id")
            source_property_id = (
                str(source_property_value).strip()
                if source_property_value is not None
                else None
            )

            mapped_property = None
            if source_property_id:
                mapped_property = (
                    db.query(PlatformMigrationItem)
                    .filter(
                        PlatformMigrationItem.run_id == run.id,
                        PlatformMigrationItem.organization_id == run.organization_id,
                        PlatformMigrationItem.provider == "APPFOLIO",
                        PlatformMigrationItem.resource == "PROPERTIES",
                        PlatformMigrationItem.source_id == source_property_id,
                    )
                    .first()
                )
                if mapped_property is not None:
                    target_property = (
                        db.query(Property)
                        .filter(
                            Property.id == mapped_property.target_id,
                            Property.organization_id == run.organization_id,
                            Property.is_active.is_(True),
                            Property.deleted_at.is_(None),
                        )
                        .first()
                    )
                    if (
                        mapped_property.target_entity != "PROPERTY"
                        or target_property is None
                    ):
                        errors.append(
                            "Mapped source Property ID does not resolve to an active same-organization Property."
                        )

            if errors:
                disposition = "INVALID"
                invalid += 1
            else:
                valid += 1
                mapped_unit = None
                if source_id:
                    mapped_unit = (
                        db.query(PlatformMigrationItem)
                        .filter(
                            PlatformMigrationItem.run_id == run.id,
                            PlatformMigrationItem.organization_id == run.organization_id,
                            PlatformMigrationItem.provider == "APPFOLIO",
                            PlatformMigrationItem.resource == "UNITS",
                            PlatformMigrationItem.source_id == source_id,
                        )
                        .first()
                    )
                if mapped_unit is not None:
                    disposition = "ALREADY_MAPPED"
                    warnings.append(
                        f"Source Unit ID is already mapped to {mapped_unit.target_entity} #{mapped_unit.target_id}."
                    )
                elif not source_property_id:
                    disposition = "REVIEW"
                    warnings.append(
                        "Property ID was not supplied; Unit-to-Property linkage must be resolved before dry run or commit."
                    )
                elif mapped_property is None:
                    disposition = "REVIEW"
                    warnings.append(
                        f"Source Property ID {source_property_id} has no durable Property mapping in this migration run."
                    )
                elif not source_id:
                    disposition = "REVIEW"
                    warnings.append(
                        "Unit ID was not supplied; durable Unit source identity must be resolved before dry run or commit."
                    )
                else:
                    exact = (
                        db.query(Unit)
                        .filter(
                            Unit.property_id == target_property.id,
                            Unit.unit_number == str(unit_name).strip(),
                            Unit.is_active.is_(True),
                            Unit.deleted_at.is_(None),
                        )
                        .first()
                    )
                    candidate = exact
                    if candidate is None:
                        candidate = (
                            db.query(Unit)
                            .filter(
                                Unit.property_id == target_property.id,
                                func.lower(Unit.unit_number)
                                == str(unit_name).strip().lower(),
                                Unit.is_active.is_(True),
                                Unit.deleted_at.is_(None),
                            )
                            .first()
                        )
                    if candidate is not None:
                        disposition = "POSSIBLE_MATCH"
                        possible_matches += 1
                        match_kind = (
                            "exact unit-number"
                            if candidate.unit_number == str(unit_name).strip()
                            else "case-insensitive unit-number"
                        )
                        warnings.append(
                            f"Possible existing target unit match: local unit #{candidate.id} ({match_kind})."
                        )
                    else:
                        disposition = "NEW"
        elif resource == "TENANTS":
            normalized_data = {
                field: _mapped_value(source_row, mapping, field)
                for field in TENANT_ALIASES
                if field in mapping
            }
            if ambiguous:
                errors.append(
                    "Ambiguous automatic mapping requires explicit mapping for: "
                    + ", ".join(sorted(ambiguous))
                )
            if missing_required:
                errors.append(
                    "Missing required source columns: " + ", ".join(missing_required)
                )

            tenant_name_value = normalized_data.get("tenant_name")
            tenant_name = (
                str(tenant_name_value).strip()
                if tenant_name_value is not None
                else None
            )
            if not tenant_name:
                errors.append("tenant_name is required.")

            source_id_value = normalized_data.get("source_id")
            source_id = str(source_id_value).strip() if source_id_value is not None else None
            if source_id:
                if source_id in seen_source_ids:
                    errors.append("Duplicate AppFolio Tenant ID in this staged upload.")
                    duplicates += 1
                else:
                    seen_source_ids.add(source_id)

            source_unit_value = normalized_data.get("source_unit_id")
            source_unit_id = (
                str(source_unit_value).strip()
                if source_unit_value is not None
                else None
            )
            source_property_value = normalized_data.get("source_property_id")
            source_property_id = (
                str(source_property_value).strip()
                if source_property_value is not None
                else None
            )

            mapped_unit = None
            target_unit = None
            if source_unit_id:
                mapped_unit = (
                    db.query(PlatformMigrationItem)
                    .filter(
                        PlatformMigrationItem.run_id == run.id,
                        PlatformMigrationItem.organization_id == run.organization_id,
                        PlatformMigrationItem.provider == "APPFOLIO",
                        PlatformMigrationItem.resource == "UNITS",
                        PlatformMigrationItem.source_id == source_unit_id,
                    )
                    .first()
                )
                if mapped_unit is not None:
                    target_unit = (
                        db.query(Unit)
                        .join(Property, Property.id == Unit.property_id)
                        .filter(
                            Unit.id == mapped_unit.target_id,
                            Unit.is_active.is_(True),
                            Unit.deleted_at.is_(None),
                            Property.organization_id == run.organization_id,
                            Property.is_active.is_(True),
                            Property.deleted_at.is_(None),
                        )
                        .first()
                    )
                    if mapped_unit.target_entity != "UNIT" or target_unit is None:
                        errors.append(
                            "Mapped source Unit ID does not resolve to an active same-organization Unit."
                        )

            mapped_property = None
            target_property = None
            if source_property_id:
                mapped_property = (
                    db.query(PlatformMigrationItem)
                    .filter(
                        PlatformMigrationItem.run_id == run.id,
                        PlatformMigrationItem.organization_id == run.organization_id,
                        PlatformMigrationItem.provider == "APPFOLIO",
                        PlatformMigrationItem.resource == "PROPERTIES",
                        PlatformMigrationItem.source_id == source_property_id,
                    )
                    .first()
                )
                if mapped_property is not None:
                    target_property = (
                        db.query(Property)
                        .filter(
                            Property.id == mapped_property.target_id,
                            Property.organization_id == run.organization_id,
                            Property.is_active.is_(True),
                            Property.deleted_at.is_(None),
                        )
                        .first()
                    )
                    if mapped_property.target_entity != "PROPERTY" or target_property is None:
                        errors.append(
                            "Mapped source Property ID does not resolve to an active same-organization Property."
                        )

            if (
                target_unit is not None
                and target_property is not None
                and target_unit.property_id != target_property.id
            ):
                errors.append(
                    "Mapped Tenant Unit and Property source IDs resolve to different target Properties."
                )

            if errors:
                disposition = "INVALID"
                invalid += 1
            else:
                valid += 1
                mapped_tenant = None
                if source_id:
                    mapped_tenant = (
                        db.query(PlatformMigrationItem)
                        .filter(
                            PlatformMigrationItem.run_id == run.id,
                            PlatformMigrationItem.organization_id == run.organization_id,
                            PlatformMigrationItem.provider == "APPFOLIO",
                            PlatformMigrationItem.resource == "TENANTS",
                            PlatformMigrationItem.source_id == source_id,
                        )
                        .first()
                    )

                if mapped_tenant is not None:
                    disposition = "ALREADY_MAPPED"
                    warnings.append(
                        f"Source Tenant ID is already mapped to {mapped_tenant.target_entity} #{mapped_tenant.target_id}."
                    )
                else:
                    review_required = False
                    if not source_id:
                        review_required = True
                        warnings.append(
                            "Tenant ID was not supplied; durable Tenant source identity must be resolved before any future commit."
                        )
                    if not source_unit_id:
                        review_required = True
                        warnings.append(
                            "Unit ID was not supplied; Tenant-to-Unit relationship cannot be resolved safely."
                        )
                    elif mapped_unit is None:
                        review_required = True
                        warnings.append(
                            "Source Unit ID has no durable Unit mapping yet; relationship remains review-only."
                        )
                    if not source_property_id:
                        review_required = True
                        warnings.append(
                            "Property ID was not supplied; Tenant-to-Property relationship cannot be resolved safely."
                        )
                    elif mapped_property is None:
                        review_required = True
                        warnings.append(
                            "Source Property ID has no durable Property mapping yet; relationship remains review-only."
                        )

                    exact_email = _single_source_email(normalized_data.get("emails"))
                    candidate = None
                    if exact_email:
                        candidate = (
                            db.query(User)
                            .filter(
                                User.organization_id == run.organization_id,
                                User.role == UserRole.TENANT,
                                User.is_active.is_(True),
                                User.deleted_at.is_(None),
                                func.lower(User.email) == exact_email,
                            )
                            .first()
                        )
                    if candidate is not None:
                        possible_matches += 1
                        warnings.append(
                            f"Possible existing target tenant match by one exact source email: local tenant #{candidate.id}; explicit review is required before any future commit."
                        )
                        disposition = "REVIEW" if review_required else "POSSIBLE_MATCH"
                    elif review_required:
                        disposition = "REVIEW"
                    else:
                        disposition = "NEW"

                    if target_unit is not None and target_property is not None:
                        warnings.append(
                            "Unit/Property source IDs resolve consistently, but this staging batch does not establish occupancy or lease liability."
                        )
                    if normalized_data.get("move_in") or normalized_data.get("move_out"):
                        warnings.append(
                            "Move-in/move-out values are preserved as source evidence only; no occupancy event is created."
                        )
                    if normalized_data.get("lease_from") or normalized_data.get("lease_to"):
                        warnings.append(
                            "Lease From/To values are preserved as source evidence only; no Lease record is created."
                        )
        elif resource == "LEASE_OCCUPANCY":
            normalized_data = {
                field: _mapped_value(source_row, mapping, field)
                for field in LEASE_OCCUPANCY_ALIASES
                if field in mapping
            }
            if ambiguous:
                errors.append(
                    "Ambiguous automatic mapping requires explicit mapping for: "
                    + ", ".join(sorted(ambiguous))
                )
            if missing_required:
                errors.append(
                    "Missing required source columns: " + ", ".join(missing_required)
                )

            tenant_source = _clean_cell(normalized_data.get("source_tenant_id"))
            property_source = _clean_cell(normalized_data.get("source_property_id"))
            unit_source = _clean_cell(normalized_data.get("source_unit_id"))
            tenant_name = _clean_cell(normalized_data.get("tenant_name"))

            if not tenant_name:
                errors.append("tenant_name is required.")

            missing_identity = False
            for value, label in (
                (tenant_source, "Tenant ID"),
                (property_source, "Property ID"),
                (unit_source, "Unit ID"),
            ):
                if not value:
                    missing_identity = True
                    warnings.append(
                        f"{label} was not supplied; Lease/occupancy relationship identity remains review-only."
                    )

            if tenant_source and property_source and unit_source:
                duplicate_key = (
                    str(tenant_source),
                    str(property_source),
                    str(unit_source),
                    str(normalized_data.get("lease_from") or ""),
                    str(normalized_data.get("lease_to") or ""),
                    str(normalized_data.get("move_in") or ""),
                    str(normalized_data.get("move_out") or ""),
                )
                if duplicate_key in seen_lease_occupancy_keys:
                    errors.append("Duplicate Lease/occupancy source relationship row in this staged upload.")
                    duplicates += 1
                else:
                    seen_lease_occupancy_keys.add(duplicate_key)

            tenant_mapping = None
            target_tenant = None
            if tenant_source:
                tenant_mapping = (
                    db.query(PlatformMigrationItem)
                    .filter(
                        PlatformMigrationItem.run_id == run.id,
                        PlatformMigrationItem.organization_id == run.organization_id,
                        PlatformMigrationItem.provider == "APPFOLIO",
                        PlatformMigrationItem.resource == "TENANTS",
                        PlatformMigrationItem.source_id == str(tenant_source),
                    )
                    .first()
                )
                if tenant_mapping is not None:
                    if tenant_mapping.target_entity != "TENANT_USER":
                        errors.append("Mapped source Tenant ID has an inconsistent target type.")
                    else:
                        target_tenant = (
                            db.query(User)
                            .filter(
                                User.id == tenant_mapping.target_id,
                                User.organization_id == run.organization_id,
                                User.role == UserRole.TENANT,
                                User.is_active.is_(True),
                                User.deleted_at.is_(None),
                            )
                            .first()
                        )
                        if target_tenant is None:
                            errors.append("Mapped source Tenant ID does not resolve to an active same-organization TENANT user.")

            property_mapping = None
            target_property = None
            if property_source:
                property_mapping = (
                    db.query(PlatformMigrationItem)
                    .filter(
                        PlatformMigrationItem.run_id == run.id,
                        PlatformMigrationItem.organization_id == run.organization_id,
                        PlatformMigrationItem.provider == "APPFOLIO",
                        PlatformMigrationItem.resource == "PROPERTIES",
                        PlatformMigrationItem.source_id == str(property_source),
                    )
                    .first()
                )
                if property_mapping is not None:
                    if property_mapping.target_entity != "PROPERTY":
                        errors.append("Mapped source Property ID has an inconsistent target type.")
                    else:
                        target_property = (
                            db.query(Property)
                            .filter(
                                Property.id == property_mapping.target_id,
                                Property.organization_id == run.organization_id,
                                Property.is_active.is_(True),
                                Property.deleted_at.is_(None),
                            )
                            .first()
                        )
                        if target_property is None:
                            errors.append("Mapped source Property ID does not resolve to an active same-organization Property.")

            unit_mapping = None
            target_unit = None
            if unit_source:
                unit_mapping = (
                    db.query(PlatformMigrationItem)
                    .filter(
                        PlatformMigrationItem.run_id == run.id,
                        PlatformMigrationItem.organization_id == run.organization_id,
                        PlatformMigrationItem.provider == "APPFOLIO",
                        PlatformMigrationItem.resource == "UNITS",
                        PlatformMigrationItem.source_id == str(unit_source),
                    )
                    .first()
                )
                if unit_mapping is not None:
                    if unit_mapping.target_entity != "UNIT":
                        errors.append("Mapped source Unit ID has an inconsistent target type.")
                    else:
                        target_unit = (
                            db.query(Unit)
                            .join(Property, Property.id == Unit.property_id)
                            .filter(
                                Unit.id == unit_mapping.target_id,
                                Unit.is_active.is_(True),
                                Unit.deleted_at.is_(None),
                                Property.organization_id == run.organization_id,
                                Property.is_active.is_(True),
                                Property.deleted_at.is_(None),
                            )
                            .first()
                        )
                        if target_unit is None:
                            errors.append("Mapped source Unit ID does not resolve to an active same-organization Unit.")

            if target_unit is not None and target_property is not None and target_unit.property_id != target_property.id:
                errors.append("Mapped Lease/occupancy Unit and Property source IDs resolve to different target Properties.")

            if errors:
                disposition = "INVALID"
                invalid += 1
            else:
                valid += 1
                unresolved = missing_identity
                if tenant_source and tenant_mapping is None:
                    unresolved = True
                    warnings.append("Source Tenant ID has no durable TENANTS mapping yet; relationship remains review-only.")
                if property_source and property_mapping is None:
                    unresolved = True
                    warnings.append("Source Property ID has no durable PROPERTIES mapping yet; relationship remains review-only.")
                if unit_source and unit_mapping is None:
                    unresolved = True
                    warnings.append("Source Unit ID has no durable UNITS mapping yet; relationship remains review-only.")

                if not unresolved and target_tenant and target_property and target_unit:
                    warnings.append(
                        "Tenant/Unit/Property source relationships resolve through durable mappings for review, but no Lease or occupancy record is created."
                    )
                warnings.append(
                    "Lease/move dates, status, rent and deposit are preserved as source evidence only; they do not establish occupancy, rent liability, security-deposit receipt or accounting balances."
                )
                disposition = "REVIEW"

        elif resource == "OWNERS":
            normalized_data = {
                field: _mapped_value(source_row, mapping, field)
                for field in OWNER_ALIASES
                if field in mapping
            }
            if ambiguous:
                errors.append(
                    "Ambiguous automatic mapping requires explicit mapping for: "
                    + ", ".join(sorted(ambiguous))
                )
            if missing_required:
                errors.append(
                    "Missing required source columns: " + ", ".join(missing_required)
                )

            owner_name = normalized_data.get("name")
            if owner_name is None or not str(owner_name).strip():
                errors.append("name is required.")

            source_id_value = normalized_data.get("source_id")
            source_id = str(source_id_value).strip() if source_id_value is not None else None
            if source_id:
                if source_id in seen_source_ids:
                    errors.append("Duplicate AppFolio Owner ID in this staged upload.")
                    duplicates += 1
                else:
                    seen_source_ids.add(source_id)

            email_value = normalized_data.get("email")
            email = str(email_value).strip() if email_value is not None else None
            property_ids_value = normalized_data.get("properties_owned_ids")
            property_ids = (
                str(property_ids_value).strip()
                if property_ids_value is not None
                else None
            )

            if errors:
                disposition = "INVALID"
                invalid += 1
            else:
                valid += 1
                mapped_owner = None
                if source_id:
                    mapped_owner = (
                        db.query(PlatformMigrationItem)
                        .filter(
                            PlatformMigrationItem.run_id == run.id,
                            PlatformMigrationItem.organization_id == run.organization_id,
                            PlatformMigrationItem.provider == "APPFOLIO",
                            PlatformMigrationItem.resource == "OWNERS",
                            PlatformMigrationItem.source_id == source_id,
                        )
                        .first()
                    )

                if mapped_owner is not None:
                    disposition = "ALREADY_MAPPED"
                    warnings.append(
                        f"Source Owner ID is already mapped to {mapped_owner.target_entity} #{mapped_owner.target_id}."
                    )
                else:
                    review_required = False
                    if not source_id:
                        review_required = True
                        warnings.append(
                            "Owner ID was not supplied; durable Owner source identity must be resolved before dry run or commit."
                        )
                    if not email:
                        review_required = True
                        warnings.append(
                            "Owner email was not supplied; target owner identity/contact mapping requires explicit review."
                        )
                    if not property_ids:
                        review_required = True
                        warnings.append(
                            "Properties Owned IDs were not supplied; owner-to-property relationships require explicit review."
                        )

                    candidate = None
                    if email:
                        candidate = (
                            db.query(User)
                            .filter(
                                User.organization_id == run.organization_id,
                                User.role == UserRole.OWNER,
                                User.is_active.is_(True),
                                User.deleted_at.is_(None),
                                func.lower(User.email) == email.lower(),
                            )
                            .first()
                        )
                    if candidate is not None:
                        possible_matches += 1
                        warnings.append(
                            f"Possible existing target owner match by exact email: local owner #{candidate.id}."
                        )
                        disposition = "REVIEW" if review_required else "POSSIBLE_MATCH"
                    elif review_required:
                        disposition = "REVIEW"
                    else:
                        disposition = "NEW"
        elif resource == "VENDORS":
            normalized_data = {
                field: _mapped_value(source_row, mapping, field)
                for field in VENDOR_ALIASES
                if field in mapping
            }
            if ambiguous:
                errors.append(
                    "Ambiguous automatic mapping requires explicit mapping for: "
                    + ", ".join(sorted(ambiguous))
                )
            if missing_required:
                errors.append(
                    "Missing required source columns: " + ", ".join(missing_required)
                )

            company_value = normalized_data.get("company_name")
            company_name = (
                str(company_value).strip() if company_value is not None else None
            )
            if not company_name:
                errors.append("company_name is required.")

            source_id_value = normalized_data.get("source_id")
            source_id = str(source_id_value).strip() if source_id_value is not None else None
            if source_id:
                if source_id in seen_source_ids:
                    errors.append("Duplicate AppFolio Vendor ID in this staged upload.")
                    duplicates += 1
                else:
                    seen_source_ids.add(source_id)

            email_value = normalized_data.get("email")
            email = str(email_value).strip() if email_value is not None else None

            if errors:
                disposition = "INVALID"
                invalid += 1
            else:
                valid += 1
                mapped_vendor = None
                if source_id:
                    mapped_vendor = (
                        db.query(PlatformMigrationItem)
                        .filter(
                            PlatformMigrationItem.run_id == run.id,
                            PlatformMigrationItem.organization_id == run.organization_id,
                            PlatformMigrationItem.provider == "APPFOLIO",
                            PlatformMigrationItem.resource == "VENDORS",
                            PlatformMigrationItem.source_id == source_id,
                        )
                        .first()
                    )

                if mapped_vendor is not None:
                    disposition = "ALREADY_MAPPED"
                    warnings.append(
                        f"Source Vendor ID is already mapped to {mapped_vendor.target_entity} #{mapped_vendor.target_id}."
                    )
                else:
                    review_required = False
                    if not source_id:
                        review_required = True
                        warnings.append(
                            "Vendor ID was not supplied; durable Vendor source identity must be resolved before dry run or commit."
                        )

                    candidate = None
                    if email:
                        candidate = (
                            db.query(Vendor)
                            .filter(
                                Vendor.organization_id == run.organization_id,
                                Vendor.is_active.is_(True),
                                Vendor.deleted_at.is_(None),
                                func.lower(Vendor.business_email) == email.lower(),
                            )
                            .first()
                        )
                    if candidate is None and company_name:
                        candidate = (
                            db.query(Vendor)
                            .filter(
                                Vendor.organization_id == run.organization_id,
                                Vendor.is_active.is_(True),
                                Vendor.deleted_at.is_(None),
                                func.lower(Vendor.company_name) == company_name.lower(),
                            )
                            .first()
                        )
                    if candidate is not None:
                        possible_matches += 1
                        warnings.append(
                            f"Possible existing target vendor match: local vendor #{candidate.id}; explicit review is required before any future commit."
                        )
                        disposition = "REVIEW" if review_required else "POSSIBLE_MATCH"
                    elif review_required:
                        disposition = "REVIEW"
                    else:
                        disposition = "NEW"
        elif resource == "GL_ACCOUNTS":
            normalized_data = {
                field: _mapped_value(source_row, mapping, field)
                for field in GL_ACCOUNT_ALIASES
                if field in mapping
            }
            if ambiguous:
                errors.append(
                    "Ambiguous automatic mapping requires explicit mapping for: "
                    + ", ".join(sorted(ambiguous))
                )
            if missing_required:
                errors.append(
                    "Missing required source columns: " + ", ".join(missing_required)
                )

            account_number_value = normalized_data.get("account_number")
            account_name_value = normalized_data.get("account_name")
            account_type_value = normalized_data.get("account_type")
            account_number = (
                str(account_number_value).strip()
                if account_number_value is not None
                else None
            )
            account_name = (
                str(account_name_value).strip()
                if account_name_value is not None
                else None
            )
            account_type = (
                str(account_type_value).strip()
                if account_type_value is not None
                else None
            )
            if not account_number:
                errors.append("account_number is required.")
            if not account_name:
                errors.append("account_name is required.")
            if not account_type:
                errors.append("account_type is required.")

            source_id_value = normalized_data.get("source_id")
            source_id = (
                str(source_id_value).strip()
                if source_id_value is not None
                else None
            )
            if source_id:
                if source_id in seen_source_ids:
                    errors.append("Duplicate AppFolio GL Account ID in this staged upload.")
                    duplicates += 1
                else:
                    seen_source_ids.add(source_id)

            if account_number and account_name and account_type:
                duplicate_key = (account_number, account_name, account_type)
                if duplicate_key in seen_gl_account_keys:
                    errors.append(
                        "Duplicate AppFolio GL account number/name/type row in this staged upload."
                    )
                    duplicates += 1
                else:
                    seen_gl_account_keys.add(duplicate_key)

            if errors:
                disposition = "INVALID"
                invalid += 1
            else:
                valid += 1
                mapped_item = None
                if source_id:
                    mapped_item = (
                        db.query(PlatformMigrationItem)
                        .filter(
                            PlatformMigrationItem.run_id == run.id,
                            PlatformMigrationItem.organization_id == run.organization_id,
                            PlatformMigrationItem.provider == "APPFOLIO",
                            PlatformMigrationItem.resource == "GL_ACCOUNTS",
                            PlatformMigrationItem.source_id == source_id,
                        )
                        .first()
                    )
                if mapped_item is not None:
                    disposition = "ALREADY_MAPPED"
                    warnings.append(
                        f"Source GL Account ID is already mapped to {mapped_item.target_entity} #{mapped_item.target_id}."
                    )
                else:
                    disposition = "REVIEW"
                    if not source_id:
                        warnings.append(
                            "A stable GL Account ID was not supplied; the source account number/code is preserved but is not promoted to durable source identity automatically."
                        )
                    warnings.append(
                        "AppFolio account Type/FundAccount/corporate/parent/offset/property fields are preserved as source evidence only; no target GL classification or key-account semantics are inferred."
                    )
                    warnings.append(
                        "This staging batch creates or updates no GLAccount, key-account configuration, journal entry, GL transaction or accounting balance."
                    )
        elif resource == "GENERAL_LEDGER":
            normalized_data = {
                field: _mapped_value(source_row, mapping, field)
                for field in GENERAL_LEDGER_ALIASES
                if field in mapping
            }
            if ambiguous:
                errors.append(
                    "Ambiguous automatic mapping requires explicit mapping for: "
                    + ", ".join(sorted(ambiguous))
                )
            if missing_required:
                errors.append(
                    "Missing required source columns: " + ", ".join(missing_required)
                )

            source_line_value = normalized_data.get("source_id")
            source_id = (
                str(source_line_value).strip()
                if source_line_value is not None
                else None
            )
            if source_id:
                if source_id in seen_general_ledger_line_ids:
                    errors.append("Duplicate AppFolio General Ledger LineItemId in this staged upload.")
                    duplicates += 1
                else:
                    seen_general_ledger_line_ids.add(source_id)

            source_gl_value = normalized_data.get("source_gl_account_id")
            source_gl_account_id = (
                str(source_gl_value).strip()
                if source_gl_value is not None
                else None
            )
            posted_date_value = normalized_data.get("posted_date")
            posted_date = (
                str(posted_date_value).strip()
                if posted_date_value is not None
                else None
            )
            if not source_gl_account_id:
                errors.append("source_gl_account_id is required.")
            if not posted_date:
                errors.append("posted_date is required.")

            for money_field in ("debit", "credit"):
                value = normalized_data.get(money_field)
                if value in (None, ""):
                    errors.append(f"{money_field} is required.")
                    continue
                try:
                    Decimal(str(value).replace(",", "").strip())
                except (InvalidOperation, ValueError):
                    errors.append(f"{money_field} must be a numeric source amount.")

            mapped_gl = None
            if source_gl_account_id:
                mapped_gl = (
                    db.query(PlatformMigrationItem)
                    .filter(
                        PlatformMigrationItem.run_id == run.id,
                        PlatformMigrationItem.organization_id == run.organization_id,
                        PlatformMigrationItem.provider == "APPFOLIO",
                        PlatformMigrationItem.resource == "GL_ACCOUNTS",
                        PlatformMigrationItem.source_id == source_gl_account_id,
                    )
                    .first()
                )
                if mapped_gl is None:
                    warnings.append(
                        "General Ledger GL Account relationship is unresolved; commit remains blocked until the source GL Account is durably mapped."
                    )
                elif mapped_gl.target_entity != "GL_ACCOUNT":
                    errors.append(
                        "Mapped source GL Account ID does not resolve to a GL_ACCOUNT target."
                    )
                else:
                    target_gl = (
                        db.query(GLAccount)
                        .filter(
                            GLAccount.id == mapped_gl.target_id,
                            GLAccount.organization_id == run.organization_id,
                            GLAccount.is_active.is_(True),
                            GLAccount.deleted_at.is_(None),
                        )
                        .first()
                    )
                    if target_gl is None:
                        errors.append(
                            "Mapped source GL Account ID no longer resolves to an active same-organization GL Account."
                        )

            for field_name, resource_name, target_entity in (
                ("source_property_id", "PROPERTIES", "PROPERTY"),
                ("source_unit_id", "UNITS", "UNIT"),
            ):
                raw_value = normalized_data.get(field_name)
                source_ref = str(raw_value).strip() if raw_value not in (None, "") else None
                if not source_ref:
                    continue
                linked = (
                    db.query(PlatformMigrationItem)
                    .filter(
                        PlatformMigrationItem.run_id == run.id,
                        PlatformMigrationItem.organization_id == run.organization_id,
                        PlatformMigrationItem.provider == "APPFOLIO",
                        PlatformMigrationItem.resource == resource_name,
                        PlatformMigrationItem.source_id == source_ref,
                    )
                    .first()
                )
                if linked is None:
                    warnings.append(
                        f"General Ledger {resource_name[:-1].title()} relationship {source_ref} is not yet durably mapped."
                    )
                elif linked.target_entity != target_entity:
                    errors.append(
                        f"Mapped source {resource_name[:-1].title()} ID resolves to the wrong target type."
                    )

            if errors:
                disposition = "INVALID"
                invalid += 1
            else:
                valid += 1
                disposition = "REVIEW"
                if not source_id:
                    warnings.append(
                        "LineItemId was not supplied; no durable General Ledger row identity is synthesized from dates, descriptions, references, transaction IDs or amounts."
                    )
                if mapped_gl is not None:
                    warnings.append(
                        "The source GL Account relationship is durably mapped; this row remains review-only until a separate accounting-history reconciliation batch is verified."
                    )
                warnings.append(
                    "Debit, credit, date, description, reference, remarks and transaction type are preserved as source evidence only; no accounting posting or balance is created."
                )
                warnings.append(
                    "This staging batch creates or updates no GLTransaction, GLEntry, journal entry, Receipt, Bill, Charge or accounting balance."
                )
        elif resource == "BILLS":
            normalized_data = {
                field: _mapped_value(source_row, mapping, field)
                for field in BILL_ALIASES
                if field in mapping
            }
            if ambiguous:
                errors.append(
                    "Ambiguous automatic mapping requires explicit mapping for: "
                    + ", ".join(sorted(ambiguous))
                )
            if missing_required:
                errors.append(
                    "Missing required source columns: " + ", ".join(missing_required)
                )

            source_bill_value = normalized_data.get("source_id")
            source_id = (
                str(source_bill_value).strip()
                if source_bill_value is not None
                else None
            )
            if source_id:
                if source_id in seen_source_ids:
                    errors.append("Duplicate AppFolio Bill ID in this staged upload.")
                    duplicates += 1
                else:
                    seen_source_ids.add(source_id)

            source_vendor_value = normalized_data.get("source_vendor_id")
            source_vendor_id = (
                str(source_vendor_value).strip()
                if source_vendor_value is not None
                else None
            )
            due_date = normalized_data.get("due_date")
            total_amount = normalized_data.get("total_amount")
            if not source_vendor_id:
                errors.append("source_vendor_id is required.")
            if due_date in (None, ""):
                errors.append("due_date is required.")
            if total_amount in (None, ""):
                errors.append("total_amount is required.")
            else:
                try:
                    amount = Decimal(str(total_amount).replace(",", "").strip())
                    if amount < 0:
                        errors.append("total_amount must be a nonnegative source amount.")
                except (InvalidOperation, ValueError):
                    errors.append("total_amount must be a numeric source amount.")

            mapped_vendor = None
            if source_vendor_id:
                mapped_vendor = (
                    db.query(PlatformMigrationItem)
                    .filter(
                        PlatformMigrationItem.run_id == run.id,
                        PlatformMigrationItem.organization_id == run.organization_id,
                        PlatformMigrationItem.provider == "APPFOLIO",
                        PlatformMigrationItem.resource == "VENDORS",
                        PlatformMigrationItem.source_id == source_vendor_id,
                    )
                    .first()
                )
                if mapped_vendor is None:
                    warnings.append(
                        "Bill VendorId has no durable VENDORS mapping yet; Vendor relationship remains unresolved."
                    )
                elif mapped_vendor.target_entity != "VENDOR":
                    errors.append(
                        "Mapped source VendorId does not resolve to a VENDOR target."
                    )
                else:
                    target_vendor = (
                        db.query(Vendor)
                        .filter(
                            Vendor.id == mapped_vendor.target_id,
                            Vendor.organization_id == run.organization_id,
                            Vendor.is_active.is_(True),
                            Vendor.deleted_at.is_(None),
                        )
                        .first()
                    )
                    if target_vendor is None:
                        errors.append(
                            "Mapped source VendorId no longer resolves to an active same-organization Vendor."
                        )

            source_property_value = normalized_data.get("source_property_id")
            source_property_id = (
                str(source_property_value).strip()
                if source_property_value not in (None, "")
                else None
            )
            if source_property_id:
                mapped_property = (
                    db.query(PlatformMigrationItem)
                    .filter(
                        PlatformMigrationItem.run_id == run.id,
                        PlatformMigrationItem.organization_id == run.organization_id,
                        PlatformMigrationItem.provider == "APPFOLIO",
                        PlatformMigrationItem.resource == "PROPERTIES",
                        PlatformMigrationItem.source_id == source_property_id,
                    )
                    .first()
                )
                if mapped_property is None:
                    warnings.append(
                        "Bill PropertyId has no durable PROPERTIES mapping yet; Property relationship remains unresolved."
                    )
                elif mapped_property.target_entity != "PROPERTY":
                    errors.append(
                        "Mapped source PropertyId does not resolve to a PROPERTY target."
                    )
                else:
                    target_property = (
                        db.query(Property)
                        .filter(
                            Property.id == mapped_property.target_id,
                            Property.organization_id == run.organization_id,
                            Property.is_active.is_(True),
                            Property.deleted_at.is_(None),
                        )
                        .first()
                    )
                    if target_property is None:
                        errors.append(
                            "Mapped source PropertyId no longer resolves to an active same-organization Property."
                        )

            if errors:
                disposition = "INVALID"
                invalid += 1
            else:
                valid += 1
                disposition = "REVIEW"
                if not source_id:
                    warnings.append(
                        "Bill ID was not supplied; no durable Bill source identity is synthesized from reference, vendor, dates or amount."
                    )
                if mapped_vendor is not None:
                    warnings.append(
                        "The Bill Vendor relationship is durably mapped; the staged row remains review-only."
                    )
                if source_property_id and not any(
                    "Property relationship remains unresolved" in warning
                    for warning in warnings
                ):
                    warnings.append(
                        "The Bill Property relationship is durably mapped; the staged row remains review-only."
                    )
                if normalized_data.get("source_work_order_id") not in (None, ""):
                    warnings.append(
                        "WorkOrderId is preserved as source evidence only; no Work Order relationship is inferred without a durable migration mapping."
                    )
                warnings.append(
                    "ApprovalStatus, ManagementCompanyAsPayee, dates, reference, remarks and amount remain source evidence only; paid/unpaid, approval, check/payment and GL posting state are not inferred."
                )
                warnings.append(
                    "This staging batch creates or updates no Bill, BillLine, Check, Vendor, WorkOrder, GLTransaction, GLEntry, Charge or accounting balance."
                )
        elif resource == "WORK_ORDERS":
            normalized_data = {
                field: _mapped_value(source_row, mapping, field)
                for field in WORK_ORDER_ALIASES
                if field in mapping
            }
            if ambiguous:
                errors.append(
                    "Ambiguous automatic mapping requires explicit mapping for: "
                    + ", ".join(sorted(ambiguous))
                )
            if missing_required:
                errors.append(
                    "Missing required source columns: " + ", ".join(missing_required)
                )

            source_work_order_value = normalized_data.get("source_id")
            source_id = (
                str(source_work_order_value).strip()
                if source_work_order_value is not None
                else None
            )
            if source_id:
                if source_id in seen_source_ids:
                    errors.append("Duplicate AppFolio Work Order ID in this staged upload.")
                    duplicates += 1
                else:
                    seen_source_ids.add(source_id)

            source_property_value = normalized_data.get("source_property_id")
            source_property_id = (
                str(source_property_value).strip()
                if source_property_value not in (None, "")
                else None
            )
            source_unit_value = normalized_data.get("source_unit_id")
            source_unit_id = (
                str(source_unit_value).strip()
                if source_unit_value not in (None, "")
                else None
            )
            source_vendor_value = normalized_data.get("source_vendor_id")
            source_vendor_id = (
                str(source_vendor_value).strip()
                if source_vendor_value not in (None, "")
                else None
            )

            for field_name in ("status", "job_description"):
                value = normalized_data.get(field_name)
                if value is None or not str(value).strip():
                    errors.append(f"{field_name} is required.")

            mapped_property = None
            target_property = None
            if source_property_id:
                mapped_property = (
                    db.query(PlatformMigrationItem)
                    .filter(
                        PlatformMigrationItem.run_id == run.id,
                        PlatformMigrationItem.organization_id == run.organization_id,
                        PlatformMigrationItem.provider == "APPFOLIO",
                        PlatformMigrationItem.resource == "PROPERTIES",
                        PlatformMigrationItem.source_id == source_property_id,
                    )
                    .first()
                )
                if mapped_property is None:
                    warnings.append(
                        "Work Order Property relationship is unresolved; PropertyId is preserved as source evidence."
                    )
                elif mapped_property.target_entity != "PROPERTY":
                    errors.append(
                        "Mapped Work Order PropertyId resolves to the wrong target type."
                    )
                else:
                    target_property = (
                        db.query(Property)
                        .filter(
                            Property.id == mapped_property.target_id,
                            Property.organization_id == run.organization_id,
                            Property.is_active.is_(True),
                            Property.deleted_at.is_(None),
                        )
                        .first()
                    )
                    if target_property is None:
                        errors.append(
                            "Mapped Work Order PropertyId no longer resolves to an active same-organization Property."
                        )

            mapped_unit = None
            target_unit = None
            if source_unit_id:
                mapped_unit = (
                    db.query(PlatformMigrationItem)
                    .filter(
                        PlatformMigrationItem.run_id == run.id,
                        PlatformMigrationItem.organization_id == run.organization_id,
                        PlatformMigrationItem.provider == "APPFOLIO",
                        PlatformMigrationItem.resource == "UNITS",
                        PlatformMigrationItem.source_id == source_unit_id,
                    )
                    .first()
                )
                if mapped_unit is None:
                    warnings.append(
                        "Work Order Unit relationship is unresolved; UnitId is preserved as source evidence."
                    )
                elif mapped_unit.target_entity != "UNIT":
                    errors.append(
                        "Mapped Work Order UnitId resolves to the wrong target type."
                    )
                else:
                    target_unit = (
                        db.query(Unit)
                        .join(Property, Property.id == Unit.property_id)
                        .filter(
                            Unit.id == mapped_unit.target_id,
                            Unit.is_active.is_(True),
                            Unit.deleted_at.is_(None),
                            Property.organization_id == run.organization_id,
                            Property.is_active.is_(True),
                            Property.deleted_at.is_(None),
                        )
                        .first()
                    )
                    if target_unit is None:
                        errors.append(
                            "Mapped Work Order UnitId no longer resolves to an active same-organization Unit."
                        )

            if (
                target_property is not None
                and target_unit is not None
                and target_unit.property_id != target_property.id
            ):
                errors.append(
                    "Mapped Work Order PropertyId and UnitId resolve to different target Properties."
                )

            mapped_vendor = None
            if source_vendor_id:
                mapped_vendor = (
                    db.query(PlatformMigrationItem)
                    .filter(
                        PlatformMigrationItem.run_id == run.id,
                        PlatformMigrationItem.organization_id == run.organization_id,
                        PlatformMigrationItem.provider == "APPFOLIO",
                        PlatformMigrationItem.resource == "VENDORS",
                        PlatformMigrationItem.source_id == source_vendor_id,
                    )
                    .first()
                )
                if mapped_vendor is None:
                    warnings.append(
                        "Work Order Vendor relationship is unresolved; VendorId is preserved as source evidence."
                    )
                elif mapped_vendor.target_entity != "VENDOR":
                    errors.append(
                        "Mapped Work Order VendorId resolves to the wrong target type."
                    )
                else:
                    target_vendor = (
                        db.query(Vendor)
                        .filter(
                            Vendor.id == mapped_vendor.target_id,
                            Vendor.organization_id == run.organization_id,
                            Vendor.is_active.is_(True),
                            Vendor.deleted_at.is_(None),
                        )
                        .first()
                    )
                    if target_vendor is None:
                        errors.append(
                            "Mapped Work Order VendorId no longer resolves to an active same-organization Vendor."
                        )

            if errors:
                disposition = "INVALID"
                invalid += 1
            else:
                valid += 1
                disposition = "REVIEW"
                if not source_id:
                    warnings.append(
                        "Work Order ID was not supplied; no durable Work Order identity is synthesized from property, unit, vendor, status, dates or description."
                    )
                warnings.append(
                    "Status, priority, schedule, completion/cancellation dates, permission-to-enter, AssignedUsers and VendorTrade remain source evidence only; target workflow state and assignments are not inferred."
                )
                warnings.append(
                    "This staging batch creates or updates no WorkOrder, Bill, Charge, GLTransaction, GLEntry, inventory, purchase order or vendor-payment record."
                )
        elif resource == "CHARGES":
            normalized_data = {
                field: _mapped_value(source_row, mapping, field)
                for field in CHARGE_ALIASES
                if field in mapping
            }
            if ambiguous:
                errors.append(
                    "Ambiguous automatic mapping requires explicit mapping for: "
                    + ", ".join(sorted(ambiguous))
                )
            if missing_required:
                errors.append(
                    "Missing required source columns: " + ", ".join(missing_required)
                )

            source_charge_value = normalized_data.get("source_id")
            source_id = (
                str(source_charge_value).strip()
                if source_charge_value is not None
                else None
            )
            if source_id:
                if source_id in seen_source_ids:
                    errors.append("Duplicate AppFolio Charge ID in this staged upload.")
                    duplicates += 1
                else:
                    seen_source_ids.add(source_id)

            amount_due = normalized_data.get("amount_due")
            if amount_due in (None, ""):
                errors.append("amount_due is required.")
            else:
                try:
                    Decimal(str(amount_due).replace(",", "").strip())
                except (InvalidOperation, ValueError):
                    errors.append("amount_due must be a numeric source amount.")

            charged_on_value = normalized_data.get("charged_on")
            charged_on = (
                str(charged_on_value).strip()
                if charged_on_value is not None
                else None
            )
            if not charged_on:
                errors.append("charged_on is required.")

            source_gl_value = normalized_data.get("source_gl_account_id")
            source_gl_account_id = (
                str(source_gl_value).strip()
                if source_gl_value is not None
                else None
            )
            if not source_gl_account_id:
                errors.append("source_gl_account_id is required.")

            source_occupancy_value = normalized_data.get("source_occupancy_id")
            source_occupancy_id = (
                str(source_occupancy_value).strip()
                if source_occupancy_value is not None
                else None
            )
            if not source_occupancy_id:
                errors.append("source_occupancy_id is required.")

            mapped_gl = None
            if source_gl_account_id:
                mapped_gl = (
                    db.query(PlatformMigrationItem)
                    .filter(
                        PlatformMigrationItem.run_id == run.id,
                        PlatformMigrationItem.organization_id == run.organization_id,
                        PlatformMigrationItem.provider == "APPFOLIO",
                        PlatformMigrationItem.resource == "GL_ACCOUNTS",
                        PlatformMigrationItem.source_id == source_gl_account_id,
                    )
                    .first()
                )
                if mapped_gl is None:
                    warnings.append(
                        "Charge GL Account relationship is unresolved; the source GlAccountId has no durable GL_ACCOUNTS mapping yet."
                    )
                elif mapped_gl.target_entity != "GL_ACCOUNT":
                    errors.append(
                        "Mapped source Charge GlAccountId does not resolve to a GL_ACCOUNT target."
                    )
                else:
                    target_gl = (
                        db.query(GLAccount)
                        .filter(
                            GLAccount.id == mapped_gl.target_id,
                            GLAccount.organization_id == run.organization_id,
                            GLAccount.is_active.is_(True),
                            GLAccount.deleted_at.is_(None),
                        )
                        .first()
                    )
                    if target_gl is None:
                        errors.append(
                            "Mapped source Charge GlAccountId no longer resolves to an active same-organization GL Account."
                        )

            if errors:
                disposition = "INVALID"
                invalid += 1
            else:
                valid += 1
                disposition = "REVIEW"
                if not source_id:
                    warnings.append(
                        "Charge ID was not supplied; no durable Charge source identity is synthesized from occupancy, GL account, date, description or amount."
                    )
                if mapped_gl is not None:
                    warnings.append(
                        "The Charge GL Account relationship is durably mapped; the staged row remains review-only."
                    )
                warnings.append(
                    "OccupancyId is preserved as source evidence only because no durable occupancy source-to-target mapping is verified in the migration architecture."
                )
                warnings.append(
                    "AmountDue is preserved as the source current outstanding amount; it is not treated as original target Charge amount, amount_paid, is_paid, rent liability or reconciled balance."
                )
                warnings.append(
                    "ChargedOn and Description are source evidence only; no tenant, lease, property/unit, payer liability, payment application or GL posting is inferred."
                )
                warnings.append(
                    "This staging batch creates or updates no Charge, RentInvoice, Receipt, Payment, GLTransaction, GLEntry or customer balance."
                )
        else:
            normalized_data = _normalized_source_row(source_row)
            errors.append(
                "Report type could not be detected safely; choose PROPERTIES, UNITS, TENANTS, LEASE_OCCUPANCY, OWNERS, VENDORS, GL_ACCOUNTS, GENERAL_LEDGER, BILLS, CHARGES or WORK_ORDERS and supply explicit column mapping."
            )
            disposition = "INVALID"
            invalid += 1

        if warnings:
            warning_rows += 1

        db.add(
            PlatformMigrationStagedRow(
                upload_id=upload.id,
                run_id=run.id,
                organization_id=run.organization_id,
                provider="APPFOLIO",
                resource=resource,
                row_number=row_number,
                source_id=source_id,
                disposition=disposition,
                row_fingerprint=_row_fingerprint(source_row),
                normalized_data=normalized_data,
                warnings=warnings,
                errors=errors,
            )
        )

    summary = {
        "total": len(parsed.rows),
        "valid": valid,
        "warnings": warning_rows,
        "invalid": invalid,
        "duplicates": duplicates,
        "possible_existing_matches": possible_matches,
        "missing_required_columns": missing_required,
        "ambiguous_mapping_fields": sorted(ambiguous),
    }
    if resource == "UNITS":
        unit_rows = (
            db.query(PlatformMigrationStagedRow)
            .filter(
                PlatformMigrationStagedRow.upload_id == upload.id,
                PlatformMigrationStagedRow.resource == "UNITS",
            )
            .all()
        )
        summary["unresolved_property_links"] = sum(
            1
            for row in unit_rows
            if any("Property" in warning for warning in (row.warnings or []))
        )
        summary["missing_unit_source_ids"] = sum(
            1
            for row in unit_rows
            if any("Unit ID was not supplied" in warning for warning in (row.warnings or []))
        )
    if resource == "OWNERS":
        owner_rows = (
            db.query(PlatformMigrationStagedRow)
            .filter(
                PlatformMigrationStagedRow.upload_id == upload.id,
                PlatformMigrationStagedRow.resource == "OWNERS",
            )
            .all()
        )
        summary["missing_owner_source_ids"] = sum(
            1
            for row in owner_rows
            if any("Owner ID was not supplied" in warning for warning in (row.warnings or []))
        )
        summary["missing_owner_emails"] = sum(
            1
            for row in owner_rows
            if any("Owner email was not supplied" in warning for warning in (row.warnings or []))
        )
        summary["missing_owner_property_ids"] = sum(
            1
            for row in owner_rows
            if any("Properties Owned IDs were not supplied" in warning for warning in (row.warnings or []))
        )
    if resource == "TENANTS":
        tenant_rows = (
            db.query(PlatformMigrationStagedRow)
            .filter(
                PlatformMigrationStagedRow.upload_id == upload.id,
                PlatformMigrationStagedRow.resource == "TENANTS",
            )
            .all()
        )
        summary["missing_tenant_source_ids"] = sum(
            1
            for row in tenant_rows
            if any("Tenant ID was not supplied" in warning for warning in (row.warnings or []))
        )
        summary["missing_tenant_unit_ids"] = sum(
            1
            for row in tenant_rows
            if any("Unit ID was not supplied" in warning for warning in (row.warnings or []))
        )
        summary["missing_tenant_property_ids"] = sum(
            1
            for row in tenant_rows
            if any("Property ID was not supplied" in warning for warning in (row.warnings or []))
        )
        summary["unresolved_tenant_relationships"] = sum(
            1
            for row in tenant_rows
            if any(
                "relationship" in warning.lower()
                for warning in (row.warnings or [])
            )
        )
    if resource == "LEASE_OCCUPANCY":
        lease_rows = (
            db.query(PlatformMigrationStagedRow)
            .filter(
                PlatformMigrationStagedRow.upload_id == upload.id,
                PlatformMigrationStagedRow.resource == "LEASE_OCCUPANCY",
            )
            .all()
        )
        summary["lease_occupancy_rows"] = len(lease_rows)
        summary["relationships_ready_for_review"] = sum(
            1
            for row in lease_rows
            if row.disposition == "REVIEW"
            and any(
                "resolve through durable mappings for review" in warning
                for warning in (row.warnings or [])
            )
        )
        summary["unresolved_lease_occupancy_relationships"] = sum(
            1
            for row in lease_rows
            if any(
                "relationship remains review-only" in warning
                for warning in (row.warnings or [])
            )
        )
        summary["customer_lease_mutation"] = False
        summary["accounting_mutation"] = False
    if resource == "VENDORS":
        vendor_rows = (
            db.query(PlatformMigrationStagedRow)
            .filter(
                PlatformMigrationStagedRow.upload_id == upload.id,
                PlatformMigrationStagedRow.resource == "VENDORS",
            )
            .all()
        )
        summary["missing_vendor_source_ids"] = sum(
            1
            for row in vendor_rows
            if any("Vendor ID was not supplied" in warning for warning in (row.warnings or []))
        )
    if resource == "GL_ACCOUNTS":
        gl_rows = (
            db.query(PlatformMigrationStagedRow)
            .filter(
                PlatformMigrationStagedRow.upload_id == upload.id,
                PlatformMigrationStagedRow.resource == "GL_ACCOUNTS",
            )
            .all()
        )
        summary["gl_account_rows"] = len(gl_rows)
        summary["missing_gl_account_source_ids"] = sum(
            1
            for row in gl_rows
            if any(
                "stable GL Account ID was not supplied" in warning
                for warning in (row.warnings or [])
            )
        )
        summary["gl_account_target_mutation"] = False
        summary["accounting_history_mutation"] = False
    if resource == "GENERAL_LEDGER":
        ledger_rows = (
            db.query(PlatformMigrationStagedRow)
            .filter(
                PlatformMigrationStagedRow.upload_id == upload.id,
                PlatformMigrationStagedRow.resource == "GENERAL_LEDGER",
            )
            .all()
        )
        summary["general_ledger_rows"] = len(ledger_rows)
        summary["missing_general_ledger_line_ids"] = sum(
            1
            for row in ledger_rows
            if any(
                "LineItemId was not supplied" in warning
                for warning in (row.warnings or [])
            )
        )
        summary["unresolved_general_ledger_gl_accounts"] = sum(
            1
            for row in ledger_rows
            if any(
                "GL Account relationship is unresolved" in warning
                for warning in (row.warnings or [])
            )
        )
        summary["accounting_history_mutation"] = False
        summary["customer_accounting_mutation"] = False
    if resource == "BILLS":
        bill_rows = (
            db.query(PlatformMigrationStagedRow)
            .filter(
                PlatformMigrationStagedRow.upload_id == upload.id,
                PlatformMigrationStagedRow.resource == "BILLS",
            )
            .all()
        )
        summary["bill_rows"] = len(bill_rows)
        summary["missing_bill_source_ids"] = sum(
            1
            for row in bill_rows
            if any("Bill ID was not supplied" in warning for warning in (row.warnings or []))
        )
        summary["unresolved_bill_vendor_relationships"] = sum(
            1
            for row in bill_rows
            if any("Vendor relationship remains unresolved" in warning for warning in (row.warnings or []))
        )
        summary["unresolved_bill_property_relationships"] = sum(
            1
            for row in bill_rows
            if any("Property relationship remains unresolved" in warning for warning in (row.warnings or []))
        )
        summary["bill_mutation"] = False
        summary["accounting_history_mutation"] = False
        summary["payment_state_inferred"] = False
    if resource == "WORK_ORDERS":
        work_order_rows = (
            db.query(PlatformMigrationStagedRow)
            .filter(
                PlatformMigrationStagedRow.upload_id == upload.id,
                PlatformMigrationStagedRow.resource == "WORK_ORDERS",
            )
            .all()
        )
        summary["work_order_rows"] = len(work_order_rows)
        summary["missing_work_order_source_ids"] = sum(
            1
            for row in work_order_rows
            if any("Work Order ID was not supplied" in warning for warning in (row.warnings or []))
        )
        summary["unresolved_work_order_properties"] = sum(
            1
            for row in work_order_rows
            if any("Property relationship is unresolved" in warning for warning in (row.warnings or []))
        )
        summary["unresolved_work_order_units"] = sum(
            1
            for row in work_order_rows
            if any("Unit relationship is unresolved" in warning for warning in (row.warnings or []))
        )
        summary["unresolved_work_order_vendors"] = sum(
            1
            for row in work_order_rows
            if any("Vendor relationship is unresolved" in warning for warning in (row.warnings or []))
        )
        summary["work_order_mutation"] = False
        summary["bill_mutation"] = False
        summary["charge_mutation"] = False
        summary["accounting_history_mutation"] = False
        summary["workflow_state_inferred"] = False
    if resource == "CHARGES":
        charge_rows = (
            db.query(PlatformMigrationStagedRow)
            .filter(
                PlatformMigrationStagedRow.upload_id == upload.id,
                PlatformMigrationStagedRow.resource == "CHARGES",
            )
            .all()
        )
        summary["charge_rows"] = len(charge_rows)
        summary["missing_charge_source_ids"] = sum(
            1
            for row in charge_rows
            if any("Charge ID was not supplied" in warning for warning in (row.warnings or []))
        )
        summary["unresolved_charge_gl_accounts"] = sum(
            1
            for row in charge_rows
            if any("GL Account relationship is unresolved" in warning for warning in (row.warnings or []))
        )
        summary["unresolved_charge_occupancies"] = sum(
            1
            for row in charge_rows
            if any("OccupancyId is preserved as source evidence only" in warning for warning in (row.warnings or []))
        )
        summary["charge_mutation"] = False
        summary["rent_invoice_mutation"] = False
        summary["receipt_payment_mutation"] = False
        summary["accounting_history_mutation"] = False
        summary["payment_state_inferred"] = False
        summary["tenant_liability_inferred"] = False
    if resource == "UNKNOWN" or missing_required or ambiguous:
        upload.status = "MAPPING_REQUIRED"
    elif invalid:
        upload.status = "STAGED_WITH_ERRORS"
    elif warning_rows:
        upload.status = "REVIEW_REQUIRED"
    else:
        upload.status = "STAGED"
    upload.validation_summary = summary

    # Any new staging/mapping state makes an earlier dry run stale.
    run.last_dry_run_fingerprint = None
    run.last_dry_run_summary = None
    run.status = "STAGED"
    db.flush()
    return StagedUploadResult(upload=upload, replayed=False)
