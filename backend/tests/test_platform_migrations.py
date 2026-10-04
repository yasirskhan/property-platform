from __future__ import annotations

import json
from types import SimpleNamespace

import pytest
from fastapi import HTTPException, Response
from pydantic import ValidationError
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

import init_db  # noqa: F401
from app.core.database import Base
from app.core.security import hash_password
from app.models.audit_log import AuditLog
from app.models.platform_migration import (
    PlatformMigrationCorrectionRule,
    PlatformMigrationItem,
    PlatformMigrationRun,
    PlatformMigrationStagedRow,
    PlatformMigrationUpload,
)
from app.models.platform_user import PlatformUser, PlatformUserRole
from app.models.property import Property, PropertyOwner, Unit
from app.models.lease import Lease
from app.models.charge import Charge
from app.models.bill import Bill
from app.models.gl_transaction import GLTransaction
from app.models.gl_account import GLAccount
from app.models.work_order import WorkOrder
from app.models.user import Organization, User, UserRole
from app.models.vendor import Vendor
from app.routers import platform_migrations as api
from app.schemas.platform_migration import (
    AppFolioMigrationRunCreateIn,
    AppFolioMigrationCorrectionRuleCreateIn,
    AppFolioMigrationCorrectionRuleUpdateIn,
    AppFolioStagedRowCorrectionIn,
    AppFolioPropertyCommitIn,
    AppFolioPropertyDryRunIn,
    AppFolioStagedPropertyCommitIn,
    AppFolioStagedRowResolutionIn,
    AppFolioStagedUnitCommitIn,
    AppFolioStagedUnitResolutionIn,
    AppFolioStagedOwnerResolutionIn,
    AppFolioStagedOwnerCommitIn,
    AppFolioStagedVendorResolutionIn,
    AppFolioStagedTenantResolutionIn,
    AppFolioStagedLeaseOccupancyResolutionIn,
    AppFolioStagedGLAccountResolutionIn,
    AppFolioStagedGeneralLedgerResolutionIn,
    AppFolioStagedBillResolutionIn,
    AppFolioStagedChargeResolutionIn,
    AppFolioStagedWorkOrderResolutionIn,
    AppFolioStagedGLAccountCommitIn,
    AppFolioStagedTenantCommitIn,
    AppFolioStagedVendorCommitIn,
)


def _session():
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(bind=engine)
    return sessionmaker(bind=engine, expire_on_commit=False)(), engine


def _platform_user(db, role: PlatformUserRole) -> PlatformUser:
    row = PlatformUser(
        email=f"{role.value}-{db.query(PlatformUser).count()}@example.com",
        hashed_password=hash_password("migration-test-password"),
        first_name="Platform",
        last_name="Migration",
        role=role,
        is_active=True,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def _org(db, *, name: str = "Migration Target", active: bool = True) -> Organization:
    row = Organization(
        name=name,
        slug=name.lower().replace(" ", "-"),
        is_active=active,
    )
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def _valid_record(**changes):
    record = {
        "Id": "AF-100",
        "Name": "Lake Apartments",
        "Address1": "10 Lake Ave",
        "Address2": "Suite 1",
        "City": "Cleveland",
        "State": "OH",
        "Zip": "44113",
        "PropertyType": "Multi Family",
        "HiddenAt": None,
    }
    record.update(changes)
    return record


def test_appfolio_run_roles_scope_and_no_store_reads():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        support = _platform_user(db, PlatformUserRole.PLATFORM_SUPPORT)
        sales = _platform_user(db, PlatformUserRole.PLATFORM_SALES)
        org = _org(db)

        created = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="account-123",
            ),
            db=db,
            current_user=admin,
        )
        assert created.provider == "APPFOLIO"
        assert created.status == "DRAFT"

        response = Response()
        listed = api.list_runs(
            response=response,
            organization_id=org.id,
            limit=100,
            db=db,
            current_user=support,
        )
        assert [row.id for row in listed] == [created.id]
        assert response.headers["cache-control"] == "no-store"

        response = Response()
        fetched = api.get_run(
            created.id,
            response=response,
            db=db,
            current_user=support,
        )
        assert fetched.id == created.id
        assert response.headers["cache-control"] == "no-store"

        with pytest.raises(HTTPException) as exc:
            api.create_run(
                AppFolioMigrationRunCreateIn(
                    organization_id=org.id,
                    source_account_ref="blocked",
                ),
                db=db,
                current_user=support,
            )
        assert exc.value.status_code == 403

        with pytest.raises(HTTPException) as exc:
            api.list_runs(
                response=Response(),
                organization_id=None,
                limit=100,
                db=db,
                current_user=sales,
            )
        assert exc.value.status_code == 403

        org.is_active = False
        db.commit()
        with pytest.raises(HTTPException) as exc:
            api.get_run(
                created.id,
                response=Response(),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 404
    finally:
        db.close()
        engine.dispose()


def test_appfolio_property_dry_run_is_idempotent_and_target_non_mutating():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db)
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="portfolio-A",
            ),
            db=db,
            current_user=admin,
        )

        payload = AppFolioPropertyDryRunIn(
            records=[
                _valid_record(),
                _valid_record(Id="AF-200", HiddenAt="2026-01-01"),
                _valid_record(Id="AF-300", Name=""),
                _valid_record(Id="AF-100", Name="Duplicate"),
            ]
        )
        before_properties = db.query(Property).count()

        first = api.dry_run_appfolio_properties(
            run.id,
            payload,
            db=db,
            current_user=admin,
        )
        assert first.total == 4
        assert first.importable == 1
        assert first.skipped_hidden == 1
        assert first.invalid == 2
        assert first.replayed is False
        assert first.rows[0].mapped["property_type"] == "multi_family"
        assert db.query(Property).count() == before_properties

        stored = db.get(PlatformMigrationRun, run.id)
        assert stored.status == "DRY_RUN_READY"
        assert stored.last_dry_run_fingerprint == first.fingerprint
        assert stored.last_dry_run_summary["resource"] == "PROPERTIES"

        audit_before = (
            db.query(AuditLog)
            .filter(
                AuditLog.entity_type == "platform_migration_run",
                AuditLog.entity_id == run.id,
            )
            .count()
        )

        second = api.dry_run_appfolio_properties(
            run.id,
            payload,
            db=db,
            current_user=admin,
        )
        assert second.replayed is True
        assert second.fingerprint == first.fingerprint
        assert db.query(Property).count() == before_properties

        audit_after = (
            db.query(AuditLog)
            .filter(
                AuditLog.entity_type == "platform_migration_run",
                AuditLog.entity_id == run.id,
            )
            .count()
        )
        assert audit_after == audit_before
    finally:
        db.close()
        engine.dispose()


def test_appfolio_dry_run_warns_on_existing_target_and_unknown_property_type():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_DEV)
        org = _org(db)
        db.add(
            Property(
                organization_id=org.id,
                name="Lake Apartments",
                property_type="multi_family",
                address_line1="10 Lake Ave",
                address_line2="Suite 1",
                city="Cleveland",
                state="OH",
                zip_code="44113",
                country="USA",
                is_active=True,
            )
        )
        db.commit()

        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="portfolio-B",
            ),
            db=db,
            current_user=admin,
        )
        result = api.dry_run_appfolio_properties(
            run.id,
            AppFolioPropertyDryRunIn(
                include_hidden=True,
                records=[_valid_record(PropertyType="Mystery Type", HiddenAt="yes")],
            ),
            db=db,
            current_user=admin,
        )
        assert result.importable == 1
        assert result.skipped_hidden == 0
        assert result.warning_count == 2
        assert result.rows[0].mapped["property_type"] == "other"
        assert any("Unrecognized AppFolio PropertyType" in x for x in result.rows[0].warnings)
        assert any("Possible existing target property match" in x for x in result.rows[0].warnings)
        assert db.query(Property).count() == 1
    finally:
        db.close()
        engine.dispose()


def test_appfolio_schemas_reject_credentials_and_raw_transport_fields():
    with pytest.raises(ValidationError):
        AppFolioMigrationRunCreateIn(
            organization_id=1,
            source_account_ref="portfolio",
            api_key="secret",
        )

    with pytest.raises(ValidationError):
        AppFolioPropertyDryRunIn(
            records=[_valid_record()],
            access_token="secret",
        )



def test_appfolio_property_commit_is_atomic_mapped_and_replay_safe():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db)
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="portfolio-commit",
            ),
            db=db,
            current_user=admin,
        )
        records = [
            _valid_record(),
            _valid_record(
                Id="AF-200",
                Name="River Apartments",
                Address1="20 River Rd",
                Address2=None,
            ),
            _valid_record(Id="AF-HIDDEN", HiddenAt="2026-09-01"),
        ]
        dry = api.dry_run_appfolio_properties(
            run.id,
            AppFolioPropertyDryRunIn(records=records),
            db=db,
            current_user=admin,
        )
        assert dry.importable == 2
        assert dry.skipped_hidden == 1
        before_commit_audit = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_properties_committed",
        ).count()

        committed = api.commit_appfolio_properties(
            run.id,
            AppFolioPropertyCommitIn(
                records=records,
                fingerprint=dry.fingerprint,
            ),
            db=db,
            current_user=admin,
        )
        assert committed.replayed is False
        assert committed.committed == 2
        assert committed.skipped_hidden == 1
        assert db.query(Property).filter(Property.organization_id == org.id).count() == 2
        assert db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "PROPERTIES",
        ).count() == 2
        assert db.get(PlatformMigrationRun, run.id).status == "PROPERTIES_COMMITTED"
        assert all(row.replayed is False for row in committed.rows)

        after_first_audit = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_properties_committed",
        ).count()
        assert after_first_audit == before_commit_audit + 1

        replay = api.commit_appfolio_properties(
            run.id,
            AppFolioPropertyCommitIn(
                records=records,
                fingerprint=dry.fingerprint,
            ),
            db=db,
            current_user=admin,
        )
        assert replay.replayed is True
        assert replay.committed == 0
        assert db.query(Property).filter(Property.organization_id == org.id).count() == 2
        assert all(row.replayed is True for row in replay.rows)
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_properties_committed",
        ).count() == after_first_audit
    finally:
        db.close()
        engine.dispose()


def test_appfolio_property_commit_blocks_changed_invalid_and_existing_targets():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db)
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="portfolio-blocks",
            ),
            db=db,
            current_user=admin,
        )
        valid = [_valid_record()]
        dry = api.dry_run_appfolio_properties(
            run.id,
            AppFolioPropertyDryRunIn(records=valid),
            db=db,
            current_user=admin,
        )

        changed = [_valid_record(Name="Changed after preview")]
        with pytest.raises(HTTPException) as exc:
            api.commit_appfolio_properties(
                run.id,
                AppFolioPropertyCommitIn(
                    records=changed,
                    fingerprint=dry.fingerprint,
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert db.query(Property).count() == 0

        invalid = [_valid_record(Name="")]
        invalid_dry = api.dry_run_appfolio_properties(
            run.id,
            AppFolioPropertyDryRunIn(records=invalid),
            db=db,
            current_user=admin,
        )
        with pytest.raises(HTTPException) as exc:
            api.commit_appfolio_properties(
                run.id,
                AppFolioPropertyCommitIn(
                    records=invalid,
                    fingerprint=invalid_dry.fingerprint,
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert db.query(Property).count() == 0

        existing = Property(
            organization_id=org.id,
            name="Lake Apartments",
            property_type="multi_family",
            address_line1="10 Lake Ave",
            address_line2="Suite 1",
            city="Cleveland",
            state="OH",
            zip_code="44113",
            country="USA",
            is_active=True,
        )
        db.add(existing)
        db.commit()
        conflict_dry = api.dry_run_appfolio_properties(
            run.id,
            AppFolioPropertyDryRunIn(records=valid),
            db=db,
            current_user=admin,
        )
        assert conflict_dry.warning_count == 1
        with pytest.raises(HTTPException) as exc:
            api.commit_appfolio_properties(
                run.id,
                AppFolioPropertyCommitIn(
                    records=valid,
                    fingerprint=conflict_dry.fingerprint,
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert db.query(Property).count() == 1
        assert db.query(PlatformMigrationItem).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_commit_schema_rejects_credentials_and_router_is_exposed():
    with pytest.raises(ValidationError):
        AppFolioPropertyCommitIn(
            records=[_valid_record()],
            fingerprint="0" * 64,
            access_token="secret",
        )

    from app.main import app

    paths = set(app.openapi()["paths"])
    assert "/api/platform/migrations/appfolio/runs" in paths
    assert "/api/platform/migrations/appfolio/runs/{run_id}/properties/dry-run" in paths
    assert "/api/platform/migrations/appfolio/runs/{run_id}/properties/commit" in paths



def test_appfolio_mapping_visibility_is_read_only_scoped_and_no_store():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        support = _platform_user(db, PlatformUserRole.PLATFORM_SUPPORT)
        sales = _platform_user(db, PlatformUserRole.PLATFORM_SALES)
        org = _org(db)
        other_org = _org(db, name="Other Migration Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="portfolio-visibility",
            ),
            db=db,
            current_user=admin,
        )
        records = [_valid_record()]
        dry = api.dry_run_appfolio_properties(
            run.id,
            AppFolioPropertyDryRunIn(records=records),
            db=db,
            current_user=admin,
        )
        committed = api.commit_appfolio_properties(
            run.id,
            AppFolioPropertyCommitIn(
                records=records,
                fingerprint=dry.fingerprint,
            ),
            db=db,
            current_user=admin,
        )
        assert committed.committed == 1
        target_id = committed.rows[0].target_property_id

        response = Response()
        rows = api.list_appfolio_migration_items(
            run.id,
            response=response,
            resource=None,
            limit=200,
            db=db,
            current_user=support,
        )
        assert response.headers["cache-control"] == "no-store"
        assert len(rows) == 1
        assert rows[0].source_id == "AF-100"
        assert rows[0].target_entity == "PROPERTY"
        assert rows[0].target_id == target_id
        assert rows[0].target_exists is True
        assert rows[0].target_label == "Lake Apartments"
        assert rows[0].source_fingerprint == dry.fingerprint

        filtered = api.list_appfolio_migration_items(
            run.id,
            response=Response(),
            resource="properties",
            limit=200,
            db=db,
            current_user=support,
        )
        assert [row.id for row in filtered] == [rows[0].id]

        assert api.list_appfolio_migration_items(
            run.id,
            response=Response(),
            resource="UNITS",
            limit=200,
            db=db,
            current_user=support,
        ) == []

        with pytest.raises(HTTPException) as exc:
            api.list_appfolio_migration_items(
                run.id,
                response=Response(),
                resource=None,
                limit=200,
                db=db,
                current_user=sales,
            )
        assert exc.value.status_code == 403

        # A mapping cannot leak or rebind to another organization even if a
        # same numeric target id is otherwise meaningful there.
        db.get(PlatformMigrationRun, run.id).organization_id = other_org.id
        db.flush()
        with pytest.raises(HTTPException) as exc:
            api.list_appfolio_migration_items(
                run.id,
                response=Response(),
                resource=None,
                limit=200,
                db=db,
                current_user=support,
            )
        assert exc.value.status_code == 404

        assert db.query(PlatformMigrationItem).count() == 1
        assert db.query(Property).count() == 1
    finally:
        db.rollback()
        db.close()
        engine.dispose()


def test_appfolio_mapping_visibility_marks_missing_target_without_mutation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db)
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="portfolio-recovery",
            ),
            db=db,
            current_user=admin,
        )
        item = PlatformMigrationItem(
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            resource="PROPERTIES",
            source_id="AF-MISSING",
            target_entity="PROPERTY",
            target_id=999999,
            source_fingerprint="a" * 64,
            created_by_platform_user_id=admin.id,
        )
        db.add(item)
        db.commit()

        rows = api.list_appfolio_migration_items(
            run.id,
            response=Response(),
            resource="PROPERTIES",
            limit=200,
            db=db,
            current_user=admin,
        )
        assert len(rows) == 1
        assert rows[0].target_exists is False
        assert rows[0].target_label is None
        assert db.query(Property).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 CSV/XLSX ingestion + staging foundation
# ---------------------------------------------------------------------------

import asyncio
from io import BytesIO
import json as _json

from fastapi import UploadFile
from openpyxl import Workbook

from app.models.platform_migration import PlatformMigrationStagedRow, PlatformMigrationUpload


def _upload_file(name: str, content: bytes) -> UploadFile:
    return UploadFile(filename=name, file=BytesIO(content))


def test_appfolio_csv_upload_detects_stages_and_replays_without_target_mutation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="CSV Stage Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="csv-stage",
            ),
            db=db,
            current_user=admin,
        )
        content = (
            "Property Id,Property Name,Address 1,City,State,Zip,Property Type\n"
            "AF-CSV-1,Lake CSV,10 Lake Ave,Cleveland,OH,44113,Multi Family\n"
        ).encode()

        first = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("properties.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert first.detected_resource == "PROPERTIES"
        assert first.status == "STAGED"
        assert first.replayed is False
        assert first.validation_summary == {
            "total": 1,
            "valid": 1,
            "warnings": 0,
            "invalid": 0,
            "duplicates": 0,
            "possible_existing_matches": 0,
            "missing_required_columns": [],
            "ambiguous_mapping_fields": [],
        }
        assert db.query(Property).count() == 0
        assert db.query(PlatformMigrationUpload).count() == 1
        row = db.query(PlatformMigrationStagedRow).one()
        assert row.source_id == "AF-CSV-1"
        assert row.disposition == "NEW"
        assert row.normalized_data["name"] == "Lake CSV"

        audit_before = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_file_staged",
        ).count()
        replay = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("renamed.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert replay.replayed is True
        assert replay.id == first.id
        assert replay.normalized_fingerprint == first.normalized_fingerprint
        assert db.query(PlatformMigrationUpload).count() == 1
        assert db.query(PlatformMigrationStagedRow).count() == 1
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_file_staged",
        ).count() == audit_before
    finally:
        db.close()
        engine.dispose()


def test_appfolio_explicit_mapping_changes_staging_fingerprint_and_clears_old_dry_run():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Explicit Map Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="explicit-map",
            ),
            db=db,
            current_user=admin,
        )
        valid = [_valid_record()]
        dry = api.dry_run_appfolio_properties(
            run.id,
            AppFolioPropertyDryRunIn(records=valid),
            db=db,
            current_user=admin,
        )
        assert db.get(PlatformMigrationRun, run.id).last_dry_run_fingerprint == dry.fingerprint

        content = (
            "External Key,Display Label,StreetAddr,Town,Province,Postal\n"
            "PX-1,Mapped Property,1 Main St,Cleveland,OH,44113\n"
        ).encode()
        uncertain = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("uncertain.csv", content),
                resource="PROPERTIES",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert uncertain.status == "MAPPING_REQUIRED"
        assert uncertain.validation_summary["invalid"] == 1

        mapping = {
            "source_id": "External Key",
            "name": "Display Label",
            "address_line1": "StreetAddr",
            "city": "Town",
            "state": "Province",
            "zip_code": "Postal",
        }
        mapped = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("uncertain.csv", content),
                resource="PROPERTIES",
                sheet_name=None,
                column_mapping_json=_json.dumps(mapping),
                db=db,
                current_user=admin,
            )
        )
        assert mapped.status == "STAGED"
        assert mapped.normalized_fingerprint != uncertain.normalized_fingerprint
        assert mapped.column_mapping["source_id"] == "External Key"
        assert db.get(PlatformMigrationRun, run.id).last_dry_run_fingerprint is None
        assert db.get(PlatformMigrationRun, run.id).status == "STAGED"
        assert db.query(Property).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_xlsx_is_data_only_and_formula_or_credentials_fail_closed():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_DEV)
        org = _org(db, name="XLSX Stage Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="xlsx-stage",
            ),
            db=db,
            current_user=admin,
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "Property Directory"
        ws.append(["Property Id", "Property Name", "Address 1", "City", "State", "Zip"])
        ws.append(["X-1", "XLSX Property", "1 Excel Way", "Cleveland", "OH", "44113"])
        buf = BytesIO()
        wb.save(buf)
        staged = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("properties.xlsx", buf.getvalue()),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert staged.file_format == "XLSX"
        assert staged.sheet_name == "Property Directory"
        assert staged.status == "STAGED"

        wb2 = Workbook()
        ws2 = wb2.active
        ws2.append(["Property Id", "Property Name", "Address 1", "City", "State", "Zip", "Calc"])
        ws2.append(["X-2", "Formula Property", "2 Excel Way", "Cleveland", "OH", "44113", "=1+1"])
        buf2 = BytesIO()
        wb2.save(buf2)
        with pytest.raises(HTTPException) as exc:
            asyncio.run(
                api.stage_appfolio_upload(
                    run.id,
                    file=_upload_file("formula.xlsx", buf2.getvalue()),
                    resource=None,
                    sheet_name=None,
                    column_mapping_json=None,
                    db=db,
                    current_user=admin,
                )
            )
        assert exc.value.status_code == 422
        assert "formula" in exc.value.detail.lower()

        secret_csv = (
            "Property Id,Property Name,Address 1,City,State,Zip,Client Secret\n"
            "S-1,Secret,3 Secret Way,Cleveland,OH,44113,do-not-store\n"
        ).encode()
        with pytest.raises(HTTPException) as exc:
            asyncio.run(
                api.stage_appfolio_upload(
                    run.id,
                    file=_upload_file("secret.csv", secret_csv),
                    resource=None,
                    sheet_name=None,
                    column_mapping_json=None,
                    db=db,
                    current_user=admin,
                )
            )
        assert exc.value.status_code == 422
        assert "credential" in exc.value.detail.lower()
    finally:
        db.close()
        engine.dispose()


def test_appfolio_staging_dispositions_existing_mapping_possible_match_and_scope():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        support = _platform_user(db, PlatformUserRole.PLATFORM_SUPPORT)
        org = _org(db, name="Disposition Org")
        other_org = _org(db, name="Disposition Other")
        existing = Property(
            organization_id=org.id,
            name="Existing Property",
            property_type="multi_family",
            address_line1="10 Match St",
            city="Cleveland",
            state="OH",
            zip_code="44113",
            country="USA",
            is_active=True,
        )
        db.add(existing)
        db.commit()

        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="dispositions",
            ),
            db=db,
            current_user=admin,
        )
        db.add(
            PlatformMigrationItem(
                run_id=run.id,
                organization_id=org.id,
                provider="APPFOLIO",
                resource="PROPERTIES",
                source_id="ALREADY-1",
                target_entity="PROPERTY",
                target_id=existing.id,
                source_fingerprint="b" * 64,
                created_by_platform_user_id=admin.id,
            )
        )
        db.commit()

        content = (
            "Property Id,Property Name,Address 1,City,State,Zip\n"
            "ALREADY-1,Previously Imported,1 Old St,Cleveland,OH,44113\n"
            "MATCH-1,Existing Property,10 Match St,Cleveland,OH,44113\n"
            "DUP-1,First Duplicate,20 Dup St,Cleveland,OH,44113\n"
            "DUP-1,Second Duplicate,21 Dup St,Cleveland,OH,44113\n"
        ).encode()
        staged = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("dispositions.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = api.list_appfolio_staged_rows(
            run.id,
            staged.id,
            response=Response(),
            offset=0,
            limit=200,
            db=db,
            current_user=support,
        )
        assert [row.disposition for row in rows] == [
            "ALREADY_MAPPED",
            "POSSIBLE_MATCH",
            "NEW",
            "INVALID",
        ]
        assert staged.validation_summary["duplicates"] == 1
        assert staged.validation_summary["possible_existing_matches"] == 1

        other_run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=other_org.id,
                source_account_ref="other-scope",
            ),
            db=db,
            current_user=admin,
        )
        with pytest.raises(HTTPException) as exc:
            api.list_appfolio_staged_rows(
                other_run.id,
                staged.id,
                response=Response(),
                offset=0,
                limit=200,
                db=db,
                current_user=support,
            )
        assert exc.value.status_code == 404
    finally:
        db.close()
        engine.dispose()


def test_appfolio_upload_endpoint_is_exposed_and_business_targets_unchanged():
    from app.main import app

    paths = set(app.openapi()["paths"])
    assert "/api/platform/migrations/appfolio/runs/{run_id}/uploads" in paths
    assert "/api/platform/migrations/appfolio/runs/{run_id}/uploads/{upload_id}/rows" in paths



def test_staged_property_upload_reuses_existing_dry_run_engine_and_binds_source_fingerprint():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Staged Dry Run Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="staged-dry-run",
            ),
            db=db,
            current_user=admin,
        )
        content = (
            "Property Id,Property Name,Address 1,City,State,Zip,Property Type\n"
            "STAGE-1,Staged One,100 Stage St,Cleveland,OH,44113,Multi Family\n"
        ).encode()
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("properties.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        before = db.query(Property).count()
        preview = api.dry_run_staged_appfolio_properties(
            run.id,
            upload.id,
            db=db,
            current_user=admin,
        )
        assert preview.total == 1
        assert preview.importable == 1
        assert preview.invalid == 0
        assert preview.replayed is False
        assert preview.rows[0].source_id == "STAGE-1"
        assert db.query(Property).count() == before

        stored = db.get(PlatformMigrationRun, run.id)
        assert stored.last_dry_run_fingerprint == preview.fingerprint
        staged_rows = (
            db.query(PlatformMigrationStagedRow)
            .filter(PlatformMigrationStagedRow.upload_id == upload.id)
            .order_by(PlatformMigrationStagedRow.row_number.asc())
            .all()
        )
        assert (
            stored.last_dry_run_summary["source_context_fingerprint"]
            == api._staged_review_fingerprint(upload, staged_rows)
        )
        assert (
            stored.last_dry_run_summary["source_context_fingerprint"]
            != upload.normalized_fingerprint
        )
        assert preview.fingerprint != upload.normalized_fingerprint

        audit_before = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_staged_properties_dry_run",
        ).count()
        replay = api.dry_run_staged_appfolio_properties(
            run.id,
            upload.id,
            db=db,
            current_user=admin,
        )
        assert replay.replayed is True
        assert replay.fingerprint == preview.fingerprint
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_staged_properties_dry_run",
        ).count() == audit_before
        assert db.query(Property).count() == before
    finally:
        db.close()
        engine.dispose()


def test_staged_property_dry_run_blocks_invalid_possible_match_and_review_rows():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Staged Block Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="staged-block",
            ),
            db=db,
            current_user=admin,
        )

        invalid_content = (
            "Property Id,Property Name,Address 1,City,State,Zip\n"
            "BAD-1,,1 Bad St,Cleveland,OH,44113\n"
        ).encode()
        invalid_upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("invalid.csv", invalid_content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_properties(
                run.id, invalid_upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "INVALID=1" in exc.value.detail

        existing = Property(
            organization_id=org.id,
            name="Existing Stage",
            property_type="multi_family",
            address_line1="2 Match St",
            city="Cleveland",
            state="OH",
            zip_code="44113",
            country="USA",
            is_active=True,
        )
        db.add(existing)
        db.commit()
        match_content = (
            "Property Id,Property Name,Address 1,City,State,Zip\n"
            "MATCH-STAGE,Existing Stage,2 Match St,Cleveland,OH,44113\n"
        ).encode()
        match_upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("match.csv", match_content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_properties(
                run.id, match_upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "POSSIBLE_MATCH=1" in exc.value.detail

        review_content = (
            "Property Id,Property Name,Address 1,City,State,Zip,Hidden At\n"
            "HIDDEN-STAGE,Hidden Stage,3 Hide St,Cleveland,OH,44113,2026-01-01\n"
        ).encode()
        review_upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("review.csv", review_content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_properties(
                run.id, review_upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "REVIEW=1" in exc.value.detail
    finally:
        db.close()
        engine.dispose()


def test_staged_property_dry_run_allows_already_mapped_replay_and_blocks_cross_run_upload():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Staged Replay Org")
        other_org = _org(db, name="Staged Replay Other")
        target = Property(
            organization_id=org.id,
            name="Mapped Target",
            property_type="multi_family",
            address_line1="10 Target St",
            city="Cleveland",
            state="OH",
            zip_code="44113",
            country="USA",
            is_active=True,
        )
        db.add(target)
        db.commit()
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="staged-replay",
            ),
            db=db,
            current_user=admin,
        )
        db.add(
            PlatformMigrationItem(
                run_id=run.id,
                organization_id=org.id,
                provider="APPFOLIO",
                resource="PROPERTIES",
                source_id="MAPPED-1",
                target_entity="PROPERTY",
                target_id=target.id,
                source_fingerprint="c" * 64,
                created_by_platform_user_id=admin.id,
            )
        )
        db.commit()
        content = (
            "Property Id,Property Name,Address 1,City,State,Zip\n"
            "MAPPED-1,Mapped Target,10 Target St,Cleveland,OH,44113\n"
        ).encode()
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("mapped.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        staged_row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert staged_row.disposition == "ALREADY_MAPPED"

        preview = api.dry_run_staged_appfolio_properties(
            run.id, upload.id, db=db, current_user=admin
        )
        assert preview.importable == 1
        assert preview.warning_count == 2
        assert any(
            "Possible existing target property match" in warning
            for warning in preview.rows[0].warnings
        )
        assert db.query(Property).count() == 1

        other_run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=other_org.id,
                source_account_ref="other",
            ),
            db=db,
            current_user=admin,
        )
        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_properties(
                other_run.id, upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 404
    finally:
        db.close()
        engine.dispose()


def test_staged_property_dry_run_route_is_exposed():
    from app.main import app
    assert (
        "/api/platform/migrations/appfolio/runs/{run_id}/uploads/{upload_id}/properties/dry-run"
        in set(app.openapi()["paths"])
    )


def test_staged_property_resolution_match_existing_and_controlled_commit_is_replay_safe():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Resolution Match Org")
        existing = Property(
            organization_id=org.id,
            name="Existing Match",
            property_type="multi_family",
            address_line1="10 Match Way",
            city="Cleveland",
            state="OH",
            zip_code="44113",
            country="USA",
            is_active=True,
        )
        db.add(existing)
        db.commit()
        original_name = existing.name
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="resolution-match",
            ),
            db=db,
            current_user=admin,
        )
        content = (
            "Property Id,Property Name,Address 1,City,State,Zip\n"
            "MATCH-RES,Existing Match,10 Match Way,Cleveland,OH,44113\n"
            "NEW-RES,New Property,20 New Way,Cleveland,OH,44113\n"
        ).encode()
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("resolution.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = api.list_appfolio_staged_rows(
            run.id, upload.id, response=Response(), offset=0, limit=200,
            db=db, current_user=admin,
        )
        match_row = next(row for row in rows if row.source_id == "MATCH-RES")
        assert match_row.disposition == "POSSIBLE_MATCH"

        with pytest.raises(HTTPException):
            api.dry_run_staged_appfolio_properties(
                run.id, upload.id, db=db, current_user=admin
            )

        resolved = api.resolve_staged_appfolio_property(
            run.id,
            upload.id,
            match_row.id,
            AppFolioStagedRowResolutionIn(
                action="MATCH_EXISTING",
                target_property_id=existing.id,
            ),
            db=db,
            current_user=admin,
        )
        assert resolved.resolution_action == "MATCH_EXISTING"
        assert resolved.resolution_target_id == existing.id

        preview = api.dry_run_staged_appfolio_properties(
            run.id, upload.id, db=db, current_user=admin
        )
        assert preview.total == 2 and preview.importable == 2
        committed = api.commit_staged_appfolio_properties(
            run.id,
            upload.id,
            AppFolioStagedPropertyCommitIn(fingerprint=preview.fingerprint),
            db=db,
            current_user=admin,
        )
        assert committed.committed == 1
        assert committed.matched_existing == 1
        assert committed.replayed is False
        assert db.query(Property).filter(Property.organization_id == org.id).count() == 2
        assert db.get(Property, existing.id).name == original_name
        mappings = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "PROPERTIES",
        ).all()
        assert {item.source_id for item in mappings} == {"MATCH-RES", "NEW-RES"}
        assert next(item for item in mappings if item.source_id == "MATCH-RES").target_id == existing.id

        replay = api.commit_staged_appfolio_properties(
            run.id,
            upload.id,
            AppFolioStagedPropertyCommitIn(fingerprint=preview.fingerprint),
            db=db,
            current_user=admin,
        )
        assert replay.replayed is True
        assert replay.committed == 0
        assert replay.matched_existing == 0
        assert db.query(Property).filter(Property.organization_id == org.id).count() == 2
    finally:
        db.close()
        engine.dispose()


def test_staged_property_create_new_resolution_binds_fingerprint_and_old_preview_goes_stale():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Resolution Create Org")
        existing = Property(
            organization_id=org.id,
            name="Reviewed Possible",
            property_type="multi_family",
            address_line1="30 Review Rd",
            city="Cleveland",
            state="OH",
            zip_code="44113",
            country="USA",
            is_active=True,
        )
        db.add(existing)
        db.commit()
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="resolution-create",
            ),
            db=db,
            current_user=admin,
        )
        content = (
            "Property Id,Property Name,Address 1,City,State,Zip\n"
            "CREATE-RES,Reviewed Possible,30 Review Rd,Cleveland,OH,44113\n"
        ).encode()
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("create-new.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        api.resolve_staged_appfolio_property(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedRowResolutionIn(action="CREATE_NEW"),
            db=db,
            current_user=admin,
        )
        preview_create = api.dry_run_staged_appfolio_properties(
            run.id, upload.id, db=db, current_user=admin
        )
        assert any(
            "Possible existing target property match" in warning
            for warning in preview_create.rows[0].warnings
        )

        api.resolve_staged_appfolio_property(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedRowResolutionIn(
                action="MATCH_EXISTING",
                target_property_id=existing.id,
            ),
            db=db,
            current_user=admin,
        )
        assert db.get(PlatformMigrationRun, run.id).last_dry_run_fingerprint is None
        with pytest.raises(HTTPException) as exc:
            api.commit_staged_appfolio_properties(
                run.id,
                upload.id,
                AppFolioStagedPropertyCommitIn(fingerprint=preview_create.fingerprint),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409

        preview_match = api.dry_run_staged_appfolio_properties(
            run.id, upload.id, db=db, current_user=admin
        )
        assert preview_match.fingerprint != preview_create.fingerprint
        matched = api.commit_staged_appfolio_properties(
            run.id,
            upload.id,
            AppFolioStagedPropertyCommitIn(fingerprint=preview_match.fingerprint),
            db=db,
            current_user=admin,
        )
        assert matched.committed == 0 and matched.matched_existing == 1
        assert db.query(Property).filter(Property.organization_id == org.id).count() == 1
    finally:
        db.close()
        engine.dispose()


def test_staged_property_create_new_and_skip_are_explicit_and_scoped():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        support = _platform_user(db, PlatformUserRole.PLATFORM_SUPPORT)
        org = _org(db, name="Resolution Scope Org")
        other = _org(db, name="Resolution Scope Other")
        existing = Property(
            organization_id=org.id,
            name="Possible Create",
            property_type="multi_family",
            address_line1="40 Create St",
            city="Cleveland",
            state="OH",
            zip_code="44113",
            country="USA",
            is_active=True,
        )
        foreign = Property(
            organization_id=other.id,
            name="Foreign",
            property_type="multi_family",
            address_line1="99 Foreign St",
            city="Cleveland",
            state="OH",
            zip_code="44113",
            country="USA",
            is_active=True,
        )
        db.add_all([existing, foreign])
        db.commit()
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="resolution-scope",
            ),
            db=db,
            current_user=admin,
        )
        content = (
            "Property Id,Property Name,Address 1,City,State,Zip,Hidden At\n"
            "CREATE-NEW,Possible Create,40 Create St,Cleveland,OH,44113,\n"
            "SKIP-HIDDEN,Hidden Source,50 Hidden St,Cleveland,OH,44113,2026-01-01\n"
        ).encode()
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("scope.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()

        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_property(
                run.id,
                upload.id,
                rows[0].id,
                AppFolioStagedRowResolutionIn(
                    action="MATCH_EXISTING",
                    target_property_id=foreign.id,
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 404

        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_property(
                run.id,
                upload.id,
                rows[0].id,
                AppFolioStagedRowResolutionIn(action="CREATE_NEW"),
                db=db,
                current_user=support,
            )
        assert exc.value.status_code == 403

        api.resolve_staged_appfolio_property(
            run.id,
            upload.id,
            rows[0].id,
            AppFolioStagedRowResolutionIn(action="CREATE_NEW"),
            db=db,
            current_user=admin,
        )
        api.resolve_staged_appfolio_property(
            run.id,
            upload.id,
            rows[1].id,
            AppFolioStagedRowResolutionIn(action="SKIP"),
            db=db,
            current_user=admin,
        )
        preview = api.dry_run_staged_appfolio_properties(
            run.id, upload.id, db=db, current_user=admin
        )
        assert preview.total == 1 and preview.rows[0].source_id == "CREATE-NEW"
        committed = api.commit_staged_appfolio_properties(
            run.id,
            upload.id,
            AppFolioStagedPropertyCommitIn(fingerprint=preview.fingerprint),
            db=db,
            current_user=admin,
        )
        assert committed.committed == 1
        assert db.query(Property).filter(Property.organization_id == org.id).count() == 2
        assert db.query(Property).filter(
            Property.organization_id == org.id,
            Property.name == "Hidden Source",
        ).count() == 0

        resolution_audits = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.action == "appfolio_property_resolution_changed",
        ).count()
        assert resolution_audits == 2
    finally:
        db.close()
        engine.dispose()


def test_staged_property_resolution_and_commit_routes_are_exposed():
    from app.main import app

    paths = set(app.openapi()["paths"])
    assert (
        "/api/platform/migrations/appfolio/runs/{run_id}/uploads/{upload_id}/rows/{staged_row_id}/resolution"
        in paths
    )
    assert (
        "/api/platform/migrations/appfolio/runs/{run_id}/uploads/{upload_id}/properties/commit"
        in paths
    )


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Unit Directory ingestion + staging
# ---------------------------------------------------------------------------

def _mapped_property_for_units(db, *, run, org, admin, source_id="AF-PROP-UNIT"):
    target = Property(
        organization_id=org.id,
        name="Mapped Unit Property",
        property_type="multi_family",
        address_line1="100 Unit Way",
        city="Cleveland",
        state="OH",
        zip_code="44113",
        country="USA",
        is_active=True,
    )
    db.add(target)
    db.flush()
    db.add(
        PlatformMigrationItem(
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            resource="PROPERTIES",
            source_id=source_id,
            target_entity="PROPERTY",
            target_id=target.id,
            source_fingerprint="d" * 64,
            created_by_platform_user_id=admin.id,
        )
    )
    db.commit()
    return target


def test_appfolio_unit_directory_csv_stages_against_durable_property_mapping_and_replays():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Unit Stage Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="unit-stage",
            ),
            db=db,
            current_user=admin,
        )
        target = _mapped_property_for_units(
            db, run=run, org=org, admin=admin, source_id="PROP-100"
        )
        content = (
            "Unit ID,Unit Name,Property ID,Property Name,Unit Address,"
            "Unit Street Address 1,Unit Street Address 2,Unit City,Unit State,Unit Zip\n"
            "UNIT-1,101,PROP-100,Mapped Unit Property,100 Unit Way #101,"
            "100 Unit Way,,Cleveland,OH,44113\n"
        ).encode()

        first = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("unit-directory.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert first.detected_resource == "UNITS"
        assert first.status == "STAGED"
        assert first.replayed is False
        assert first.validation_summary["total"] == 1
        assert first.validation_summary["valid"] == 1
        assert first.validation_summary["invalid"] == 0
        assert first.validation_summary["warnings"] == 0
        assert first.validation_summary["unresolved_property_links"] == 0
        assert first.validation_summary["missing_unit_source_ids"] == 0

        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == first.id
        ).one()
        assert row.resource == "UNITS"
        assert row.source_id == "UNIT-1"
        assert row.disposition == "NEW"
        assert row.normalized_data["source_property_id"] == "PROP-100"
        assert row.normalized_data["unit_name"] == "101"
        assert row.normalized_data["address_line1"] == "100 Unit Way"
        assert db.query(Unit).count() == 0
        assert db.get(Property, target.id).name == "Mapped Unit Property"

        replay = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("same-units-renamed.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert replay.replayed is True
        assert replay.id == first.id
        assert db.query(PlatformMigrationUpload).filter(
            PlatformMigrationUpload.run_id == run.id,
            PlatformMigrationUpload.detected_resource == "UNITS",
        ).count() == 1
        assert db.query(Unit).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_unit_directory_missing_or_unresolved_ids_stage_for_review_not_invention():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Unit Review Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="unit-review",
            ),
            db=db,
            current_user=admin,
        )
        _mapped_property_for_units(
            db, run=run, org=org, admin=admin, source_id="PROP-OK"
        )
        content = (
            "Unit ID,Unit Name,Property ID,Property Name,Unit Street Address 1,"
            "Unit City,Unit State,Unit Zip\n"
            ",201,PROP-OK,Mapped Unit Property,100 Unit Way,Cleveland,OH,44113\n"
            "UNIT-UNKNOWN,202,PROP-MISSING,Unknown Property,200 Missing Way,Cleveland,OH,44113\n"
        ).encode()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("unit-review.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert upload.detected_resource == "UNITS"
        assert upload.status == "REVIEW_REQUIRED"
        assert upload.validation_summary["valid"] == 2
        assert upload.validation_summary["warnings"] == 2
        assert upload.validation_summary["unresolved_property_links"] == 1
        assert upload.validation_summary["missing_unit_source_ids"] == 1

        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()
        assert [row.disposition for row in rows] == ["REVIEW", "REVIEW"]
        assert rows[0].source_id is None
        assert any("Unit ID was not supplied" in warning for warning in rows[0].warnings)
        assert rows[1].source_id == "UNIT-UNKNOWN"
        assert any("has no durable Property mapping" in warning for warning in rows[1].warnings)
        assert db.query(Unit).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_unit_directory_duplicate_unit_id_is_invalid_and_xlsx_autodetects():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Unit XLSX Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="unit-xlsx",
            ),
            db=db,
            current_user=admin,
        )
        _mapped_property_for_units(
            db, run=run, org=org, admin=admin, source_id="PROP-XLSX"
        )

        workbook = Workbook()
        ws = workbook.active
        ws.title = "Unit Directory"
        ws.append([
            "Unit ID", "Unit Name", "Property ID", "Property Name",
            "Unit Street Address 1", "Unit City", "Unit State", "Unit Zip",
        ])
        ws.append([
            "UNIT-X", "A", "PROP-XLSX", "Mapped Unit Property",
            "100 Unit Way", "Cleveland", "OH", "44113",
        ])
        ws.append([
            "UNIT-X", "B", "PROP-XLSX", "Mapped Unit Property",
            "100 Unit Way", "Cleveland", "OH", "44113",
        ])
        notes = workbook.create_sheet("Notes")
        notes.append(["Comment"])
        notes.append(["not a supported migration report"])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("unit-directory.xlsx", buffer.getvalue()),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert upload.detected_resource == "UNITS"
        assert upload.sheet_name == "Unit Directory"
        assert upload.status == "STAGED_WITH_ERRORS"
        assert upload.validation_summary["duplicates"] == 1
        assert upload.validation_summary["invalid"] == 1

        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()
        assert rows[0].disposition == "NEW"
        assert rows[1].disposition == "INVALID"
        assert any("Duplicate AppFolio Unit ID" in error for error in rows[1].errors)
        assert db.query(Unit).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 staged Unit dry run + controlled commit
# ---------------------------------------------------------------------------

def _stage_safe_unit_upload(db, *, api_user, run, property_source_id="PROP-UNIT-COMMIT", unit_id="UNIT-COMMIT-1", unit_name="301"):
    content = (
        "Unit ID,Unit Name,Property ID,Property Name,Unit Street Address 1,"
        "Unit City,Unit State,Unit Zip\n"
        f"{unit_id},{unit_name},{property_source_id},Mapped Unit Property,"
        "100 Unit Way,Cleveland,OH,44113\n"
    ).encode()
    return asyncio.run(
        api.stage_appfolio_upload(
            run.id,
            file=_upload_file("units-safe.csv", content),
            resource=None,
            sheet_name=None,
            column_mapping_json=None,
            db=db,
            current_user=api_user,
        )
    )


def test_staged_unit_dry_run_is_non_mutating_and_commit_replays_without_duplicates():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Unit Commit Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="unit-commit",
            ),
            db=db,
            current_user=admin,
        )
        target_property = _mapped_property_for_units(
            db,
            run=run,
            org=org,
            admin=admin,
            source_id="PROP-UNIT-COMMIT",
        )
        upload = _stage_safe_unit_upload(db, api_user=admin, run=run)

        before = db.query(Unit).count()
        preview = api.dry_run_staged_appfolio_units(
            run.id, upload.id, db=db, current_user=admin
        )
        assert preview.total == 1
        assert preview.importable == 1
        assert preview.invalid == 0
        assert preview.replayed is False
        assert preview.rows[0].mapped["property_id"] == target_property.id
        assert preview.rows[0].mapped["unit_number"] == "301"
        assert any("target Unit defaults apply" in warning for warning in preview.rows[0].warnings)
        assert db.query(Unit).count() == before

        committed = api.commit_staged_appfolio_units(
            run.id,
            upload.id,
            AppFolioStagedUnitCommitIn(fingerprint=preview.fingerprint),
            db=db,
            current_user=admin,
        )
        assert committed.committed == 1
        assert committed.replayed is False
        unit = db.query(Unit).filter(
            Unit.property_id == target_property.id,
            Unit.unit_number == "301",
        ).one()
        mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "UNITS",
            PlatformMigrationItem.source_id == "UNIT-COMMIT-1",
        ).one()
        assert mapping.target_entity == "UNIT"
        assert mapping.target_id == unit.id

        replay = api.commit_staged_appfolio_units(
            run.id,
            upload.id,
            AppFolioStagedUnitCommitIn(fingerprint=preview.fingerprint),
            db=db,
            current_user=admin,
        )
        assert replay.replayed is True
        assert replay.committed == 0
        assert replay.rows[0].target_unit_id == unit.id
        assert db.query(Unit).filter(Unit.property_id == target_property.id).count() == 1

        audit_count = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_staged_units_committed",
        ).count()
        assert audit_count == 1
    finally:
        db.close()
        engine.dispose()


def test_staged_unit_match_existing_resolution_is_explicit_non_overwriting_and_replay_safe():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Unit Existing Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="unit-existing",
            ),
            db=db,
            current_user=admin,
        )
        target_property = _mapped_property_for_units(
            db,
            run=run,
            org=org,
            admin=admin,
            source_id="PROP-UNIT-COMMIT",
        )
        existing = Unit(
            property_id=target_property.id,
            unit_number="301",
        )
        db.add(existing)
        db.commit()

        upload = _stage_safe_unit_upload(db, api_user=admin, run=run)
        staged = api.list_appfolio_staged_rows(
            run.id,
            upload.id,
            Response(),
            offset=0,
            limit=200,
            db=db,
            current_user=admin,
        )
        assert len(staged) == 1
        assert staged[0].disposition == "POSSIBLE_MATCH"

        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_units(
                run.id, upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "POSSIBLE_MATCH=1" in exc.value.detail

        resolved = api.resolve_staged_appfolio_unit(
            run.id,
            upload.id,
            staged[0].id,
            AppFolioStagedUnitResolutionIn(
                action="MATCH_EXISTING",
                target_unit_id=existing.id,
            ),
            db=db,
            current_user=admin,
        )
        assert resolved.resolution_action == "MATCH_EXISTING"
        assert resolved.resolution_target_unit_id == existing.id

        preview = api.dry_run_staged_appfolio_units(
            run.id, upload.id, db=db, current_user=admin
        )
        assert preview.importable == 1
        assert any(
            warning.startswith("Explicitly matched existing target unit")
            for warning in preview.rows[0].warnings
        )

        before_number = existing.unit_number
        committed = api.commit_staged_appfolio_units(
            run.id,
            upload.id,
            AppFolioStagedUnitCommitIn(fingerprint=preview.fingerprint),
            db=db,
            current_user=admin,
        )
        assert committed.committed == 0
        assert committed.matched_existing == 1
        assert db.get(Unit, existing.id).unit_number == before_number
        mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "UNITS",
            PlatformMigrationItem.source_id == "UNIT-COMMIT-1",
        ).one()
        assert mapping.target_id == existing.id

        replay = api.commit_staged_appfolio_units(
            run.id,
            upload.id,
            AppFolioStagedUnitCommitIn(fingerprint=preview.fingerprint),
            db=db,
            current_user=admin,
        )
        assert replay.replayed is True
        assert replay.committed == 0
        assert replay.matched_existing == 0
        assert db.query(Unit).filter(Unit.property_id == target_property.id).count() == 1
    finally:
        db.close()
        engine.dispose()


def test_staged_unit_create_new_and_skip_resolution_invalidate_preview_and_do_not_overwrite():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Unit Resolution Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="unit-resolution",
            ),
            db=db,
            current_user=admin,
        )
        target_property = _mapped_property_for_units(
            db,
            run=run,
            org=org,
            admin=admin,
            source_id="PROP-UNIT-COMMIT",
        )
        existing = Unit(property_id=target_property.id, unit_number="ABC")
        db.add(existing)
        db.commit()

        upload = _stage_safe_unit_upload(db, api_user=admin, run=run, unit_name="abc")
        staged = api.list_appfolio_staged_rows(
            run.id,
            upload.id,
            Response(),
            offset=0,
            limit=200,
            db=db,
            current_user=admin,
        )
        row = staged[0]

        api.resolve_staged_appfolio_unit(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedUnitResolutionIn(action="CREATE_NEW"),
            db=db,
            current_user=admin,
        )
        preview = api.dry_run_staged_appfolio_units(
            run.id, upload.id, db=db, current_user=admin
        )
        assert any(
            warning.startswith("Explicit CREATE_NEW resolution")
            for warning in preview.rows[0].warnings
        )

        api.resolve_staged_appfolio_unit(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedUnitResolutionIn(action="SKIP"),
            db=db,
            current_user=admin,
        )
        with pytest.raises(HTTPException) as stale:
            api.commit_staged_appfolio_units(
                run.id,
                upload.id,
                AppFolioStagedUnitCommitIn(fingerprint=preview.fingerprint),
                db=db,
                current_user=admin,
            )
        assert stale.value.status_code == 409
        assert "No staged Unit rows remain" in stale.value.detail

        api.resolve_staged_appfolio_unit(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedUnitResolutionIn(action="CREATE_NEW"),
            db=db,
            current_user=admin,
        )
        fresh = api.dry_run_staged_appfolio_units(
            run.id, upload.id, db=db, current_user=admin
        )
        committed = api.commit_staged_appfolio_units(
            run.id,
            upload.id,
            AppFolioStagedUnitCommitIn(fingerprint=fresh.fingerprint),
            db=db,
            current_user=admin,
        )
        assert committed.committed == 1
        assert committed.matched_existing == 0
        assert db.get(Unit, existing.id).unit_number == "ABC"
        assert db.query(Unit).filter(Unit.property_id == target_property.id).count() == 2

        resolution_audit = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == row.id,
            AuditLog.action == "appfolio_unit_resolution_changed",
        ).count()
        assert resolution_audit == 3
    finally:
        db.close()
        engine.dispose()


def test_staged_unit_commit_fails_when_property_relationship_fingerprint_changes():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Unit Relationship Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="unit-relationship",
            ),
            db=db,
            current_user=admin,
        )
        _mapped_property_for_units(
            db,
            run=run,
            org=org,
            admin=admin,
            source_id="PROP-UNIT-COMMIT",
        )
        upload = _stage_safe_unit_upload(db, api_user=admin, run=run)
        preview = api.dry_run_staged_appfolio_units(
            run.id, upload.id, db=db, current_user=admin
        )

        property_mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "PROPERTIES",
            PlatformMigrationItem.source_id == "PROP-UNIT-COMMIT",
        ).one()
        property_mapping.source_fingerprint = "e" * 64
        db.commit()

        with pytest.raises(HTTPException) as exc:
            api.commit_staged_appfolio_units(
                run.id,
                upload.id,
                AppFolioStagedUnitCommitIn(fingerprint=preview.fingerprint),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert db.query(Unit).count() == 0

        refreshed = api.dry_run_staged_appfolio_units(
            run.id, upload.id, db=db, current_user=admin
        )
        assert refreshed.fingerprint != preview.fingerprint
    finally:
        db.close()
        engine.dispose()


def test_staged_unit_dry_run_and_commit_routes_are_exposed():
    from app.main import app

    paths = set(app.openapi()["paths"])
    assert (
        "/api/platform/migrations/appfolio/runs/{run_id}/uploads/{upload_id}/units/dry-run"
        in paths
    )
    assert (
        "/api/platform/migrations/appfolio/runs/{run_id}/uploads/{upload_id}/units/commit"
        in paths
    )


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Owner Directory ingestion + staging
# ---------------------------------------------------------------------------

def test_appfolio_owner_directory_csv_stages_verified_fields_and_replays_without_customer_mutation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Owner Stage Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="owner-stage",
            ),
            db=db,
            current_user=admin,
        )
        content = (
            "Owner ID,Name,Phone Numbers,Email,Properties Owned,Properties Owned IDs\n"
            "OWNER-1,Jane Owner,216-555-0100,jane.owner@example.com,Lake Apartments,PROP-100\n"
        ).encode()

        before_users = db.query(User).count()
        first = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("owner-directory.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert first.detected_resource == "OWNERS"
        assert first.status == "STAGED"
        assert first.replayed is False
        assert first.validation_summary["total"] == 1
        assert first.validation_summary["valid"] == 1
        assert first.validation_summary["invalid"] == 0
        assert first.validation_summary["warnings"] == 0
        assert first.validation_summary["missing_owner_source_ids"] == 0
        assert first.validation_summary["missing_owner_emails"] == 0
        assert first.validation_summary["missing_owner_property_ids"] == 0

        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == first.id
        ).one()
        assert row.resource == "OWNERS"
        assert row.source_id == "OWNER-1"
        assert row.disposition == "NEW"
        assert row.normalized_data == {
            "source_id": "OWNER-1",
            "name": "Jane Owner",
            "phone_numbers": "216-555-0100",
            "email": "jane.owner@example.com",
            "properties_owned": "Lake Apartments",
            "properties_owned_ids": "PROP-100",
        }
        assert db.query(User).count() == before_users

        replay = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("renamed-owner-directory.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert replay.replayed is True
        assert replay.id == first.id
        assert db.query(PlatformMigrationUpload).filter(
            PlatformMigrationUpload.run_id == run.id,
            PlatformMigrationUpload.detected_resource == "OWNERS",
        ).count() == 1
        assert db.query(User).count() == before_users
    finally:
        db.close()
        engine.dispose()


def test_appfolio_owner_directory_missing_identity_or_relationship_stages_review_and_exact_email_is_possible_match():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Owner Review Org")
        existing = User(
            organization_id=org.id,
            role=UserRole.OWNER,
            email="existing.owner@example.com",
            first_name="Existing",
            last_name="Owner",
            hashed_password="x",
            is_active=True,
        )
        db.add(existing)
        db.commit()
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="owner-review",
            ),
            db=db,
            current_user=admin,
        )
        content = (
            "Owner ID,Name,Phone Numbers,Email,Properties Owned,Properties Owned IDs\n"
            "OWNER-MATCH,Existing Owner,216-555-0101,existing.owner@example.com,Lake Apartments,PROP-100\n"
            ",Missing Stable Id,216-555-0102,missing.id@example.com,Lake Apartments,PROP-100\n"
            "OWNER-NOEMAIL,No Email,216-555-0103,,Lake Apartments,PROP-100\n"
            "OWNER-NOPROP,No Property Refs,216-555-0104,noprop@example.com,,\n"
        ).encode()

        before_users = db.query(User).count()
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("owner-review.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert upload.detected_resource == "OWNERS"
        assert upload.status == "REVIEW_REQUIRED"
        assert upload.validation_summary["valid"] == 4
        assert upload.validation_summary["warnings"] == 4
        assert upload.validation_summary["possible_existing_matches"] == 1
        assert upload.validation_summary["missing_owner_source_ids"] == 1
        assert upload.validation_summary["missing_owner_emails"] == 1
        assert upload.validation_summary["missing_owner_property_ids"] == 1

        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()
        assert [row.disposition for row in rows] == [
            "POSSIBLE_MATCH", "REVIEW", "REVIEW", "REVIEW"
        ]
        assert any("exact email" in warning for warning in rows[0].warnings)
        assert any("Owner ID was not supplied" in warning for warning in rows[1].warnings)
        assert any("Owner email was not supplied" in warning for warning in rows[2].warnings)
        assert any("Properties Owned IDs were not supplied" in warning for warning in rows[3].warnings)
        assert db.query(User).count() == before_users
        assert db.get(User, existing.id).email == "existing.owner@example.com"
    finally:
        db.close()
        engine.dispose()


def test_appfolio_owner_directory_duplicate_source_id_invalid_and_multisheet_xlsx_autodetects():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Owner XLSX Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="owner-xlsx",
            ),
            db=db,
            current_user=admin,
        )

        workbook = Workbook()
        ws = workbook.active
        ws.title = "Owner Directory"
        ws.append([
            "Owner ID", "Name", "Phone Numbers", "Email",
            "Properties Owned", "Properties Owned IDs",
        ])
        ws.append([
            "OWNER-X", "Owner One", "216-555-0105", "owner.one@example.com",
            "Lake Apartments", "PROP-100",
        ])
        ws.append([
            "OWNER-X", "Owner Duplicate", "216-555-0106", "owner.two@example.com",
            "River Apartments", "PROP-200",
        ])
        notes = workbook.create_sheet("Notes")
        notes.append(["Comment"])
        notes.append(["not a supported migration report"])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("owner-directory.xlsx", buffer.getvalue()),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert upload.detected_resource == "OWNERS"
        assert upload.sheet_name == "Owner Directory"
        assert upload.status == "STAGED_WITH_ERRORS"
        assert upload.validation_summary["duplicates"] == 1
        assert upload.validation_summary["invalid"] == 1

        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()
        assert rows[0].disposition == "NEW"
        assert rows[1].disposition == "INVALID"
        assert any("Duplicate AppFolio Owner ID" in error for error in rows[1].errors)
        assert db.query(User).filter(User.role == UserRole.OWNER).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Vendor Directory ingestion + staging
# ---------------------------------------------------------------------------

def test_appfolio_vendor_directory_csv_stages_verified_fields_and_replays_without_vendor_mutation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Vendor Stage Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="vendor-stage",
            ),
            db=db,
            current_user=admin,
        )
        content = (
            "Vendor ID,Company Name,Address,Phone Numbers,Email,Send 1099?,"
            "Liability Insurance Expiration,Workers Comp Expiration,"
            "EPA Certification Expiration,State License Expiration,Contract Expiration\n"
            "VENDOR-1,ABC Plumbing,10 Trade St Cleveland OH 44113,216-555-0100,"
            "service@abc.example,Yes,2027-01-31,2027-02-28,2027-03-31,2027-04-30,2027-12-31\n"
        ).encode()

        before = db.query(Vendor).count()
        first = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("vendor-directory.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert first.detected_resource == "VENDORS"
        assert first.status == "STAGED"
        assert first.validation_summary["valid"] == 1
        assert first.validation_summary["invalid"] == 0
        assert first.validation_summary["warnings"] == 0
        assert first.validation_summary["missing_vendor_source_ids"] == 0

        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == first.id
        ).one()
        assert row.source_id == "VENDOR-1"
        assert row.disposition == "NEW"
        assert row.normalized_data["company_name"] == "ABC Plumbing"
        assert row.normalized_data["email"] == "service@abc.example"
        assert row.normalized_data["send_1099"] == "Yes"
        assert row.normalized_data["liability_insurance_expiration"] == "2027-01-31"
        assert row.normalized_data["workers_comp_expiration"] == "2027-02-28"
        assert row.normalized_data["epa_certification_expiration"] == "2027-03-31"
        assert row.normalized_data["state_license_expiration"] == "2027-04-30"
        assert row.normalized_data["contract_expiration"] == "2027-12-31"
        assert db.query(Vendor).count() == before

        replay = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("vendor-directory-renamed.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert replay.replayed is True
        assert replay.id == first.id
        assert db.query(Vendor).count() == before
    finally:
        db.close()
        engine.dispose()


def test_appfolio_vendor_directory_missing_source_id_and_existing_target_are_reviewed_not_overwritten():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Vendor Review Org")
        existing = Vendor(
            organization_id=org.id,
            company_name="Existing Vendor",
            business_email="existing.vendor@example.com",
            is_active=True,
        )
        db.add(existing)
        db.commit()
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="vendor-review",
            ),
            db=db,
            current_user=admin,
        )
        content = (
            "Vendor ID,Company Name,Address,Phone Numbers,Email,Send 1099?\n"
            "V-MATCH,Existing Vendor,20 Match St,216-555-0111,existing.vendor@example.com,No\n"
            ",Missing Source Id,30 Review St,216-555-0112,new.vendor@example.com,No\n"
        ).encode()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("vendor-review.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert upload.detected_resource == "VENDORS"
        assert upload.status == "REVIEW_REQUIRED"
        assert upload.validation_summary["possible_existing_matches"] == 1
        assert upload.validation_summary["missing_vendor_source_ids"] == 1

        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()
        assert [row.disposition for row in rows] == ["POSSIBLE_MATCH", "REVIEW"]
        assert any("Possible existing target vendor match" in w for w in rows[0].warnings)
        assert any("Vendor ID was not supplied" in w for w in rows[1].warnings)
        assert db.query(Vendor).count() == 1
        assert db.get(Vendor, existing.id).company_name == "Existing Vendor"
    finally:
        db.close()
        engine.dispose()


def test_appfolio_vendor_directory_duplicate_source_id_invalid_and_xlsx_autodetects():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Vendor XLSX Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="vendor-xlsx",
            ),
            db=db,
            current_user=admin,
        )

        workbook = Workbook()
        ws = workbook.active
        ws.title = "Vendor Directory"
        ws.append([
            "Vendor ID", "Company Name", "Address", "Phone Numbers", "Email",
            "Send 1099?", "Liability Insurance Expiration",
        ])
        ws.append([
            "VENDOR-X", "Vendor One", "1 Vendor St", "216-555-0201",
            "one@vendor.example", "Yes", "2027-01-31",
        ])
        ws.append([
            "VENDOR-X", "Vendor Duplicate", "2 Vendor St", "216-555-0202",
            "two@vendor.example", "No", "2027-02-28",
        ])
        notes = workbook.create_sheet("Notes")
        notes.append(["Comment"])
        notes.append(["not a supported migration report"])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("vendor-directory.xlsx", buffer.getvalue()),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert upload.detected_resource == "VENDORS"
        assert upload.sheet_name == "Vendor Directory"
        assert upload.status == "STAGED_WITH_ERRORS"
        assert upload.validation_summary["duplicates"] == 1
        assert upload.validation_summary["invalid"] == 1

        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()
        assert rows[0].disposition == "NEW"
        assert rows[1].disposition == "INVALID"
        assert any("Duplicate AppFolio Vendor ID" in e for e in rows[1].errors)
        assert db.query(Vendor).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Owner/Vendor staged reconciliation foundation
# ---------------------------------------------------------------------------

def test_appfolio_owner_possible_match_resolution_is_typed_audited_and_fingerprint_invalidating():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Owner Resolve Org")
        existing = User(
            organization_id=org.id,
            role=UserRole.OWNER,
            email="owner.resolve@example.com",
            first_name="Existing",
            last_name="Owner",
            hashed_password="x",
            is_active=True,
        )
        db.add(existing)
        db.commit()

        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="owner-resolve",
            ),
            db=db,
            current_user=admin,
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "owner-directory.csv",
                    (
                        "Owner ID,Name,Phone Numbers,Email,Properties Owned,Properties Owned IDs\n"
                        "OWNER-R1,Existing Owner,216-555-0101,owner.resolve@example.com,"
                        "Lake Apartments,PROP-1\n"
                    ).encode(),
                ),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.disposition == "POSSIBLE_MATCH"
        before_fingerprint = api._staged_review_fingerprint(upload, [row])

        run.last_dry_run_fingerprint = "a" * 64
        run.last_dry_run_summary = {"old": True}
        db.commit()
        original_email = existing.email

        resolved = api.resolve_staged_appfolio_owner(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedOwnerResolutionIn(
                action="MATCH_EXISTING",
                target_owner_user_id=existing.id,
            ),
            db=db,
            current_user=admin,
        )
        assert resolved.resolution_action == "MATCH_EXISTING"
        assert resolved.resolution_target_owner_user_id == existing.id
        assert resolved.resolution_target_id is None
        assert resolved.resolution_target_unit_id is None
        assert resolved.resolution_target_vendor_id is None
        assert db.get(User, existing.id).email == original_email
        db.refresh(run)
        assert run.last_dry_run_fingerprint is None
        assert run.last_dry_run_summary is None
        after_fingerprint = api._staged_review_fingerprint(upload, [resolved])
        assert after_fingerprint != before_fingerprint
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == row.id,
            AuditLog.action == "appfolio_owner_resolution_changed",
        ).count() == 1

        changed = api.resolve_staged_appfolio_owner(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedOwnerResolutionIn(action="CREATE_NEW"),
            db=db,
            current_user=admin,
        )
        assert changed.resolution_action == "CREATE_NEW"
        assert changed.resolution_target_owner_user_id is None
        assert db.get(User, existing.id).email == original_email
    finally:
        db.close()
        engine.dispose()


def test_appfolio_owner_resolution_rejects_foreign_target_and_missing_source_identity():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Owner Resolve Local")
        foreign_org = _org(db, name="Owner Resolve Foreign")
        foreign_owner = User(
            organization_id=foreign_org.id,
            role=UserRole.OWNER,
            email="foreign.owner.resolve@example.com",
            first_name="Foreign",
            last_name="Owner",
            hashed_password="x",
            is_active=True,
        )
        db.add(foreign_owner)
        db.commit()
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="owner-resolve-scope",
            ),
            db=db,
            current_user=admin,
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "owner-review.csv",
                    (
                        "Owner ID,Name,Phone Numbers,Email,Properties Owned,Properties Owned IDs\n"
                        "OWNER-SCOPE,Scope Owner,216-555-0102,scope.owner@example.com,Lake,PROP-1\n"
                        ",Missing ID,216-555-0103,missing.id@example.com,Lake,PROP-1\n"
                    ).encode(),
                ),
                resource="OWNERS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()
        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_owner(
                run.id,
                upload.id,
                rows[0].id,
                AppFolioStagedOwnerResolutionIn(
                    action="MATCH_EXISTING",
                    target_owner_user_id=foreign_owner.id,
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 404
        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_owner(
                run.id,
                upload.id,
                rows[1].id,
                AppFolioStagedOwnerResolutionIn(action="CREATE_NEW"),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        skipped = api.resolve_staged_appfolio_owner(
            run.id,
            upload.id,
            rows[1].id,
            AppFolioStagedOwnerResolutionIn(action="SKIP"),
            db=db,
            current_user=admin,
        )
        assert skipped.resolution_action == "SKIP"
    finally:
        db.close()
        engine.dispose()


def test_appfolio_vendor_possible_match_resolution_is_typed_scoped_and_non_mutating():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Vendor Resolve Org")
        foreign_org = _org(db, name="Vendor Resolve Foreign")
        existing = Vendor(
            organization_id=org.id,
            company_name="Resolve Plumbing",
            business_email="resolve.vendor@example.com",
            is_active=True,
        )
        foreign = Vendor(
            organization_id=foreign_org.id,
            company_name="Foreign Plumbing",
            business_email="foreign.vendor@example.com",
            is_active=True,
        )
        db.add_all([existing, foreign])
        db.commit()
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="vendor-resolve",
            ),
            db=db,
            current_user=admin,
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "vendor-directory.csv",
                    (
                        "Vendor ID,Company Name,Address,Phone Numbers,Email,Send 1099?\n"
                        "VENDOR-R1,Resolve Plumbing,10 Trade St,216-555-0200,"
                        "resolve.vendor@example.com,Yes\n"
                    ).encode(),
                ),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.disposition == "POSSIBLE_MATCH"
        before = api._staged_review_fingerprint(upload, [row])
        original_name = existing.company_name

        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_vendor(
                run.id,
                upload.id,
                row.id,
                AppFolioStagedVendorResolutionIn(
                    action="MATCH_EXISTING",
                    target_vendor_id=foreign.id,
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 404

        resolved = api.resolve_staged_appfolio_vendor(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedVendorResolutionIn(
                action="MATCH_EXISTING",
                target_vendor_id=existing.id,
            ),
            db=db,
            current_user=admin,
        )
        assert resolved.resolution_target_vendor_id == existing.id
        assert resolved.resolution_target_owner_user_id is None
        assert db.get(Vendor, existing.id).company_name == original_name
        assert api._staged_review_fingerprint(upload, [resolved]) != before

        changed = api.resolve_staged_appfolio_vendor(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedVendorResolutionIn(action="CREATE_NEW"),
            db=db,
            current_user=admin,
        )
        assert changed.resolution_action == "CREATE_NEW"
        assert changed.resolution_target_vendor_id is None
        assert db.get(Vendor, existing.id).company_name == original_name
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == row.id,
            AuditLog.action == "appfolio_vendor_resolution_changed",
        ).count() == 2
    finally:
        db.close()
        engine.dispose()


def test_appfolio_vendor_staged_dry_run_commit_match_create_replay_and_staging_only_fields():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Vendor Commit Org")
        existing = Vendor(
            organization_id=org.id,
            company_name="Existing Electric",
            business_email="existing.electric@example.com",
            phone="do-not-overwrite",
            is_active=True,
        )
        db.add(existing)
        db.commit()
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="vendor-commit",
            ),
            db=db,
            current_user=admin,
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "vendor-directory.csv",
                    (
                        "Vendor ID,Company Name,Address,Phone Numbers,Email,Send 1099?,"
                        "Liability Insurance Expiration,Workers Comp Expiration\n"
                        "VENDOR-MATCH,Existing Electric,20 Old St,216-555-0300,"
                        "existing.electric@example.com,Yes,2027-01-31,2027-02-28\n"
                        "VENDOR-NEW,New Roofing,30 New St,216-555-0400,"
                        "new.roofing@example.com,No,2028-01-31,2028-02-28\n"
                    ).encode(),
                ),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()
        assert [row.disposition for row in rows] == ["POSSIBLE_MATCH", "NEW"]

        api.resolve_staged_appfolio_vendor(
            run.id,
            upload.id,
            rows[0].id,
            AppFolioStagedVendorResolutionIn(
                action="MATCH_EXISTING",
                target_vendor_id=existing.id,
            ),
            db=db,
            current_user=admin,
        )

        before_count = db.query(Vendor).count()
        preview = api.dry_run_staged_appfolio_vendors(
            run.id,
            upload.id,
            db=db,
            current_user=admin,
        )
        assert preview.total == 2
        assert preview.importable == 2
        assert preview.invalid == 0
        assert db.query(Vendor).count() == before_count
        assert any(
            "target fields will not be overwritten" in warning
            for warning in preview.rows[0].warnings
        )
        assert any(
            "Staging-only source values" in warning
            for warning in preview.rows[1].warnings
        )

        committed = api.commit_staged_appfolio_vendors(
            run.id,
            upload.id,
            AppFolioStagedVendorCommitIn(fingerprint=preview.fingerprint),
            db=db,
            current_user=admin,
        )
        assert committed.committed == 1
        assert committed.matched_existing == 1
        assert db.query(Vendor).count() == before_count + 1
        assert db.get(Vendor, existing.id).phone == "do-not-overwrite"

        new_vendor = db.query(Vendor).filter(
            Vendor.organization_id == org.id,
            Vendor.company_name == "New Roofing",
        ).one()
        assert new_vendor.business_email == "new.roofing@example.com"
        assert new_vendor.phone is None
        assert new_vendor.address_line1 is None
        mappings = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "VENDORS",
        ).order_by(PlatformMigrationItem.source_id).all()
        assert [(m.source_id, m.target_entity) for m in mappings] == [
            ("VENDOR-MATCH", "VENDOR"),
            ("VENDOR-NEW", "VENDOR"),
        ]

        replay_preview = api.dry_run_staged_appfolio_vendors(
            run.id,
            upload.id,
            db=db,
            current_user=admin,
        )
        replay = api.commit_staged_appfolio_vendors(
            run.id,
            upload.id,
            AppFolioStagedVendorCommitIn(fingerprint=replay_preview.fingerprint),
            db=db,
            current_user=admin,
        )
        assert replay.replayed is True
        assert replay.committed == 0
        assert replay.matched_existing == 0
        assert db.query(Vendor).count() == before_count + 1
    finally:
        db.close()
        engine.dispose()


def test_appfolio_vendor_commit_rejects_stale_fingerprint_after_resolution_change():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Vendor Stale Org")
        existing = Vendor(
            organization_id=org.id,
            company_name="Stale HVAC",
            business_email="stale.hvac@example.com",
            is_active=True,
        )
        db.add(existing)
        db.commit()
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="vendor-stale",
            ),
            db=db,
            current_user=admin,
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "vendor-directory.csv",
                    (
                        "Vendor ID,Company Name,Email,Send 1099?\n"
                        "VENDOR-S,Stale HVAC,stale.hvac@example.com,Yes\n"
                    ).encode(),
                ),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        api.resolve_staged_appfolio_vendor(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedVendorResolutionIn(action="CREATE_NEW"),
            db=db,
            current_user=admin,
        )
        preview = api.dry_run_staged_appfolio_vendors(
            run.id, upload.id, db=db, current_user=admin
        )
        api.resolve_staged_appfolio_vendor(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedVendorResolutionIn(
                action="MATCH_EXISTING",
                target_vendor_id=existing.id,
            ),
            db=db,
            current_user=admin,
        )
        with pytest.raises(HTTPException) as exc:
            api.commit_staged_appfolio_vendors(
                run.id,
                upload.id,
                AppFolioStagedVendorCommitIn(fingerprint=preview.fingerprint),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "VENDORS",
        ).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_owner_staged_dry_run_commit_maps_existing_without_user_or_ownership_mutation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Owner Commit Org")
        existing = User(
            organization_id=org.id,
            role=UserRole.OWNER,
            email="owner.commit@example.com",
            first_name="Existing",
            last_name="Owner",
            hashed_password="x",
            is_active=True,
        )
        db.add(existing)
        db.commit()
        original_name = (existing.first_name, existing.last_name, existing.email)
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="owner-commit",
            ),
            db=db,
            current_user=admin,
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "owner-directory.csv",
                    (
                        "Owner ID,Name,Phone Numbers,Email,Properties Owned,Properties Owned IDs\n"
                        "OWNER-100,Source Display Name,216-555-0100,"
                        "owner.commit@example.com,Lake House,PROP-100\n"
                    ).encode(),
                ),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.disposition == "POSSIBLE_MATCH"

        api.resolve_staged_appfolio_owner(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedOwnerResolutionIn(
                action="MATCH_EXISTING",
                target_owner_user_id=existing.id,
            ),
            db=db,
            current_user=admin,
        )

        before_users = db.query(User).count()
        before_property_owners = db.query(PropertyOwner).count()
        preview = api.dry_run_staged_appfolio_owners(
            run.id,
            upload.id,
            db=db,
            current_user=admin,
        )
        assert preview.total == 1
        assert preview.importable == 1
        assert preview.invalid == 0
        assert db.query(User).count() == before_users
        assert db.query(PropertyOwner).count() == before_property_owners
        assert any(
            "target profile fields will not be overwritten" in warning
            for warning in preview.rows[0].warnings
        )
        assert any(
            "property ownership" in warning
            for warning in preview.rows[0].warnings
        )

        committed = api.commit_staged_appfolio_owners(
            run.id,
            upload.id,
            AppFolioStagedOwnerCommitIn(fingerprint=preview.fingerprint),
            db=db,
            current_user=admin,
        )
        assert committed.mapped_existing == 1
        assert committed.replayed is False
        assert db.query(User).count() == before_users
        assert db.query(PropertyOwner).count() == before_property_owners
        db.refresh(existing)
        assert (existing.first_name, existing.last_name, existing.email) == original_name

        mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "OWNERS",
            PlatformMigrationItem.source_id == "OWNER-100",
        ).one()
        assert mapping.target_entity == "OWNER_USER"
        assert mapping.target_id == existing.id

        replay_preview = api.dry_run_staged_appfolio_owners(
            run.id,
            upload.id,
            db=db,
            current_user=admin,
        )
        replay = api.commit_staged_appfolio_owners(
            run.id,
            upload.id,
            AppFolioStagedOwnerCommitIn(fingerprint=replay_preview.fingerprint),
            db=db,
            current_user=admin,
        )
        assert replay.replayed is True
        assert replay.mapped_existing == 0
        assert db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "OWNERS",
        ).count() == 1
        assert db.query(User).count() == before_users
    finally:
        db.close()
        engine.dispose()


def test_appfolio_owner_controlled_commit_blocks_create_new_and_stale_resolution():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Owner Commit Guard Org")
        first = User(
            organization_id=org.id,
            role=UserRole.OWNER,
            email="owner.guard@example.com",
            first_name="First",
            last_name="Owner",
            hashed_password="x",
            is_active=True,
        )
        second = User(
            organization_id=org.id,
            role=UserRole.OWNER,
            email="second.owner.guard@example.com",
            first_name="Second",
            last_name="Owner",
            hashed_password="x",
            is_active=True,
        )
        db.add_all([first, second])
        db.commit()
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="owner-commit-guard",
            ),
            db=db,
            current_user=admin,
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "owner-directory.csv",
                    (
                        "Owner ID,Name,Phone Numbers,Email,Properties Owned,Properties Owned IDs\n"
                        "OWNER-GUARD,Guard Owner,216-555-0101,"
                        "owner.guard@example.com,Lake House,PROP-100\n"
                    ).encode(),
                ),
                resource="OWNERS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()

        api.resolve_staged_appfolio_owner(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedOwnerResolutionIn(action="CREATE_NEW"),
            db=db,
            current_user=admin,
        )
        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_owners(
                run.id,
                upload.id,
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert "CREATE_NEW is intentionally unsupported" in str(exc.value.detail)
        assert db.query(User).count() == 2

        api.resolve_staged_appfolio_owner(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedOwnerResolutionIn(
                action="MATCH_EXISTING",
                target_owner_user_id=first.id,
            ),
            db=db,
            current_user=admin,
        )
        preview = api.dry_run_staged_appfolio_owners(
            run.id,
            upload.id,
            db=db,
            current_user=admin,
        )
        api.resolve_staged_appfolio_owner(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedOwnerResolutionIn(
                action="MATCH_EXISTING",
                target_owner_user_id=second.id,
            ),
            db=db,
            current_user=admin,
        )
        with pytest.raises(HTTPException) as exc:
            api.commit_staged_appfolio_owners(
                run.id,
                upload.id,
                AppFolioStagedOwnerCommitIn(fingerprint=preview.fingerprint),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "OWNERS",
        ).count() == 0
        assert db.query(User).count() == 2
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Tenant Directory ingestion + staging
# ---------------------------------------------------------------------------

def _seed_tenant_source_relationship_mappings(db, *, run, org):
    prop = Property(
        organization_id=org.id,
        name="Tenant Stage Property",
        address_line1="100 Tenant Ave",
        city="Cleveland",
        state="OH",
        zip_code="44113",
        is_active=True,
    )
    db.add(prop)
    db.flush()
    unit = Unit(
        property_id=prop.id,
        unit_number="101",
        bedrooms=1,
        bathrooms=1,
        monthly_rent=1000,
        is_active=True,
    )
    db.add(unit)
    db.flush()
    db.add_all([
        PlatformMigrationItem(
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            resource="PROPERTIES",
            source_id="PROP-TENANT-1",
            target_entity="PROPERTY",
            target_id=prop.id,
            source_fingerprint="1" * 64,
        ),
        PlatformMigrationItem(
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            resource="UNITS",
            source_id="UNIT-TENANT-1",
            target_entity="UNIT",
            target_id=unit.id,
            source_fingerprint="2" * 64,
        ),
    ])
    db.commit()
    return prop, unit


def test_appfolio_tenant_directory_csv_stages_verified_fields_relationship_context_and_replays_without_customer_mutation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Tenant Stage Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="tenant-stage",
            ),
            db=db,
            current_user=admin,
        )
        _seed_tenant_source_relationship_mappings(db, run=run, org=org)

        before_users = db.query(User).count()
        before_leases = db.query(Lease).count()
        before_charges = db.query(Charge).count()
        before_gl = db.query(GLTransaction).count()
        content = (
            "Tenant ID,Tenant,Phone Numbers,Emails,Tenant Street Address 1,"
            "Tenant Street Address 2,Tenant City,Tenant State,Tenant Zip,"
            "Property Name,Property ID,Property Address,Unit,Unit ID,"
            "Move-in,Move-out,Lease From,Lease To\n"
            "TENANT-1,Jane Tenant,216-555-0200,jane.tenant.source@example.com,"
            "5 Tenant St,,Cleveland,OH,44113,Tenant Stage Property,PROP-TENANT-1,"
            "100 Tenant Ave,101,UNIT-TENANT-1,2026-01-01,,2026-01-01,2026-12-31\n"
        ).encode()

        first = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("tenant-directory.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert first.detected_resource == "TENANTS"
        assert first.status == "REVIEW_REQUIRED"
        assert first.validation_summary["valid"] == 1
        assert first.validation_summary["invalid"] == 0
        assert first.validation_summary["missing_tenant_source_ids"] == 0
        assert first.validation_summary["missing_tenant_unit_ids"] == 0
        assert first.validation_summary["missing_tenant_property_ids"] == 0

        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == first.id
        ).one()
        assert row.resource == "TENANTS"
        assert row.source_id == "TENANT-1"
        assert row.disposition == "NEW"
        assert row.normalized_data["tenant_name"] == "Jane Tenant"
        assert row.normalized_data["source_property_id"] == "PROP-TENANT-1"
        assert row.normalized_data["source_unit_id"] == "UNIT-TENANT-1"
        assert row.normalized_data["move_in"] == "2026-01-01"
        assert row.normalized_data["lease_to"] == "2026-12-31"
        assert any("does not establish occupancy" in warning for warning in row.warnings)
        assert any("no occupancy event is created" in warning for warning in row.warnings)
        assert any("no Lease record is created" in warning for warning in row.warnings)

        assert db.query(User).count() == before_users
        assert db.query(Lease).count() == before_leases
        assert db.query(Charge).count() == before_charges
        assert db.query(GLTransaction).count() == before_gl

        replay = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("tenant-directory-renamed.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert replay.replayed is True
        assert replay.id == first.id
        assert db.query(PlatformMigrationUpload).filter(
            PlatformMigrationUpload.run_id == run.id,
            PlatformMigrationUpload.detected_resource == "TENANTS",
        ).count() == 1
        assert db.query(User).count() == before_users
    finally:
        db.close()
        engine.dispose()


def test_appfolio_tenant_directory_missing_identity_relationships_and_exact_single_email_are_reviewed_not_overwritten():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Tenant Review Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="tenant-review",
            ),
            db=db,
            current_user=admin,
        )
        _seed_tenant_source_relationship_mappings(db, run=run, org=org)
        existing = User(
            organization_id=org.id,
            role=UserRole.TENANT,
            email="existing.tenant@example.com",
            first_name="Existing",
            last_name="Tenant",
            hashed_password="x",
            is_active=True,
        )
        db.add(existing)
        db.commit()
        original_name = (existing.first_name, existing.last_name, existing.email)

        content = (
            "Tenant ID,Tenant,Phone Numbers,Emails,Property Name,Property ID,Unit,Unit ID\n"
            "TENANT-MATCH,Existing Source Name,216-555-0201,existing.tenant@example.com,"
            "Tenant Stage Property,PROP-TENANT-1,101,UNIT-TENANT-1\n"
            ",Missing Stable ID,216-555-0202,missing.id@example.com,"
            "Tenant Stage Property,PROP-TENANT-1,101,UNIT-TENANT-1\n"
            "TENANT-NOREF,No Refs,216-555-0203,norefs@example.com,Tenant Stage Property,,101,\n"
            "TENANT-MULTIEMAIL,Multi Email,216-555-0204,"
            "existing.tenant@example.com;other@example.com,"
            "Tenant Stage Property,PROP-TENANT-1,101,UNIT-TENANT-1\n"
        ).encode()

        before_users = db.query(User).count()
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("tenant-review.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()

        assert upload.detected_resource == "TENANTS"
        assert upload.status == "REVIEW_REQUIRED"
        assert [row.disposition for row in rows] == [
            "POSSIBLE_MATCH", "REVIEW", "REVIEW", "NEW"
        ]
        assert upload.validation_summary["possible_existing_matches"] == 1
        assert upload.validation_summary["missing_tenant_source_ids"] == 1
        assert upload.validation_summary["missing_tenant_unit_ids"] == 1
        assert upload.validation_summary["missing_tenant_property_ids"] == 1
        assert any(
            "exact source email" in warning for warning in rows[0].warnings
        )
        assert not any(
            "Possible existing target tenant match" in warning
            for warning in rows[3].warnings
        )
        db.refresh(existing)
        assert (existing.first_name, existing.last_name, existing.email) == original_name
        assert db.query(User).count() == before_users
        assert db.query(Lease).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_tenant_directory_duplicate_id_and_contradictory_mapped_unit_property_fail_closed():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Tenant Guard Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="tenant-guard",
            ),
            db=db,
            current_user=admin,
        )
        prop, unit = _seed_tenant_source_relationship_mappings(db, run=run, org=org)
        other_prop = Property(
            organization_id=org.id,
            name="Other Tenant Property",
            address_line1="200 Other Ave",
            city="Cleveland",
            state="OH",
            zip_code="44113",
            is_active=True,
        )
        db.add(other_prop)
        db.flush()
        db.add(
            PlatformMigrationItem(
                run_id=run.id,
                organization_id=org.id,
                provider="APPFOLIO",
                resource="PROPERTIES",
                source_id="PROP-TENANT-2",
                target_entity="PROPERTY",
                target_id=other_prop.id,
                source_fingerprint="3" * 64,
            )
        )
        db.commit()

        content = (
            "Tenant ID,Tenant,Emails,Property ID,Unit ID\n"
            "TENANT-DUP,One Tenant,one@example.com,PROP-TENANT-1,UNIT-TENANT-1\n"
            "TENANT-DUP,Duplicate Tenant,two@example.com,PROP-TENANT-1,UNIT-TENANT-1\n"
            "TENANT-MISMATCH,Mismatch Tenant,three@example.com,PROP-TENANT-2,UNIT-TENANT-1\n"
        ).encode()
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("tenant-guard.csv", content),
                resource="TENANTS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()
        assert rows[0].disposition == "NEW"
        assert rows[1].disposition == "INVALID"
        assert any("Duplicate AppFolio Tenant ID" in error for error in rows[1].errors)
        assert rows[2].disposition == "INVALID"
        assert any(
            "resolve to different target Properties" in error
            for error in rows[2].errors
        )
        assert upload.validation_summary["duplicates"] == 1
        assert upload.validation_summary["invalid"] == 2
        assert db.query(User).filter(User.role == UserRole.TENANT).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_tenant_directory_xlsx_auto_detection_and_replay():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Tenant XLSX Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="tenant-xlsx",
            ),
            db=db,
            current_user=admin,
        )
        _seed_tenant_source_relationship_mappings(db, run=run, org=org)

        workbook = Workbook()
        ws = workbook.active
        ws.title = "Instructions"
        ws.append(["Read me"])
        tenant_ws = workbook.create_sheet("Tenant Directory")
        tenant_ws.append([
            "Tenant ID", "Tenant", "Emails", "Property ID", "Unit ID",
            "Move-in", "Lease From", "Lease To"
        ])
        tenant_ws.append([
            "TENANT-XLSX-1", "XLSX Tenant", "xlsx.tenant@example.com",
            "PROP-TENANT-1", "UNIT-TENANT-1",
            "2026-02-01", "2026-02-01", "2027-01-31"
        ])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()
        content = buffer.getvalue()

        first = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("tenant-directory.xlsx", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert first.detected_resource == "TENANTS"
        assert first.sheet_name == "Tenant Directory"
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == first.id
        ).one()
        assert row.source_id == "TENANT-XLSX-1"
        assert row.disposition == "NEW"

        replay = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("tenant-directory-copy.xlsx", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert replay.replayed is True
        assert replay.id == first.id
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Tenant staged reconciliation
# ---------------------------------------------------------------------------

def test_appfolio_tenant_possible_match_resolution_is_typed_scoped_and_non_mutating():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Tenant Resolve Org")
        foreign_org = _org(db, name="Tenant Resolve Foreign")
        existing = User(
            organization_id=org.id,
            role=UserRole.TENANT,
            email="tenant.resolve@example.com",
            first_name="Existing",
            last_name="Tenant",
            hashed_password="x",
            is_active=True,
        )
        foreign = User(
            organization_id=foreign_org.id,
            role=UserRole.TENANT,
            email="foreign.tenant.resolve@example.com",
            first_name="Foreign",
            last_name="Tenant",
            hashed_password="x",
            is_active=True,
        )
        db.add_all([existing, foreign])
        db.commit()
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="tenant-resolve",
            ),
            db=db,
            current_user=admin,
        )
        _seed_tenant_source_relationship_mappings(db, run=run, org=org)
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "tenant-directory.csv",
                    (
                        "Tenant ID,Tenant Name,Email,Property ID,Property Name,Unit ID,Unit\n"
                        "TENANT-R1,Source Tenant,tenant.resolve@example.com,PROP-TENANT-1,"
                        "Tenant Stage Property,UNIT-TENANT-1,101\n"
                    ).encode(),
                ),
                resource="TENANTS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.disposition == "POSSIBLE_MATCH"
        before = api._staged_review_fingerprint(upload, [row])
        original = (existing.first_name, existing.last_name, existing.email)

        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_tenant(
                run.id,
                upload.id,
                row.id,
                AppFolioStagedTenantResolutionIn(
                    action="MATCH_EXISTING",
                    target_tenant_user_id=foreign.id,
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 404

        run.last_dry_run_fingerprint = "a" * 64
        run.last_dry_run_summary = {"old": True}
        db.commit()
        resolved = api.resolve_staged_appfolio_tenant(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedTenantResolutionIn(
                action="MATCH_EXISTING",
                target_tenant_user_id=existing.id,
            ),
            db=db,
            current_user=admin,
        )
        assert resolved.resolution_action == "MATCH_EXISTING"
        assert resolved.resolution_target_tenant_user_id == existing.id
        assert resolved.resolution_target_id is None
        assert resolved.resolution_target_unit_id is None
        assert resolved.resolution_target_owner_user_id is None
        assert resolved.resolution_target_vendor_id is None
        assert (existing.first_name, existing.last_name, existing.email) == original
        db.refresh(run)
        assert run.last_dry_run_fingerprint is None
        assert run.last_dry_run_summary is None
        assert api._staged_review_fingerprint(upload, [resolved]) != before
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == row.id,
            AuditLog.action == "appfolio_tenant_resolution_changed",
        ).count() == 1

        changed = api.resolve_staged_appfolio_tenant(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedTenantResolutionIn(action="CREATE_NEW"),
            db=db,
            current_user=admin,
        )
        assert changed.resolution_action == "CREATE_NEW"
        assert changed.resolution_target_tenant_user_id is None
        assert db.query(User).filter(User.organization_id == org.id, User.role == UserRole.TENANT).count() == 1
    finally:
        db.close()
        engine.dispose()


def test_appfolio_tenant_review_rows_may_only_skip_and_create_no_customer_records():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Tenant Review Resolve")
        existing = User(
            organization_id=org.id,
            role=UserRole.TENANT,
            email="tenant.review@example.com",
            first_name="Review",
            last_name="Tenant",
            hashed_password="x",
            is_active=True,
        )
        db.add(existing)
        db.commit()
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="tenant-review-resolve",
            ),
            db=db,
            current_user=admin,
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "tenant-review.csv",
                    (
                        "Tenant ID,Tenant Name,Email,Property ID,Property Name,Unit ID,Unit\n"
                        "TENANT-REVIEW,Needs Relationship,tenant.review@example.com,PROP-MISSING,"
                        "Unknown Property,UNIT-MISSING,999\n"
                    ).encode(),
                ),
                resource="TENANTS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.disposition == "REVIEW"

        for payload in (
            AppFolioStagedTenantResolutionIn(
                action="MATCH_EXISTING",
                target_tenant_user_id=existing.id,
            ),
            AppFolioStagedTenantResolutionIn(action="CREATE_NEW"),
        ):
            with pytest.raises(HTTPException) as exc:
                api.resolve_staged_appfolio_tenant(
                    run.id,
                    upload.id,
                    row.id,
                    payload,
                    db=db,
                    current_user=admin,
                )
            assert exc.value.status_code == 409

        before_users = db.query(User).count()
        skipped = api.resolve_staged_appfolio_tenant(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedTenantResolutionIn(action="SKIP"),
            db=db,
            current_user=admin,
        )
        assert skipped.resolution_action == "SKIP"
        assert skipped.resolution_target_tenant_user_id is None
        assert db.query(User).count() == before_users
        assert db.query(Lease).count() == 0
        assert db.query(Charge).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Tenant staged dry run + controlled mapping commit
# ---------------------------------------------------------------------------

def test_staged_tenant_dry_run_and_controlled_mapping_commit_is_replay_safe_and_non_mutating():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Tenant Commit Org")
        existing = User(
            organization_id=org.id,
            role=UserRole.TENANT,
            email="tenant.commit@example.com",
            first_name="Existing",
            last_name="Tenant",
            hashed_password="x",
            is_active=True,
        )
        db.add(existing)
        db.commit()
        original = (existing.first_name, existing.last_name, existing.email)
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="tenant-controlled-commit",
            ),
            db=db,
            current_user=admin,
        )
        _seed_tenant_source_relationship_mappings(db, run=run, org=org)
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "tenant-commit.csv",
                    (
                        "Tenant ID,Tenant Name,Email,Property ID,Property Name,Unit ID,Unit,"
                        "Move-in,Lease From,Lease To\n"
                        "TENANT-C1,Source Display Name,tenant.commit@example.com,PROP-TENANT-1,"
                        "Tenant Stage Property,UNIT-TENANT-1,101,2026-01-01,2026-01-01,2026-12-31\n"
                    ).encode(),
                ),
                resource="TENANTS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.disposition == "POSSIBLE_MATCH"
        api.resolve_staged_appfolio_tenant(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedTenantResolutionIn(
                action="MATCH_EXISTING",
                target_tenant_user_id=existing.id,
            ),
            db=db,
            current_user=admin,
        )

        before_users = db.query(User).count()
        before_leases = db.query(Lease).count()
        before_charges = db.query(Charge).count()
        before_gl = db.query(GLTransaction).count()
        preview = api.dry_run_staged_appfolio_tenants(
            run.id, upload.id, db=db, current_user=admin
        )
        assert preview.total == 1
        assert preview.importable == 1
        assert preview.invalid == 0
        assert preview.rows[0].mapped == {"target_tenant_user_id": existing.id}
        assert db.query(User).count() == before_users

        committed = api.commit_staged_appfolio_tenants(
            run.id,
            upload.id,
            AppFolioStagedTenantCommitIn(fingerprint=preview.fingerprint),
            db=db,
            current_user=admin,
        )
        assert committed.mapped_existing == 1
        assert committed.replayed is False
        assert committed.rows[0].target_tenant_user_id == existing.id
        mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "TENANTS",
            PlatformMigrationItem.source_id == "TENANT-C1",
        ).one()
        assert mapping.target_entity == "TENANT_USER"
        assert mapping.target_id == existing.id
        db.refresh(existing)
        assert (existing.first_name, existing.last_name, existing.email) == original
        assert db.query(User).count() == before_users
        assert db.query(Lease).count() == before_leases
        assert db.query(Charge).count() == before_charges
        assert db.query(GLTransaction).count() == before_gl

        replay = api.commit_staged_appfolio_tenants(
            run.id,
            upload.id,
            AppFolioStagedTenantCommitIn(fingerprint=preview.fingerprint),
            db=db,
            current_user=admin,
        )
        assert replay.replayed is True
        assert replay.mapped_existing == 0
        assert db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "TENANTS",
            PlatformMigrationItem.source_id == "TENANT-C1",
        ).count() == 1
    finally:
        db.close()
        engine.dispose()


def test_staged_tenant_dry_run_rejects_create_new_and_relationship_mapping_change_stales_commit():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Tenant Fingerprint Org")
        existing = User(
            organization_id=org.id,
            role=UserRole.TENANT,
            email="tenant.fingerprint@example.com",
            first_name="Existing",
            last_name="Fingerprint",
            hashed_password="x",
            is_active=True,
        )
        db.add(existing)
        db.commit()
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="tenant-fingerprint",
            ),
            db=db,
            current_user=admin,
        )
        _seed_tenant_source_relationship_mappings(db, run=run, org=org)
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "tenant-fingerprint.csv",
                    (
                        "Tenant ID,Tenant Name,Email,Property ID,Property Name,Unit ID,Unit\n"
                        "TENANT-F1,Source Fingerprint,tenant.fingerprint@example.com,PROP-TENANT-1,"
                        "Tenant Stage Property,UNIT-TENANT-1,101\n"
                    ).encode(),
                ),
                resource="TENANTS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.disposition == "POSSIBLE_MATCH"

        api.resolve_staged_appfolio_tenant(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedTenantResolutionIn(action="CREATE_NEW"),
            db=db,
            current_user=admin,
        )
        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_tenants(
                run.id, upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "CREATE_NEW is intentionally unsupported" in exc.value.detail

        api.resolve_staged_appfolio_tenant(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedTenantResolutionIn(
                action="MATCH_EXISTING",
                target_tenant_user_id=existing.id,
            ),
            db=db,
            current_user=admin,
        )
        preview = api.dry_run_staged_appfolio_tenants(
            run.id, upload.id, db=db, current_user=admin
        )
        unit_mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "UNITS",
            PlatformMigrationItem.source_id == "UNIT-TENANT-1",
        ).one()
        unit_mapping.source_fingerprint = "9" * 64
        db.commit()

        with pytest.raises(HTTPException) as exc:
            api.commit_staged_appfolio_tenants(
                run.id,
                upload.id,
                AppFolioStagedTenantCommitIn(fingerprint=preview.fingerprint),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert "does not match the supplied Tenant dry-run fingerprint" in exc.value.detail
        assert db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "TENANTS",
        ).count() == 0
        assert db.query(Lease).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_staged_tenant_dry_run_and_commit_routes_are_exposed():
    from app.main import app
    paths = set(app.openapi()["paths"])
    assert (
        "/api/platform/migrations/appfolio/runs/{run_id}/uploads/{upload_id}/tenants/dry-run"
        in paths
    )
    assert (
        "/api/platform/migrations/appfolio/runs/{run_id}/uploads/{upload_id}/tenants/commit"
        in paths
    )


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Lease / occupancy staging foundation
# ---------------------------------------------------------------------------

def _seed_lease_occupancy_source_mappings(db, *, run, org):
    prop, unit = _seed_tenant_source_relationship_mappings(db, run=run, org=org)
    tenant = User(
        organization_id=org.id,
        role=UserRole.TENANT,
        email="lease.occupancy.tenant@example.com",
        first_name="Lease",
        last_name="Occupancy",
        hashed_password="x",
        is_active=True,
    )
    db.add(tenant)
    db.flush()
    db.add(
        PlatformMigrationItem(
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            resource="TENANTS",
            source_id="TENANT-LEASE-1",
            target_entity="TENANT_USER",
            target_id=tenant.id,
            source_fingerprint="4" * 64,
        )
    )
    db.commit()
    return prop, unit, tenant


def test_appfolio_lease_occupancy_explicit_csv_stages_relationship_evidence_without_customer_mutation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Lease Occupancy Stage Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="lease-occupancy-stage",
            ),
            db=db,
            current_user=admin,
        )
        _seed_lease_occupancy_source_mappings(db, run=run, org=org)

        before_users = db.query(User).count()
        before_leases = db.query(Lease).count()
        before_charges = db.query(Charge).count()
        before_gl = db.query(GLTransaction).count()
        content = (
            "Tenant ID,Tenant,Property ID,Property Name,Unit ID,Unit,Status,Rent,Deposit,"
            "Move-in,Move-out,Lease From,Lease To\n"
            "TENANT-LEASE-1,Lease Occupancy,PROP-TENANT-1,Tenant Stage Property,"
            "UNIT-TENANT-1,101,Current,1250.00,500.00,2026-01-01,,2026-01-01,2026-12-31\n"
        ).encode()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("lease-occupancy.csv", content),
                resource="LEASE_OCCUPANCY",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert upload.detected_resource == "LEASE_OCCUPANCY"
        assert upload.status == "REVIEW_REQUIRED"
        assert upload.validation_summary["lease_occupancy_rows"] == 1
        assert upload.validation_summary["relationships_ready_for_review"] == 1
        assert upload.validation_summary["unresolved_lease_occupancy_relationships"] == 0
        assert upload.validation_summary["customer_lease_mutation"] is False
        assert upload.validation_summary["accounting_mutation"] is False

        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.resource == "LEASE_OCCUPANCY"
        assert row.source_id is None
        assert row.disposition == "REVIEW"
        assert row.normalized_data["source_tenant_id"] == "TENANT-LEASE-1"
        assert row.normalized_data["source_property_id"] == "PROP-TENANT-1"
        assert row.normalized_data["source_unit_id"] == "UNIT-TENANT-1"
        assert row.normalized_data["status"] == "Current"
        assert row.normalized_data["rent"] == "1250.00"
        assert row.normalized_data["deposit"] == "500.00"
        assert row.normalized_data["lease_from"] == "2026-01-01"
        assert any("no Lease or occupancy record is created" in w for w in row.warnings)
        assert any("source evidence only" in w for w in row.warnings)

        assert db.query(User).count() == before_users
        assert db.query(Lease).count() == before_leases
        assert db.query(Charge).count() == before_charges
        assert db.query(GLTransaction).count() == before_gl

        replay = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("lease-occupancy-copy.csv", content),
                resource="LEASE_OCCUPANCY",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert replay.replayed is True
        assert replay.id == upload.id
    finally:
        db.close()
        engine.dispose()


def test_appfolio_tenant_directory_still_auto_detects_tenants_while_lease_occupancy_is_explicit_only():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Lease Detection Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="lease-detection",
            ),
            db=db,
            current_user=admin,
        )
        _seed_lease_occupancy_source_mappings(db, run=run, org=org)
        content = (
            "Tenant ID,Tenant,Property ID,Unit ID,Lease From,Lease To\n"
            "TENANT-LEASE-1,Lease Occupancy,PROP-TENANT-1,UNIT-TENANT-1,"
            "2026-01-01,2026-12-31\n"
        ).encode()

        automatic = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("tenant-directory.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert automatic.detected_resource == "TENANTS"

        explicit = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("lease-evidence.csv", content),
                resource="LEASE_OCCUPANCY",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert explicit.detected_resource == "LEASE_OCCUPANCY"
        assert explicit.id != automatic.id
    finally:
        db.close()
        engine.dispose()


def test_appfolio_lease_occupancy_missing_or_contradictory_relationships_fail_closed_without_lease_creation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Lease Relationship Guard Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="lease-relationship-guard",
            ),
            db=db,
            current_user=admin,
        )
        _, unit, _ = _seed_lease_occupancy_source_mappings(db, run=run, org=org)

        unresolved_content = (
            "Tenant ID,Tenant,Property ID,Unit ID,Lease From,Lease To\n"
            "TENANT-MISSING,Missing Mapping,PROP-MISSING,UNIT-MISSING,2026-01-01,2026-12-31\n"
        ).encode()
        unresolved = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("lease-unresolved.csv", unresolved_content),
                resource="LEASE_OCCUPANCY",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        unresolved_row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == unresolved.id
        ).one()
        assert unresolved_row.disposition == "REVIEW"
        assert unresolved.validation_summary["unresolved_lease_occupancy_relationships"] == 1

        other_prop = Property(
            organization_id=org.id,
            name="Lease Other Property",
            address_line1="200 Lease Other St",
            city="Cleveland",
            state="OH",
            zip_code="44113",
            is_active=True,
        )
        db.add(other_prop)
        db.flush()
        property_mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "PROPERTIES",
            PlatformMigrationItem.source_id == "PROP-TENANT-1",
        ).one()
        property_mapping.target_id = other_prop.id
        db.commit()

        contradictory_content = (
            "Tenant ID,Tenant,Property ID,Unit ID,Lease From,Lease To\n"
            "TENANT-LEASE-1,Lease Occupancy,PROP-TENANT-1,UNIT-TENANT-1,"
            "2026-01-01,2026-12-31\n"
        ).encode()
        contradictory = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("lease-contradictory.csv", contradictory_content),
                resource="LEASE_OCCUPANCY",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        contradictory_row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == contradictory.id
        ).one()
        assert contradictory_row.disposition == "INVALID"
        assert any(
            "resolve to different target Properties" in error
            for error in contradictory_row.errors
        )
        assert db.query(Lease).count() == 0
        assert unit.property_id != other_prop.id
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Lease / occupancy staged reconciliation
# ---------------------------------------------------------------------------

def test_appfolio_lease_occupancy_reconciliation_accepts_durable_relationship_without_customer_mutation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Lease Occupancy Resolve Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="lease-occupancy-resolve",
            ),
            db=db,
            current_user=admin,
        )
        prop, unit, tenant = _seed_lease_occupancy_source_mappings(
            db, run=run, org=org
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "lease-occupancy-resolve.csv",
                    (
                        "Tenant ID,Tenant,Property ID,Unit ID,Lease From,Lease To,Rent,Deposit\n"
                        "TENANT-LEASE-1,Lease Occupancy,PROP-TENANT-1,UNIT-TENANT-1,"
                        "2026-01-01,2026-12-31,1250.00,500.00\n"
                    ).encode(),
                ),
                resource="LEASE_OCCUPANCY",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.disposition == "REVIEW"
        before_fingerprint = api._staged_review_fingerprint(upload, [row])
        before_leases = db.query(Lease).count()
        before_charges = db.query(Charge).count()
        before_gl = db.query(GLTransaction).count()

        run.last_dry_run_fingerprint = "a" * 64
        run.last_dry_run_summary = {"old": True}
        db.commit()

        resolved = api.resolve_staged_appfolio_lease_occupancy(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedLeaseOccupancyResolutionIn(
                action="ACCEPT_RELATIONSHIP"
            ),
            db=db,
            current_user=admin,
        )
        assert resolved.resolution_action == "ACCEPT_RELATIONSHIP"
        assert resolved.resolution_target_id == prop.id
        assert resolved.resolution_target_unit_id == unit.id
        assert resolved.resolution_target_tenant_user_id == tenant.id
        assert resolved.resolution_target_owner_user_id is None
        assert resolved.resolution_target_vendor_id is None
        assert api._staged_review_fingerprint(upload, [resolved]) != before_fingerprint

        db.refresh(run)
        assert run.last_dry_run_fingerprint is None
        assert run.last_dry_run_summary is None
        assert run.status == "STAGED"
        assert db.query(Lease).count() == before_leases
        assert db.query(Charge).count() == before_charges
        assert db.query(GLTransaction).count() == before_gl
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == row.id,
            AuditLog.action == "appfolio_lease_occupancy_resolution_changed",
        ).count() == 1

        replay = api.resolve_staged_appfolio_lease_occupancy(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedLeaseOccupancyResolutionIn(
                action="ACCEPT_RELATIONSHIP"
            ),
            db=db,
            current_user=admin,
        )
        assert replay.id == row.id
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == row.id,
            AuditLog.action == "appfolio_lease_occupancy_resolution_changed",
        ).count() == 1
    finally:
        db.close()
        engine.dispose()


def test_appfolio_lease_occupancy_unresolved_relationship_can_only_be_skipped():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Lease Occupancy Skip Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="lease-occupancy-skip",
            ),
            db=db,
            current_user=admin,
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "lease-occupancy-unresolved.csv",
                    (
                        "Tenant ID,Tenant,Property ID,Unit ID,Lease From,Lease To\n"
                        "TENANT-NO-MAP,Unresolved Tenant,PROP-NO-MAP,UNIT-NO-MAP,"
                        "2026-01-01,2026-12-31\n"
                    ).encode(),
                ),
                resource="LEASE_OCCUPANCY",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.disposition == "REVIEW"

        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_lease_occupancy(
                run.id,
                upload.id,
                row.id,
                AppFolioStagedLeaseOccupancyResolutionIn(
                    action="ACCEPT_RELATIONSHIP"
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409

        skipped = api.resolve_staged_appfolio_lease_occupancy(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedLeaseOccupancyResolutionIn(action="SKIP"),
            db=db,
            current_user=admin,
        )
        assert skipped.resolution_action == "SKIP"
        assert skipped.resolution_target_id is None
        assert skipped.resolution_target_unit_id is None
        assert skipped.resolution_target_tenant_user_id is None
        assert db.query(Lease).count() == 0
        assert db.query(Charge).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_lease_occupancy_reconciliation_revalidates_mapping_dependencies_and_rejects_invalid_rows():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Lease Occupancy Revalidate Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="lease-occupancy-revalidate",
            ),
            db=db,
            current_user=admin,
        )
        _seed_lease_occupancy_source_mappings(db, run=run, org=org)
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "lease-occupancy-revalidate.csv",
                    (
                        "Tenant ID,Tenant,Property ID,Unit ID,Lease From,Lease To\n"
                        "TENANT-LEASE-1,Lease Occupancy,PROP-TENANT-1,UNIT-TENANT-1,"
                        "2026-01-01,2026-12-31\n"
                    ).encode(),
                ),
                resource="LEASE_OCCUPANCY",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()

        tenant_mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "TENANTS",
            PlatformMigrationItem.source_id == "TENANT-LEASE-1",
        ).one()
        tenant_mapping.target_entity = "PROPERTY"
        db.commit()

        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_lease_occupancy(
                run.id,
                upload.id,
                row.id,
                AppFolioStagedLeaseOccupancyResolutionIn(
                    action="ACCEPT_RELATIONSHIP"
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409

        row.errors = ["contradictory source relationship"]
        row.disposition = "INVALID"
        tenant_mapping.target_entity = "TENANT_USER"
        db.commit()
        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_lease_occupancy(
                run.id,
                upload.id,
                row.id,
                AppFolioStagedLeaseOccupancyResolutionIn(action="SKIP"),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert db.query(Lease).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio GL Accounts source-schema verification + staging
# ---------------------------------------------------------------------------

def test_appfolio_gl_accounts_csv_stages_verified_source_fields_without_accounting_mutation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="GL Account Stage Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="gl-account-stage",
            ),
            db=db,
            current_user=admin,
        )
        before_accounts = db.query(GLAccount).count()
        before_transactions = db.query(GLTransaction).count()

        content = (
            "GL Account ID,Number,Name,Type,FundAccount,IsCorporateAccount,"
            "OffsetAccountId,ParentGlAccountId,PropertyIds,LastUpdatedAt\n"
            "GL-4100,4100,Rent Income,Income,Operating,false,,,PROP-1,2026-09-30T12:00:00Z\n"
        ).encode()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("chart-of-accounts.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert upload.detected_resource == "GL_ACCOUNTS"
        assert upload.status == "REVIEW_REQUIRED"
        assert upload.validation_summary["gl_account_rows"] == 1
        assert upload.validation_summary["missing_gl_account_source_ids"] == 0
        assert upload.validation_summary["gl_account_target_mutation"] is False
        assert upload.validation_summary["accounting_history_mutation"] is False

        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.resource == "GL_ACCOUNTS"
        assert row.source_id == "GL-4100"
        assert row.disposition == "REVIEW"
        assert row.normalized_data == {
            "source_id": "GL-4100",
            "account_number": "4100",
            "account_name": "Rent Income",
            "account_type": "Income",
            "fund_account": "Operating",
            "is_corporate_account": "false",
            "offset_account_id": None,
            "parent_gl_account_id": None,
            "property_ids": "PROP-1",
            "last_updated_at": "2026-09-30T12:00:00Z",
        }
        assert any(
            "preserved as source evidence only" in warning
            for warning in row.warnings
        )
        assert any(
            "creates or updates no GLAccount" in warning
            for warning in row.warnings
        )
        assert db.query(GLAccount).count() == before_accounts
        assert db.query(GLTransaction).count() == before_transactions
    finally:
        db.close()
        engine.dispose()


def test_appfolio_gl_accounts_missing_source_id_remains_review_and_duplicate_source_id_is_invalid():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="GL Account Review Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="gl-account-review",
            ),
            db=db,
            current_user=admin,
        )

        no_id = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "chart-no-id.csv",
                    (
                        "Number,Name,Type,FundAccount\n"
                        "1000,Operating Cash,Cash,Operating\n"
                    ).encode(),
                ),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == no_id.id
        ).one()
        assert no_id.detected_resource == "GL_ACCOUNTS"
        assert row.source_id is None
        assert row.disposition == "REVIEW"
        assert no_id.validation_summary["missing_gl_account_source_ids"] == 1
        assert any(
            "not promoted to durable source identity automatically" in warning
            for warning in row.warnings
        )

        duplicate = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "chart-duplicate.csv",
                    (
                        "GL Account ID,Number,Name,Type\n"
                        "GL-1,1000,Operating Cash,Cash\n"
                        "GL-1,1100,Savings Cash,Cash\n"
                    ).encode(),
                ),
                resource="GL_ACCOUNTS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == duplicate.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()
        assert rows[0].disposition == "REVIEW"
        assert rows[1].disposition == "INVALID"
        assert any(
            "Duplicate AppFolio GL Account ID" in error
            for error in rows[1].errors
        )
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_gl_accounts_xlsx_multisheet_autodetects_and_replays():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="GL Account XLSX Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="gl-account-xlsx",
            ),
            db=db,
            current_user=admin,
        )

        workbook = Workbook()
        ws = workbook.active
        ws.title = "Chart of Accounts"
        ws.append([
            "GL Account ID", "Number", "Name", "Type",
            "Fund Account", "Is Corporate Account",
        ])
        ws.append([
            "GL-6100", "6100", "Repairs", "Expense",
            "Operating", False,
        ])
        notes = workbook.create_sheet("Notes")
        notes.append(["Comment"])
        notes.append(["not a supported migration report"])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        first = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("chart.xlsx", buffer.getvalue()),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert first.detected_resource == "GL_ACCOUNTS"
        assert first.file_format == "XLSX"
        assert first.sheet_name == "Chart of Accounts"
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == first.id
        ).one()
        assert row.source_id == "GL-6100"
        assert row.normalized_data["account_number"] == "6100"
        assert row.normalized_data["account_name"] == "Repairs"
        assert row.normalized_data["account_type"] == "Expense"

        replay = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("chart-copy.xlsx", buffer.getvalue()),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert replay.replayed is True
        assert replay.id == first.id
        assert db.query(PlatformMigrationUpload).filter(
            PlatformMigrationUpload.run_id == run.id,
            PlatformMigrationUpload.detected_resource == "GL_ACCOUNTS",
        ).count() == 1
        assert db.query(GLAccount).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio GL Account staged reconciliation
# ---------------------------------------------------------------------------

def test_appfolio_gl_account_reconciliation_matches_existing_without_accounting_mutation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="GL Resolve Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="gl-resolve",
            ),
            db=db,
            current_user=admin,
        )
        target = GLAccount(
            organization_id=org.id,
            gl_number="4100",
            name="Existing Rent Income",
            account_type="INCOME",
            is_active=True,
        )
        db.add(target)
        db.commit()
        db.refresh(target)

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "gl-resolve.csv",
                    (
                        "GL Account ID,Number,Name,Type,FundAccount\n"
                        "GL-4100,4100,Rent Income,Income,Operating\n"
                    ).encode(),
                ),
                resource="GL_ACCOUNTS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.disposition == "REVIEW"
        before_fingerprint = api._staged_review_fingerprint(upload, [row])
        before_accounts = db.query(GLAccount).count()
        before_transactions = db.query(GLTransaction).count()

        run.last_dry_run_fingerprint = "a" * 64
        run.last_dry_run_summary = {"old": True}
        db.commit()

        resolved = api.resolve_staged_appfolio_gl_account(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedGLAccountResolutionIn(
                action="MATCH_EXISTING",
                target_gl_account_id=target.id,
            ),
            db=db,
            current_user=admin,
        )
        assert resolved.resolution_action == "MATCH_EXISTING"
        assert resolved.resolution_target_gl_account_id == target.id
        assert resolved.resolution_target_id is None
        assert resolved.resolution_target_unit_id is None
        assert resolved.resolution_target_owner_user_id is None
        assert resolved.resolution_target_vendor_id is None
        assert resolved.resolution_target_tenant_user_id is None
        assert api._staged_review_fingerprint(upload, [resolved]) != before_fingerprint

        db.refresh(run)
        assert run.last_dry_run_fingerprint is None
        assert run.last_dry_run_summary is None
        assert run.status == "STAGED"
        assert db.query(GLAccount).count() == before_accounts
        assert db.query(GLTransaction).count() == before_transactions
        assert db.get(GLAccount, target.id).gl_number == "4100"
        assert db.get(GLAccount, target.id).account_type == "INCOME"
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == row.id,
            AuditLog.action == "appfolio_gl_account_resolution_changed",
        ).count() == 1

        replay = api.resolve_staged_appfolio_gl_account(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedGLAccountResolutionIn(
                action="MATCH_EXISTING",
                target_gl_account_id=target.id,
            ),
            db=db,
            current_user=admin,
        )
        assert replay.id == row.id
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == row.id,
            AuditLog.action == "appfolio_gl_account_resolution_changed",
        ).count() == 1
    finally:
        db.close()
        engine.dispose()


def test_appfolio_gl_account_missing_source_id_may_only_be_skipped():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="GL Missing ID Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="gl-missing-id",
            ),
            db=db,
            current_user=admin,
        )
        target = GLAccount(
            organization_id=org.id,
            gl_number="1000",
            name="Operating Cash",
            account_type="ASSET",
            is_active=True,
        )
        db.add(target)
        db.commit()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "gl-no-id.csv",
                    (
                        "Number,Name,Type\n"
                        "1000,Operating Cash,Cash\n"
                    ).encode(),
                ),
                resource="GL_ACCOUNTS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.source_id is None
        assert row.disposition == "REVIEW"

        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_gl_account(
                run.id,
                upload.id,
                row.id,
                AppFolioStagedGLAccountResolutionIn(
                    action="MATCH_EXISTING",
                    target_gl_account_id=target.id,
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409

        skipped = api.resolve_staged_appfolio_gl_account(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedGLAccountResolutionIn(action="SKIP"),
            db=db,
            current_user=admin,
        )
        assert skipped.resolution_action == "SKIP"
        assert skipped.resolution_target_gl_account_id is None
        assert db.query(GLAccount).count() == 1
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_gl_account_reconciliation_rejects_cross_org_inactive_and_invalid_rows():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="GL Scope Org")
        other = _org(db, name="GL Foreign Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="gl-scope",
            ),
            db=db,
            current_user=admin,
        )
        local = GLAccount(
            organization_id=org.id,
            gl_number="6100",
            name="Repairs",
            account_type="EXPENSE",
            is_active=False,
        )
        foreign = GLAccount(
            organization_id=other.id,
            gl_number="6100",
            name="Foreign Repairs",
            account_type="EXPENSE",
            is_active=True,
        )
        db.add_all([local, foreign])
        db.commit()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "gl-scope.csv",
                    (
                        "GL Account ID,Number,Name,Type\n"
                        "GL-6100,6100,Repairs,Expense\n"
                    ).encode(),
                ),
                resource="GL_ACCOUNTS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()

        for target_id in (local.id, foreign.id):
            with pytest.raises(HTTPException) as exc:
                api.resolve_staged_appfolio_gl_account(
                    run.id,
                    upload.id,
                    row.id,
                    AppFolioStagedGLAccountResolutionIn(
                        action="MATCH_EXISTING",
                        target_gl_account_id=target_id,
                    ),
                    db=db,
                    current_user=admin,
                )
            assert exc.value.status_code == 404

        row.errors = ["contradictory source classification"]
        row.disposition = "INVALID"
        db.commit()
        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_gl_account(
                run.id,
                upload.id,
                row.id,
                AppFolioStagedGLAccountResolutionIn(action="SKIP"),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio GL Account staged dry run + controlled mapping commit
# ---------------------------------------------------------------------------

def test_staged_gl_account_dry_run_and_mapping_commit_is_exact_replay_safe_and_non_mutating():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="GL Mapping Commit Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="gl-mapping-commit",
            ),
            db=db,
            current_user=admin,
        )
        target = GLAccount(
            organization_id=org.id,
            gl_number="4100",
            name="Local Rent Income",
            account_type="INCOME",
            is_active=True,
        )
        db.add(target)
        db.commit()
        db.refresh(target)

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "gl-mapping.csv",
                    (
                        "GL Account ID,Number,Name,Type,FundAccount\n"
                        "GL-COMMIT-4100,4100,Rent Income,Income,Operating\n"
                    ).encode(),
                ),
                resource="GL_ACCOUNTS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        api.resolve_staged_appfolio_gl_account(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedGLAccountResolutionIn(
                action="MATCH_EXISTING",
                target_gl_account_id=target.id,
            ),
            db=db,
            current_user=admin,
        )

        before_accounts = db.query(GLAccount).count()
        before_transactions = db.query(GLTransaction).count()
        before_target = (
            db.get(GLAccount, target.id).gl_number,
            db.get(GLAccount, target.id).name,
            db.get(GLAccount, target.id).account_type,
        )

        preview = api.dry_run_staged_appfolio_gl_accounts(
            run.id,
            upload.id,
            db=db,
            current_user=admin,
        )
        assert preview.total == 1
        assert preview.importable == 1
        assert preview.invalid == 0
        assert preview.rows[0].mapped == {"target_gl_account_id": target.id}
        assert db.query(GLAccount).count() == before_accounts
        assert db.query(GLTransaction).count() == before_transactions

        committed = api.commit_staged_appfolio_gl_accounts(
            run.id,
            upload.id,
            AppFolioStagedGLAccountCommitIn(fingerprint=preview.fingerprint),
            db=db,
            current_user=admin,
        )
        assert committed.mapped_existing == 1
        assert committed.replayed is False
        assert committed.rows[0].target_gl_account_id == target.id
        mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "GL_ACCOUNTS",
            PlatformMigrationItem.source_id == "GL-COMMIT-4100",
        ).one()
        assert mapping.target_entity == "GL_ACCOUNT"
        assert mapping.target_id == target.id

        replay = api.commit_staged_appfolio_gl_accounts(
            run.id,
            upload.id,
            AppFolioStagedGLAccountCommitIn(fingerprint=preview.fingerprint),
            db=db,
            current_user=admin,
        )
        assert replay.replayed is True
        assert replay.mapped_existing == 0
        assert replay.rows[0].replayed is True
        assert db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "GL_ACCOUNTS",
            PlatformMigrationItem.source_id == "GL-COMMIT-4100",
        ).count() == 1
        assert db.query(GLAccount).count() == before_accounts
        assert db.query(GLTransaction).count() == before_transactions
        assert (
            db.get(GLAccount, target.id).gl_number,
            db.get(GLAccount, target.id).name,
            db.get(GLAccount, target.id).account_type,
        ) == before_target
    finally:
        db.close()
        engine.dispose()


def test_staged_gl_account_commit_rejects_stale_resolution_fingerprint():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="GL Stale Mapping Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="gl-stale",
            ),
            db=db,
            current_user=admin,
        )
        first = GLAccount(
            organization_id=org.id,
            gl_number="6100",
            name="Repairs One",
            account_type="EXPENSE",
            is_active=True,
        )
        second = GLAccount(
            organization_id=org.id,
            gl_number="6110",
            name="Repairs Two",
            account_type="EXPENSE",
            is_active=True,
        )
        db.add_all([first, second])
        db.commit()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "gl-stale.csv",
                    (
                        "GL Account ID,Number,Name,Type\n"
                        "GL-STALE-6100,6100,Repairs,Expense\n"
                    ).encode(),
                ),
                resource="GL_ACCOUNTS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        api.resolve_staged_appfolio_gl_account(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedGLAccountResolutionIn(
                action="MATCH_EXISTING",
                target_gl_account_id=first.id,
            ),
            db=db,
            current_user=admin,
        )
        preview = api.dry_run_staged_appfolio_gl_accounts(
            run.id,
            upload.id,
            db=db,
            current_user=admin,
        )

        api.resolve_staged_appfolio_gl_account(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedGLAccountResolutionIn(
                action="MATCH_EXISTING",
                target_gl_account_id=second.id,
            ),
            db=db,
            current_user=admin,
        )
        with pytest.raises(HTTPException) as exc:
            api.commit_staged_appfolio_gl_accounts(
                run.id,
                upload.id,
                AppFolioStagedGLAccountCommitIn(
                    fingerprint=preview.fingerprint
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "GL_ACCOUNTS",
        ).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_staged_gl_account_dry_run_revalidates_target_and_skip_only_rows_do_not_commit():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="GL Revalidate Mapping Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="gl-revalidate",
            ),
            db=db,
            current_user=admin,
        )
        target = GLAccount(
            organization_id=org.id,
            gl_number="1000",
            name="Operating Cash",
            account_type="ASSET",
            is_active=True,
        )
        db.add(target)
        db.commit()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "gl-revalidate.csv",
                    (
                        "GL Account ID,Number,Name,Type\n"
                        "GL-VALID-1000,1000,Operating Cash,Cash\n"
                    ).encode(),
                ),
                resource="GL_ACCOUNTS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        api.resolve_staged_appfolio_gl_account(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedGLAccountResolutionIn(
                action="MATCH_EXISTING",
                target_gl_account_id=target.id,
            ),
            db=db,
            current_user=admin,
        )
        target.is_active = False
        db.commit()
        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_gl_accounts(
                run.id,
                upload.id,
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409

        target.is_active = True
        db.commit()
        api.resolve_staged_appfolio_gl_account(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedGLAccountResolutionIn(action="SKIP"),
            db=db,
            current_user=admin,
        )
        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_gl_accounts(
                run.id,
                upload.id,
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert "No staged GL Account rows remain" in exc.value.detail
        assert db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "GL_ACCOUNTS",
        ).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio General Ledger history ingestion + staging
# ---------------------------------------------------------------------------

def _seed_general_ledger_source_mappings(db, *, run, org):
    prop = Property(
        organization_id=org.id,
        name="Ledger Stage Property",
        address_line1="900 Ledger Ave",
        city="Cleveland",
        state="OH",
        zip_code="44113",
        is_active=True,
    )
    db.add(prop)
    db.flush()
    unit = Unit(
        property_id=prop.id,
        unit_number="GL-101",
        bedrooms=1,
        bathrooms=1,
        monthly_rent=1000,
        is_active=True,
    )
    account = GLAccount(
        organization_id=org.id,
        gl_number="4100",
        name="Rent Income",
        account_type="INCOME",
        is_active=True,
    )
    db.add_all([unit, account])
    db.flush()
    for resource, source_id, target_entity, target_id in (
        ("PROPERTIES", "PROP-GL-1", "PROPERTY", prop.id),
        ("UNITS", "UNIT-GL-1", "UNIT", unit.id),
        ("GL_ACCOUNTS", "GL-4100", "GL_ACCOUNT", account.id),
    ):
        db.add(
            PlatformMigrationItem(
                run_id=run.id,
                organization_id=org.id,
                provider="APPFOLIO",
                resource=resource,
                source_id=source_id,
                target_entity=target_entity,
                target_id=target_id,
                source_fingerprint=f"seed-{resource}-{source_id}",
            )
        )
    db.commit()
    return prop, unit, account


def test_appfolio_general_ledger_csv_staging_preserves_verified_source_evidence_and_replays():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="General Ledger Stage Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="general-ledger-stage",
            ),
            db=db,
            current_user=admin,
        )
        _seed_general_ledger_source_mappings(db, run=run, org=org)
        before_gl = db.query(GLTransaction).count()
        before_charges = db.query(Charge).count()
        before_accounts = db.query(GLAccount).count()

        content = (
            "LineItemId,TransactionId,GlAccountId,PropertyId,UnitId,Date,Debit,Credit,"
            "Description,Reference,Remarks,TransactionType\n"
            "LINE-1,TX-100,GL-4100,PROP-GL-1,UNIT-GL-1,2026-01-15,0.00,1250.00,"
            "January rent,RCP-100,Imported source evidence,Receipt\n"
        ).encode()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("general-ledger.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert upload.detected_resource == "GENERAL_LEDGER"
        assert upload.status == "REVIEW_REQUIRED"
        assert upload.validation_summary["general_ledger_rows"] == 1
        assert upload.validation_summary["missing_general_ledger_line_ids"] == 0
        assert upload.validation_summary["unresolved_general_ledger_gl_accounts"] == 0
        assert upload.validation_summary["accounting_history_mutation"] is False
        assert upload.validation_summary["customer_accounting_mutation"] is False

        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.resource == "GENERAL_LEDGER"
        assert row.source_id == "LINE-1"
        assert row.disposition == "REVIEW"
        assert row.normalized_data["transaction_id"] == "TX-100"
        assert row.normalized_data["source_gl_account_id"] == "GL-4100"
        assert row.normalized_data["source_property_id"] == "PROP-GL-1"
        assert row.normalized_data["source_unit_id"] == "UNIT-GL-1"
        assert row.normalized_data["posted_date"] == "2026-01-15"
        assert row.normalized_data["debit"] == "0.00"
        assert row.normalized_data["credit"] == "1250.00"
        assert row.normalized_data["transaction_type"] == "Receipt"
        assert any("source evidence only" in warning for warning in row.warnings)
        assert any("separate accounting-history reconciliation batch" in warning for warning in row.warnings)

        replay = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("general-ledger-copy.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert replay.replayed is True
        assert replay.id == upload.id
        assert db.query(GLTransaction).count() == before_gl
        assert db.query(Charge).count() == before_charges
        assert db.query(GLAccount).count() == before_accounts
    finally:
        db.close()
        engine.dispose()


def test_appfolio_general_ledger_xlsx_detects_verified_fields_and_flags_identity_relationship_issues():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="General Ledger XLSX Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="general-ledger-xlsx",
            ),
            db=db,
            current_user=admin,
        )

        workbook = Workbook()
        ws = workbook.active
        ws.title = "General Ledger"
        ws.append([
            "LineItemId", "TransactionId", "GlAccountId", "PropertyId", "UnitId",
            "Date", "Debit", "Credit", "Description", "Reference", "Remarks",
            "TransactionType",
        ])
        ws.append([
            None, "TX-MISSING", "GL-NOT-MAPPED", None, None,
            "2026-02-01", "25.00", "0.00", "Repair", "BILL-1", None, "Bill",
        ])
        ws.append([
            "LINE-DUP", "TX-1", "GL-NOT-MAPPED", None, None,
            "2026-02-02", "10.00", "0.00", "Repair A", "BILL-2", None, "Bill",
        ])
        ws.append([
            "LINE-DUP", "TX-2", "GL-NOT-MAPPED", None, None,
            "2026-02-03", "15.00", "0.00", "Repair B", "BILL-3", None, "Bill",
        ])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        before_gl = db.query(GLTransaction).count()
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("general-ledger.xlsx", buffer.getvalue()),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert upload.detected_resource == "GENERAL_LEDGER"
        assert upload.validation_summary["general_ledger_rows"] == 3
        assert upload.validation_summary["missing_general_ledger_line_ids"] == 1
        assert upload.validation_summary["unresolved_general_ledger_gl_accounts"] == 3
        assert upload.validation_summary["duplicates"] == 1

        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()
        assert rows[0].disposition == "REVIEW"
        assert rows[0].source_id is None
        assert any("LineItemId was not supplied" in warning for warning in rows[0].warnings)
        assert any("GL Account relationship is unresolved" in warning for warning in rows[0].warnings)
        assert rows[1].disposition == "REVIEW"
        assert rows[2].disposition == "INVALID"
        assert any("Duplicate AppFolio General Ledger LineItemId" in error for error in rows[2].errors)
        assert db.query(GLTransaction).count() == before_gl
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio General Ledger staged relationship reconciliation
# ---------------------------------------------------------------------------

def test_appfolio_general_ledger_relationship_acceptance_is_scoped_idempotent_and_non_mutating():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="General Ledger Resolve Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="general-ledger-resolve",
            ),
            db=db,
            current_user=admin,
        )
        prop, unit, account = _seed_general_ledger_source_mappings(
            db, run=run, org=org
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "general-ledger-resolve.csv",
                    (
                        "LineItemId,TransactionId,GlAccountId,PropertyId,UnitId,Date,Debit,Credit\n"
                        "LINE-RESOLVE-1,TX-RESOLVE-1,GL-4100,PROP-GL-1,UNIT-GL-1,"
                        "2026-03-01,0.00,1250.00\n"
                    ).encode(),
                ),
                resource="GENERAL_LEDGER",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        before_fingerprint = api._staged_review_fingerprint(upload, [row])
        before_gl = db.query(GLTransaction).count()
        before_accounts = db.query(GLAccount).count()
        before_charges = db.query(Charge).count()

        resolved = api.resolve_staged_appfolio_general_ledger(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedGeneralLedgerResolutionIn(
                action="ACCEPT_RELATIONSHIP"
            ),
            db=db,
            current_user=admin,
        )
        assert resolved.resolution_action == "ACCEPT_RELATIONSHIP"
        assert resolved.resolution_target_gl_account_id == account.id
        assert resolved.resolution_target_id == prop.id
        assert resolved.resolution_target_unit_id == unit.id
        assert resolved.resolution_target_owner_user_id is None
        assert resolved.resolution_target_vendor_id is None
        assert resolved.resolution_target_tenant_user_id is None
        assert api._staged_review_fingerprint(upload, [resolved]) != before_fingerprint
        db.refresh(run)
        assert run.last_dry_run_fingerprint is None
        assert run.last_dry_run_summary is None

        audit_count = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == row.id,
            AuditLog.action == "appfolio_general_ledger_resolution_changed",
        ).count()
        assert audit_count == 1

        replay = api.resolve_staged_appfolio_general_ledger(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedGeneralLedgerResolutionIn(
                action="ACCEPT_RELATIONSHIP"
            ),
            db=db,
            current_user=admin,
        )
        assert replay.id == row.id
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == row.id,
            AuditLog.action == "appfolio_general_ledger_resolution_changed",
        ).count() == audit_count

        assert db.query(GLTransaction).count() == before_gl
        assert db.query(GLAccount).count() == before_accounts
        assert db.query(Charge).count() == before_charges

    finally:
        db.close()
        engine.dispose()


def test_appfolio_general_ledger_missing_identity_or_unresolved_gl_may_only_skip():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="General Ledger Resolve Guard Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="general-ledger-guard",
            ),
            db=db,
            current_user=admin,
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "general-ledger-guard.csv",
                    (
                        "LineItemId,TransactionId,GlAccountId,Date,Debit,Credit\n"
                        ",TX-NO-LINE,GL-MISSING,2026-03-02,10.00,0.00\n"
                        "LINE-NO-GLMAP,TX-NO-GLMAP,GL-MISSING,2026-03-03,20.00,0.00\n"
                    ).encode(),
                ),
                resource="GENERAL_LEDGER",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()

        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_general_ledger(
                run.id,
                upload.id,
                rows[0].id,
                AppFolioStagedGeneralLedgerResolutionIn(
                    action="ACCEPT_RELATIONSHIP"
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert "LineItemId" in exc.value.detail

        skipped = api.resolve_staged_appfolio_general_ledger(
            run.id,
            upload.id,
            rows[0].id,
            AppFolioStagedGeneralLedgerResolutionIn(action="SKIP"),
            db=db,
            current_user=admin,
        )
        assert skipped.resolution_action == "SKIP"
        assert skipped.resolution_target_gl_account_id is None

        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_general_ledger(
                run.id,
                upload.id,
                rows[1].id,
                AppFolioStagedGeneralLedgerResolutionIn(
                    action="ACCEPT_RELATIONSHIP"
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert "durable GL_ACCOUNT mapping" in exc.value.detail
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_general_ledger_relationship_rejects_unit_property_mismatch_and_invalid_rows():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="General Ledger Relationship Mismatch Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="general-ledger-mismatch",
            ),
            db=db,
            current_user=admin,
        )
        prop, unit, account = _seed_general_ledger_source_mappings(
            db, run=run, org=org
        )
        other_prop = Property(
            organization_id=org.id,
            name="Other Ledger Property",
            address_line1="901 Ledger Ave",
            city="Cleveland",
            state="OH",
            zip_code="44113",
            is_active=True,
        )
        db.add(other_prop)
        db.flush()
        other_unit = Unit(
            property_id=other_prop.id,
            unit_number="GL-OTHER",
            bedrooms=1,
            bathrooms=1,
            monthly_rent=900,
            is_active=True,
        )
        db.add(other_unit)
        db.flush()
        unit_mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "UNITS",
            PlatformMigrationItem.source_id == "UNIT-GL-1",
        ).one()
        unit_mapping.target_id = other_unit.id
        db.commit()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "general-ledger-mismatch.csv",
                    (
                        "LineItemId,TransactionId,GlAccountId,PropertyId,UnitId,Date,Debit,Credit\n"
                        "LINE-MISMATCH,TX-MISMATCH,GL-4100,PROP-GL-1,UNIT-GL-1,"
                        "2026-03-04,50.00,0.00\n"
                        "LINE-DUP,TX-DUP-1,GL-4100,PROP-GL-1,,2026-03-05,25.00,0.00\n"
                        "LINE-DUP,TX-DUP-2,GL-4100,PROP-GL-1,,2026-03-06,25.00,0.00\n"
                    ).encode(),
                ),
                resource="GENERAL_LEDGER",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()

        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_general_ledger(
                run.id,
                upload.id,
                rows[0].id,
                AppFolioStagedGeneralLedgerResolutionIn(
                    action="ACCEPT_RELATIONSHIP"
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert "same target Property" in exc.value.detail

        assert rows[2].disposition == "INVALID"
        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_general_ledger(
                run.id,
                upload.id,
                rows[2].id,
                AppFolioStagedGeneralLedgerResolutionIn(action="SKIP"),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio General Ledger staged accounting-history dry run
# ---------------------------------------------------------------------------

def test_appfolio_general_ledger_dry_run_is_exact_replay_safe_and_non_mutating():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="General Ledger Dry Run Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="general-ledger-dry-run",
            ),
            db=db,
            current_user=admin,
        )
        prop, unit, account = _seed_general_ledger_source_mappings(
            db, run=run, org=org
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "general-ledger-dry-run.csv",
                    (
                        "LineItemId,TransactionId,GlAccountId,PropertyId,UnitId,Date,Debit,Credit,"
                        "Description,Reference,Remarks,TransactionType\n"
                        "LINE-DRY-1,TX-DRY-1,GL-4100,PROP-GL-1,UNIT-GL-1,2026-04-01,"
                        "0.00,1350.00,April rent,RCP-DRY,source note,Receipt\n"
                    ).encode(),
                ),
                resource="GENERAL_LEDGER",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        api.resolve_staged_appfolio_general_ledger(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedGeneralLedgerResolutionIn(action="ACCEPT_RELATIONSHIP"),
            db=db,
            current_user=admin,
        )

        before_gl = db.query(GLTransaction).count()
        before_accounts = db.query(GLAccount).count()
        before_charges = db.query(Charge).count()

        first = api.dry_run_staged_appfolio_general_ledger(
            run.id, upload.id, db=db, current_user=admin
        )
        assert first.replayed is False
        assert first.total == 1
        assert first.importable == 1
        assert first.invalid == 0
        assert len(first.rows) == 1
        preview = first.rows[0]
        assert preview.source_id == "LINE-DRY-1"
        assert preview.source_evidence["transaction_id"] == "TX-DRY-1"
        assert preview.source_evidence["posted_date"] == "2026-04-01"
        assert preview.source_evidence["debit"] == "0.00"
        assert preview.source_evidence["credit"] == "1350.00"
        assert preview.source_evidence["transaction_type"] == "Receipt"
        assert preview.resolved_targets == {
            "gl_account_id": account.id,
            "property_id": prop.id,
            "unit_id": unit.id,
        }
        db.refresh(run)
        assert run.last_dry_run_fingerprint == first.fingerprint
        assert run.last_dry_run_summary["resource"] == "GENERAL_LEDGER"
        assert run.last_dry_run_summary["review_only"] is True
        assert run.last_dry_run_summary["accounting_history_mutation"] is False

        audit_count = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_staged_general_ledger_dry_run",
        ).count()
        assert audit_count == 1

        replay = api.dry_run_staged_appfolio_general_ledger(
            run.id, upload.id, db=db, current_user=admin
        )
        assert replay.replayed is True
        assert replay.fingerprint == first.fingerprint
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_staged_general_ledger_dry_run",
        ).count() == audit_count
        assert db.query(GLTransaction).count() == before_gl
        assert db.query(GLAccount).count() == before_accounts
        assert db.query(Charge).count() == before_charges
    finally:
        db.close()
        engine.dispose()


def test_appfolio_general_ledger_dry_run_blocks_unresolved_or_stale_relationships():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="General Ledger Dry Run Guard Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="general-ledger-dry-run-guard",
            ),
            db=db,
            current_user=admin,
        )
        _prop, _unit, account = _seed_general_ledger_source_mappings(
            db, run=run, org=org
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "general-ledger-dry-run-guard.csv",
                    (
                        "LineItemId,TransactionId,GlAccountId,PropertyId,UnitId,Date,Debit,Credit\n"
                        "LINE-GUARD-1,TX-GUARD-1,GL-4100,PROP-GL-1,UNIT-GL-1,2026-04-02,25.00,0.00\n"
                    ).encode(),
                ),
                resource="GENERAL_LEDGER",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()

        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_general_ledger(
                run.id, upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "unresolved" in exc.value.detail

        api.resolve_staged_appfolio_general_ledger(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedGeneralLedgerResolutionIn(action="ACCEPT_RELATIONSHIP"),
            db=db,
            current_user=admin,
        )
        replacement = GLAccount(
            organization_id=org.id,
            gl_number="4199",
            name="Replacement Review Account",
            account_type="INCOME",
            is_active=True,
        )
        db.add(replacement)
        db.flush()
        mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "GL_ACCOUNTS",
            PlatformMigrationItem.source_id == "GL-4100",
        ).one()
        mapping.target_id = replacement.id
        mapping.source_fingerprint = "changed-gl-mapping"
        db.commit()

        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_general_ledger(
                run.id, upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "stale or inconsistent" in exc.value.detail
        assert db.get(GLAccount, account.id).is_active is True
        assert db.query(GLTransaction).count() == 0
        assert db.query(Charge).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio General Ledger commit-readiness analysis
# ---------------------------------------------------------------------------

def test_appfolio_general_ledger_commit_readiness_groups_supplied_transaction_ids_and_replays():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="General Ledger Readiness Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="general-ledger-readiness",
            ),
            db=db,
            current_user=admin,
        )
        prop, unit, income = _seed_general_ledger_source_mappings(
            db, run=run, org=org
        )
        cash = GLAccount(
            organization_id=org.id,
            gl_number="1000",
            name="Operating Cash",
            account_type="ASSET",
            is_active=True,
        )
        db.add(cash)
        db.flush()
        db.add(
            PlatformMigrationItem(
                run_id=run.id,
                organization_id=org.id,
                provider="APPFOLIO",
                resource="GL_ACCOUNTS",
                source_id="GL-1000",
                target_entity="GL_ACCOUNT",
                target_id=cash.id,
                source_fingerprint="seed-GL_ACCOUNTS-GL-1000",
            )
        )
        db.commit()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "general-ledger-readiness.csv",
                    (
                        "LineItemId,TransactionId,GlAccountId,PropertyId,UnitId,Date,Debit,Credit,"
                        "Description,Reference,Remarks,TransactionType\n"
                        "LINE-RDY-1,TX-RDY-1,GL-1000,PROP-GL-1,UNIT-GL-1,2026-04-05,"
                        "1350.00,0.00,Cash side,RDY-1,source evidence,Receipt\n"
                        "LINE-RDY-2,TX-RDY-1,GL-4100,PROP-GL-1,UNIT-GL-1,2026-04-05,"
                        "0.00,1350.00,Income side,RDY-1,source evidence,Receipt\n"
                    ).encode(),
                ),
                resource="GENERAL_LEDGER",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()
        for row in rows:
            api.resolve_staged_appfolio_general_ledger(
                run.id,
                upload.id,
                row.id,
                AppFolioStagedGeneralLedgerResolutionIn(action="ACCEPT_RELATIONSHIP"),
                db=db,
                current_user=admin,
            )

        dry = api.dry_run_staged_appfolio_general_ledger(
            run.id, upload.id, db=db, current_user=admin
        )
        before_gl = db.query(GLTransaction).count()
        before_accounts = db.query(GLAccount).count()
        before_charges = db.query(Charge).count()

        first = api.analyze_staged_appfolio_general_ledger_commit_readiness(
            run.id, upload.id, db=db, current_user=admin
        )
        assert first.replayed is False
        assert first.dry_run_fingerprint == dry.fingerprint
        assert first.group_count == 1
        assert first.line_count == 2
        group = first.groups[0]
        assert group.transaction_id == "TX-RDY-1"
        assert group.line_count == 2
        assert group.debit_total == "1350.00"
        assert group.credit_total == "1350.00"
        assert group.balanced is True
        assert {line.source_id for line in group.lines} == {"LINE-RDY-1", "LINE-RDY-2"}
        assert {line.resolved_targets["property_id"] for line in group.lines} == {prop.id}
        assert {line.resolved_targets["unit_id"] for line in group.lines} == {unit.id}
        assert {line.resolved_targets["gl_account_id"] for line in group.lines} == {
            cash.id,
            income.id,
        }

        reconciliation_response = Response()
        reconciliation = api.get_staged_appfolio_general_ledger_reconciliation(
            run.id,
            upload.id,
            response=reconciliation_response,
            db=db,
            current_user=admin,
        )
        assert reconciliation_response.headers["cache-control"] == "no-store"
        assert reconciliation.dry_run_fingerprint == dry.fingerprint
        assert reconciliation.readiness_fingerprint == first.readiness_fingerprint
        assert reconciliation.source_line_count == 2
        assert reconciliation.preview_line_count == 2
        assert reconciliation.source_debit_total == "1350.00"
        assert reconciliation.source_credit_total == "1350.00"
        assert reconciliation.preview_debit_total == "1350.00"
        assert reconciliation.preview_credit_total == "1350.00"
        assert reconciliation.line_count_match is True
        assert reconciliation.line_identity_match is True
        assert reconciliation.debit_total_match is True
        assert reconciliation.credit_total_match is True
        assert reconciliation.transaction_groups_balanced is True
        assert reconciliation.account_controls_match is True
        assert [
            control.source_gl_account_id for control in reconciliation.account_controls
        ] == ["GL-1000", "GL-4100"]
        cash_control, income_control = reconciliation.account_controls
        assert cash_control.target_gl_account_id == cash.id
        assert cash_control.source_line_count == 1
        assert cash_control.preview_line_count == 1
        assert cash_control.source_debit_total == "1350.00"
        assert cash_control.preview_debit_total == "1350.00"
        assert cash_control.source_credit_total == "0.00"
        assert cash_control.preview_credit_total == "0.00"
        assert cash_control.reconciled is True
        assert income_control.target_gl_account_id == income.id
        assert income_control.source_line_count == 1
        assert income_control.preview_line_count == 1
        assert income_control.source_debit_total == "0.00"
        assert income_control.preview_debit_total == "0.00"
        assert income_control.source_credit_total == "1350.00"
        assert income_control.preview_credit_total == "1350.00"
        assert income_control.reconciled is True
        assert "gl_account_debit_credit_totals" in reconciliation.compared_controls
        assert reconciliation.reconciled is True
        assert reconciliation.accounting_complete is False
        assert "security_deposit_total" in reconciliation.unverified_controls
        assert "receivables_total" in reconciliation.unverified_controls

        db.refresh(run)
        assert run.last_dry_run_fingerprint == dry.fingerprint
        assert run.last_dry_run_summary["ledger_commit_ready"] is True
        assert (
            run.last_dry_run_summary["ledger_commit_readiness_fingerprint"]
            == first.readiness_fingerprint
        )
        audit_count = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_staged_general_ledger_commit_readiness",
        ).count()
        assert audit_count == 1

        replay = api.analyze_staged_appfolio_general_ledger_commit_readiness(
            run.id, upload.id, db=db, current_user=admin
        )
        assert replay.replayed is True
        assert replay.readiness_fingerprint == first.readiness_fingerprint
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_staged_general_ledger_commit_readiness",
        ).count() == audit_count
        assert db.query(GLTransaction).count() == before_gl
        assert db.query(GLAccount).count() == before_accounts
        assert db.query(Charge).count() == before_charges

        cash_mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.organization_id == org.id,
            PlatformMigrationItem.provider == "APPFOLIO",
            PlatformMigrationItem.resource == "GL_ACCOUNTS",
            PlatformMigrationItem.source_id == "GL-1000",
        ).one()
        original_mapping_fingerprint = cash_mapping.source_fingerprint
        cash_mapping.source_fingerprint = "changed-GL-1000-source-fingerprint"
        db.commit()
        with pytest.raises(HTTPException) as exc:
            api.get_staged_appfolio_general_ledger_reconciliation(
                run.id,
                upload.id,
                response=Response(),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert "exact latest staged dry-run" in exc.value.detail
        assert db.query(GLTransaction).count() == before_gl
        assert db.query(GLAccount).count() == before_accounts
        assert db.query(Charge).count() == before_charges

        cash_mapping.source_fingerprint = original_mapping_fingerprint
        db.commit()
        api.dry_run_staged_appfolio_general_ledger(
            run.id, upload.id, db=db, current_user=admin
        )
        api.analyze_staged_appfolio_general_ledger_commit_readiness(
            run.id, upload.id, db=db, current_user=admin
        )

        stale_row = db.get(PlatformMigrationStagedRow, rows[0].id)
        stale_row.normalized_data = {
            **dict(stale_row.normalized_data or {}),
            "debit": "1400.00",
        }
        db.commit()
        with pytest.raises(HTTPException) as exc:
            api.get_staged_appfolio_general_ledger_reconciliation(
                run.id,
                upload.id,
                response=Response(),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert "exact latest staged dry-run" in exc.value.detail
        assert db.query(GLTransaction).count() == before_gl
        assert db.query(GLAccount).count() == before_accounts
        assert db.query(Charge).count() == before_charges
    finally:
        db.close()
        engine.dispose()


def test_appfolio_general_ledger_reconciliation_route_is_exposed_and_sales_is_denied():
    from app.main import app

    assert (
        "/api/platform/migrations/appfolio/runs/{run_id}/uploads/{upload_id}/general-ledger/reconciliation"
        in set(app.openapi()["paths"])
    )

    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        sales = _platform_user(db, PlatformUserRole.PLATFORM_SALES)
        org = _org(db, name="Ledger Reconciliation Auth Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="ledger-reconciliation-auth",
            ),
            db=db,
            current_user=admin,
        )
        with pytest.raises(HTTPException) as exc:
            api.get_staged_appfolio_general_ledger_reconciliation(
                run.id,
                999999,
                response=Response(),
                db=db,
                current_user=sales,
            )
        assert exc.value.status_code == 403
    finally:
        db.close()
        engine.dispose()


def test_appfolio_general_ledger_commit_readiness_blocks_missing_transaction_id_or_unbalanced_group():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="General Ledger Readiness Guard Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="general-ledger-readiness-guard",
            ),
            db=db,
            current_user=admin,
        )
        _seed_general_ledger_source_mappings(db, run=run, org=org)

        missing_tx = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "general-ledger-no-transaction.csv",
                    (
                        "LineItemId,TransactionId,GlAccountId,Date,Debit,Credit\n"
                        "LINE-NOTX-1,,GL-4100,2026-04-06,10.00,10.00\n"
                    ).encode(),
                ),
                resource="GENERAL_LEDGER",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == missing_tx.id
        ).one()
        api.resolve_staged_appfolio_general_ledger(
            run.id,
            missing_tx.id,
            row.id,
            AppFolioStagedGeneralLedgerResolutionIn(action="ACCEPT_RELATIONSHIP"),
            db=db,
            current_user=admin,
        )
        api.dry_run_staged_appfolio_general_ledger(
            run.id, missing_tx.id, db=db, current_user=admin
        )
        with pytest.raises(HTTPException) as exc:
            api.analyze_staged_appfolio_general_ledger_commit_readiness(
                run.id, missing_tx.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "TransactionId" in exc.value.detail

        unbalanced = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "general-ledger-unbalanced.csv",
                    (
                        "LineItemId,TransactionId,GlAccountId,Date,Debit,Credit\n"
                        "LINE-UNBAL-1,TX-UNBAL-1,GL-4100,2026-04-07,25.00,0.00\n"
                        "LINE-UNBAL-2,TX-UNBAL-1,GL-4100,2026-04-07,0.00,20.00\n"
                    ).encode(),
                ),
                resource="GENERAL_LEDGER",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == unbalanced.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()
        for staged in rows:
            api.resolve_staged_appfolio_general_ledger(
                run.id,
                unbalanced.id,
                staged.id,
                AppFolioStagedGeneralLedgerResolutionIn(action="ACCEPT_RELATIONSHIP"),
                db=db,
                current_user=admin,
            )
        api.dry_run_staged_appfolio_general_ledger(
            run.id, unbalanced.id, db=db, current_user=admin
        )
        with pytest.raises(HTTPException) as exc:
            api.analyze_staged_appfolio_general_ledger_commit_readiness(
                run.id, unbalanced.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "not balanced" in exc.value.detail
        assert db.query(GLTransaction).count() == 0
        assert db.query(Charge).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Bills / Payables CSV/XLSX ingestion + staging
# ---------------------------------------------------------------------------

def _seed_bill_source_mappings(db, *, run, org):
    prop = Property(
        organization_id=org.id,
        name="Bill Stage Property",
        address_line1="700 Bill Ave",
        city="Cleveland",
        state="OH",
        zip_code="44113",
        is_active=True,
    )
    vendor = Vendor(
        organization_id=org.id,
        company_name="Bill Stage Vendor",
        business_email="bill-stage-vendor@example.com",
        is_active=True,
    )
    db.add_all([prop, vendor])
    db.flush()
    for resource, source_id, target_entity, target_id in (
        ("PROPERTIES", "PROP-BILL-1", "PROPERTY", prop.id),
        ("VENDORS", "VENDOR-BILL-1", "VENDOR", vendor.id),
    ):
        db.add(
            PlatformMigrationItem(
                run_id=run.id,
                organization_id=org.id,
                provider="APPFOLIO",
                resource=resource,
                source_id=source_id,
                target_entity=target_entity,
                target_id=target_id,
                source_fingerprint=f"seed-{resource}-{source_id}",
            )
        )
    db.commit()
    return prop, vendor


def test_appfolio_bills_csv_stages_verified_top_level_source_evidence_without_mutation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Bills Stage Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="bills-stage",
            ),
            db=db,
            current_user=admin,
        )
        prop, vendor = _seed_bill_source_mappings(db, run=run, org=org)
        before_bills = db.query(Bill).count()
        before_gl = db.query(GLTransaction).count()
        before_charges = db.query(Charge).count()
        before_vendors = db.query(Vendor).count()

        content = (
            "Bill ID,VendorId,PropertyId,DueDate,InvoiceDate,PostingDate,Reference,"
            "Remarks,TotalAmount,ApprovalStatus,CheckMemo,AccountNumber,"
            "ManagementCompanyAsPayee,WorkOrderId,LastUpdatedAt\n"
            "BILL-100,VENDOR-BILL-1,PROP-BILL-1,2026-05-31,2026-05-15,2026-05-16,"
            "INV-100,Source remarks,450.25,Approved,May invoice,2000,false,WO-55,"
            "2026-05-17T12:00:00Z\n"
        ).encode()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("bills.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert upload.detected_resource == "BILLS"
        assert upload.status == "REVIEW_REQUIRED"
        assert upload.validation_summary["bill_rows"] == 1
        assert upload.validation_summary["missing_bill_source_ids"] == 0
        assert upload.validation_summary["unresolved_bill_vendor_relationships"] == 0
        assert upload.validation_summary["unresolved_bill_property_relationships"] == 0
        assert upload.validation_summary["bill_mutation"] is False
        assert upload.validation_summary["accounting_history_mutation"] is False
        assert upload.validation_summary["payment_state_inferred"] is False

        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.resource == "BILLS"
        assert row.source_id == "BILL-100"
        assert row.disposition == "REVIEW"
        assert row.normalized_data == {
            "source_id": "BILL-100",
            "source_vendor_id": "VENDOR-BILL-1",
            "source_property_id": "PROP-BILL-1",
            "due_date": "2026-05-31",
            "invoice_date": "2026-05-15",
            "posting_date": "2026-05-16",
            "reference": "INV-100",
            "remarks": "Source remarks",
            "total_amount": "450.25",
            "approval_status": "Approved",
            "check_memo": "May invoice",
            "account_number": "2000",
            "management_company_as_payee": "false",
            "source_work_order_id": "WO-55",
            "last_updated_at": "2026-05-17T12:00:00Z",
        }
        assert any("WorkOrderId is preserved as source evidence only" in w for w in row.warnings)
        assert any("paid/unpaid" in w for w in row.warnings)
        assert db.query(Bill).count() == before_bills
        assert db.query(GLTransaction).count() == before_gl
        assert db.query(Charge).count() == before_charges
        assert db.query(Vendor).count() == before_vendors
        assert db.get(Property, prop.id).is_active is True
        assert db.get(Vendor, vendor.id).is_active is True
    finally:
        db.close()
        engine.dispose()


def test_appfolio_bills_missing_source_id_stays_review_unresolved_links_warn_and_duplicate_id_invalid():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Bills Review Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="bills-review",
            ),
            db=db,
            current_user=admin,
        )
        content = (
            "Bill ID,VendorId,PropertyId,DueDate,TotalAmount,Reference\n"
            ",VENDOR-MISSING,PROP-MISSING,2026-06-01,100.00,NO-ID\n"
            "BILL-DUP,VENDOR-MISSING,PROP-MISSING,2026-06-02,125.00,DUP-1\n"
            "BILL-DUP,VENDOR-MISSING,PROP-MISSING,2026-06-03,130.00,DUP-2\n"
        ).encode()
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("bills-review.csv", content),
                resource="BILLS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert upload.detected_resource == "BILLS"
        assert upload.status == "STAGED_WITH_ERRORS"
        assert upload.validation_summary["missing_bill_source_ids"] == 1
        assert upload.validation_summary["unresolved_bill_vendor_relationships"] == 3
        assert upload.validation_summary["unresolved_bill_property_relationships"] == 3

        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()
        assert rows[0].disposition == "REVIEW"
        assert any("Bill ID was not supplied" in w for w in rows[0].warnings)
        assert any("Vendor relationship remains unresolved" in w for w in rows[0].warnings)
        assert any("Property relationship remains unresolved" in w for w in rows[0].warnings)
        assert rows[1].disposition == "REVIEW"
        assert rows[2].disposition == "INVALID"
        assert any("Duplicate AppFolio Bill ID" in e for e in rows[2].errors)
        assert db.query(Bill).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_bills_xlsx_multisheet_autodetects_and_replays_without_bill_creation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Bills XLSX Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="bills-xlsx",
            ),
            db=db,
            current_user=admin,
        )
        _seed_bill_source_mappings(db, run=run, org=org)

        workbook = Workbook()
        ws = workbook.active
        ws.title = "Bills"
        ws.append([
            "Bill ID", "VendorId", "PropertyId", "DueDate", "InvoiceDate",
            "TotalAmount", "ApprovalStatus",
        ])
        ws.append([
            "BILL-XLSX-1", "VENDOR-BILL-1", "PROP-BILL-1", "2026-07-01",
            "2026-06-15", "200.50", "Pending",
        ])
        notes = workbook.create_sheet("Notes")
        notes.append(["Comment"])
        notes.append(["not a supported migration report"])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        first = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("bills.xlsx", buffer.getvalue()),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert first.detected_resource == "BILLS"
        assert first.sheet_name == "Bills"
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == first.id
        ).one()
        assert row.source_id == "BILL-XLSX-1"
        assert row.normalized_data["total_amount"] == "200.50"
        assert row.normalized_data["approval_status"] == "Pending"

        replay = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("bills-copy.xlsx", buffer.getvalue()),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert replay.replayed is True
        assert replay.id == first.id
        assert db.query(Bill).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()



# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Bills staged relationship reconciliation
# ---------------------------------------------------------------------------

def test_appfolio_bills_relationship_acceptance_replays_and_does_not_mutate_bills():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Bills Resolve Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="bills-resolve",
            ),
            db=db,
            current_user=admin,
        )
        prop, vendor = _seed_bill_source_mappings(db, run=run, org=org)
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "bills-resolve.csv",
                    (
                        "Bill ID,VendorId,PropertyId,DueDate,TotalAmount,Reference\n"
                        "BILL-RESOLVE-1,VENDOR-BILL-1,PROP-BILL-1,2026-08-01,325.50,INV-RESOLVE\n"
                    ).encode(),
                ),
                resource="BILLS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        before_fingerprint = api._staged_review_fingerprint(upload, [row])
        before_bills = db.query(Bill).count()
        before_gl = db.query(GLTransaction).count()
        before_charges = db.query(Charge).count()
        before_vendors = db.query(Vendor).count()

        resolved = api.resolve_staged_appfolio_bill(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedBillResolutionIn(action="ACCEPT_RELATIONSHIP"),
            db=db,
            current_user=admin,
        )
        assert resolved.resolution_action == "ACCEPT_RELATIONSHIP"
        assert resolved.resolution_target_vendor_id == vendor.id
        assert resolved.resolution_target_id == prop.id
        assert resolved.resolution_target_unit_id is None
        assert resolved.resolution_target_owner_user_id is None
        assert resolved.resolution_target_tenant_user_id is None
        assert resolved.resolution_target_gl_account_id is None
        assert api._staged_review_fingerprint(upload, [resolved]) != before_fingerprint
        db.refresh(run)
        assert run.last_dry_run_fingerprint is None
        assert run.last_dry_run_summary is None

        audit_count = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == row.id,
            AuditLog.action == "appfolio_bill_resolution_changed",
        ).count()
        assert audit_count == 1

        replay = api.resolve_staged_appfolio_bill(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedBillResolutionIn(action="ACCEPT_RELATIONSHIP"),
            db=db,
            current_user=admin,
        )
        assert replay.id == row.id
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == row.id,
            AuditLog.action == "appfolio_bill_resolution_changed",
        ).count() == audit_count

        assert db.query(Bill).count() == before_bills
        assert db.query(GLTransaction).count() == before_gl
        assert db.query(Charge).count() == before_charges
        assert db.query(Vendor).count() == before_vendors
    finally:
        db.close()
        engine.dispose()


def test_appfolio_bills_missing_identity_or_unresolved_vendor_may_only_skip_or_fail_closed():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Bills Resolve Guard Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="bills-resolve-guard",
            ),
            db=db,
            current_user=admin,
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "bills-resolve-guard.csv",
                    (
                        "Bill ID,VendorId,PropertyId,DueDate,TotalAmount\n"
                        ",VENDOR-NO-MAP,,2026-08-02,100.00\n"
                        "BILL-NO-VENDOR-MAP,VENDOR-NO-MAP,,2026-08-03,110.00\n"
                    ).encode(),
                ),
                resource="BILLS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()

        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_bill(
                run.id,
                upload.id,
                rows[0].id,
                AppFolioStagedBillResolutionIn(action="ACCEPT_RELATIONSHIP"),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert "Bill ID" in exc.value.detail

        skipped = api.resolve_staged_appfolio_bill(
            run.id,
            upload.id,
            rows[0].id,
            AppFolioStagedBillResolutionIn(action="SKIP"),
            db=db,
            current_user=admin,
        )
        assert skipped.resolution_action == "SKIP"
        assert skipped.resolution_target_vendor_id is None
        assert skipped.resolution_target_id is None

        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_bill(
                run.id,
                upload.id,
                rows[1].id,
                AppFolioStagedBillResolutionIn(action="ACCEPT_RELATIONSHIP"),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert "durable VENDOR mapping" in exc.value.detail
        assert db.query(Bill).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_bills_acceptance_revalidates_optional_property_and_active_vendor_scope():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_DEV)
        org = _org(db, name="Bills Resolve Scope Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="bills-resolve-scope",
            ),
            db=db,
            current_user=admin,
        )
        prop, vendor = _seed_bill_source_mappings(db, run=run, org=org)
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "bills-resolve-scope.csv",
                    (
                        "Bill ID,VendorId,PropertyId,DueDate,TotalAmount\n"
                        "BILL-SCOPE-1,VENDOR-BILL-1,PROP-BILL-1,2026-08-04,120.00\n"
                    ).encode(),
                ),
                resource="BILLS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()

        property_mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "PROPERTIES",
            PlatformMigrationItem.source_id == "PROP-BILL-1",
        ).one()
        property_mapping.target_entity = "UNIT"
        db.commit()
        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_bill(
                run.id,
                upload.id,
                row.id,
                AppFolioStagedBillResolutionIn(action="ACCEPT_RELATIONSHIP"),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert "durable PROPERTY mapping" in exc.value.detail

        property_mapping.target_entity = "PROPERTY"
        vendor.is_active = False
        db.commit()
        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_bill(
                run.id,
                upload.id,
                row.id,
                AppFolioStagedBillResolutionIn(action="ACCEPT_RELATIONSHIP"),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert "Vendor is no longer active" in exc.value.detail
        assert db.get(Property, prop.id).is_active is True
        assert db.query(Bill).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Bills staged dry run
# ---------------------------------------------------------------------------

def test_appfolio_bills_dry_run_is_exact_replay_safe_and_non_mutating():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Bills Dry Run Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="bills-dry-run",
            ),
            db=db,
            current_user=admin,
        )
        prop, vendor = _seed_bill_source_mappings(db, run=run, org=org)
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "bills-dry-run.csv",
                    (
                        "Bill ID,VendorId,PropertyId,DueDate,InvoiceDate,PostingDate,"
                        "Reference,Remarks,TotalAmount,ApprovalStatus,CheckMemo,"
                        "AccountNumber,ManagementCompanyAsPayee,WorkOrderId,LastUpdatedAt\n"
                        "BILL-DRY-1,VENDOR-BILL-1,PROP-BILL-1,2026-09-01,2026-08-15,"
                        "2026-08-16,INV-DRY,source remark,325.50,Pending,repair memo,"
                        "5000,false,WO-DRY-1,2026-08-20T12:00:00Z\n"
                    ).encode(),
                ),
                resource="BILLS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        api.resolve_staged_appfolio_bill(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedBillResolutionIn(action="ACCEPT_RELATIONSHIP"),
            db=db,
            current_user=admin,
        )

        before_bills = db.query(Bill).count()
        before_gl = db.query(GLTransaction).count()
        before_charges = db.query(Charge).count()
        before_vendors = db.query(Vendor).count()

        first = api.dry_run_staged_appfolio_bills(
            run.id, upload.id, db=db, current_user=admin
        )
        assert first.replayed is False
        assert first.total == 1
        assert first.importable == 1
        assert first.invalid == 0
        assert len(first.rows) == 1
        preview = first.rows[0]
        assert preview.source_id == "BILL-DRY-1"
        assert preview.source_evidence["source_vendor_id"] == "VENDOR-BILL-1"
        assert preview.source_evidence["source_property_id"] == "PROP-BILL-1"
        assert preview.source_evidence["due_date"] == "2026-09-01"
        assert preview.source_evidence["invoice_date"] == "2026-08-15"
        assert preview.source_evidence["posting_date"] == "2026-08-16"
        assert preview.source_evidence["reference"] == "INV-DRY"
        assert preview.source_evidence["remarks"] == "source remark"
        assert preview.source_evidence["total_amount"] == "325.50"
        assert preview.source_evidence["approval_status"] == "Pending"
        assert preview.source_evidence["check_memo"] == "repair memo"
        assert preview.source_evidence["account_number"] == "5000"
        assert preview.source_evidence["management_company_as_payee"] == "false"
        assert preview.source_evidence["source_work_order_id"] == "WO-DRY-1"
        assert preview.source_evidence["last_updated_at"] == "2026-08-20T12:00:00Z"
        assert preview.resolved_targets == {
            "vendor_id": vendor.id,
            "property_id": prop.id,
        }

        db.refresh(run)
        assert run.last_dry_run_fingerprint == first.fingerprint
        assert run.last_dry_run_summary["resource"] == "BILLS"
        assert run.last_dry_run_summary["review_only"] is True
        assert run.last_dry_run_summary["bill_mutation"] is False
        assert run.last_dry_run_summary["payment_state_inferred"] is False
        assert run.last_dry_run_summary["gl_allocation_inferred"] is False
        assert run.last_dry_run_summary["work_order_linkage_inferred"] is False

        audit_count = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_staged_bills_dry_run",
        ).count()
        assert audit_count == 1

        replay = api.dry_run_staged_appfolio_bills(
            run.id, upload.id, db=db, current_user=admin
        )
        assert replay.replayed is True
        assert replay.fingerprint == first.fingerprint
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_staged_bills_dry_run",
        ).count() == audit_count

        assert db.query(Bill).count() == before_bills
        assert db.query(GLTransaction).count() == before_gl
        assert db.query(Charge).count() == before_charges
        assert db.query(Vendor).count() == before_vendors
    finally:
        db.close()
        engine.dispose()


def test_appfolio_bills_dry_run_revalidates_relationship_mapping_fingerprints():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Bills Dry Run Stale Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="bills-dry-run-stale",
            ),
            db=db,
            current_user=admin,
        )
        prop, vendor = _seed_bill_source_mappings(db, run=run, org=org)
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "bills-dry-run-stale.csv",
                    (
                        "Bill ID,VendorId,PropertyId,DueDate,TotalAmount\n"
                        "BILL-DRY-STALE,VENDOR-BILL-1,PROP-BILL-1,2026-09-02,100.00\n"
                    ).encode(),
                ),
                resource="BILLS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        api.resolve_staged_appfolio_bill(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedBillResolutionIn(action="ACCEPT_RELATIONSHIP"),
            db=db,
            current_user=admin,
        )
        first = api.dry_run_staged_appfolio_bills(
            run.id, upload.id, db=db, current_user=admin
        )

        vendor_mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "VENDORS",
            PlatformMigrationItem.source_id == "VENDOR-BILL-1",
        ).one()
        vendor_mapping.source_fingerprint = "changed-vendor-source-fingerprint"
        db.commit()

        changed = api.dry_run_staged_appfolio_bills(
            run.id, upload.id, db=db, current_user=admin
        )
        assert changed.fingerprint != first.fingerprint
        assert changed.replayed is False
        assert changed.rows[0].resolved_targets["vendor_id"] == vendor.id
        assert changed.rows[0].resolved_targets["property_id"] == prop.id

        property_mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "PROPERTIES",
            PlatformMigrationItem.source_id == "PROP-BILL-1",
        ).one()
        property_mapping.target_entity = "UNIT"
        db.commit()
        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_bills(
                run.id, upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "Property relationship is stale" in exc.value.detail

        assert db.query(Bill).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_bills_dry_run_blocks_unresolved_rows_and_empty_accepted_set():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_DEV)
        org = _org(db, name="Bills Dry Run Guard Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="bills-dry-run-guard",
            ),
            db=db,
            current_user=admin,
        )
        _seed_bill_source_mappings(db, run=run, org=org)
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "bills-dry-run-guard.csv",
                    (
                        "Bill ID,VendorId,PropertyId,DueDate,TotalAmount\n"
                        "BILL-DRY-GUARD,VENDOR-BILL-1,PROP-BILL-1,2026-09-03,75.00\n"
                    ).encode(),
                ),
                resource="BILLS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()

        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_bills(
                run.id, upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "unresolved staged rows" in exc.value.detail

        api.resolve_staged_appfolio_bill(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedBillResolutionIn(action="SKIP"),
            db=db,
            current_user=admin,
        )
        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_bills(
                run.id, upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "at least one accepted staged row" in exc.value.detail
        assert db.query(Bill).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Bills staged reconciliation controls
# ---------------------------------------------------------------------------

def test_appfolio_bills_reconciliation_matches_identity_total_and_stale_preview_fails_closed():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Bills Reconciliation Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="bills-reconciliation",
            ),
            db=db,
            current_user=admin,
        )
        _prop, _vendor = _seed_bill_source_mappings(db, run=run, org=org)
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "bills-reconciliation.csv",
                    (
                        "Bill ID,VendorId,PropertyId,DueDate,TotalAmount,Reference\n"
                        "BILL-REC-2,VENDOR-BILL-1,PROP-BILL-1,2026-10-02,200.25,INV-REC-2\n"
                        "BILL-REC-1,VENDOR-BILL-1,PROP-BILL-1,2026-10-01,99.75,INV-REC-1\n"
                    ).encode(),
                ),
                resource="BILLS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()
        for row in rows:
            api.resolve_staged_appfolio_bill(
                run.id,
                upload.id,
                row.id,
                AppFolioStagedBillResolutionIn(action="ACCEPT_RELATIONSHIP"),
                db=db,
                current_user=admin,
            )

        dry = api.dry_run_staged_appfolio_bills(
            run.id, upload.id, db=db, current_user=admin
        )
        before_bills = db.query(Bill).count()
        before_gl = db.query(GLTransaction).count()
        before_charges = db.query(Charge).count()

        response = Response()
        reconciliation = api.get_staged_appfolio_bills_reconciliation(
            run.id,
            upload.id,
            response=response,
            db=db,
            current_user=admin,
        )
        assert response.headers["cache-control"] == "no-store"
        assert reconciliation.dry_run_fingerprint == dry.fingerprint
        assert reconciliation.source_bill_count == 2
        assert reconciliation.preview_bill_count == 2
        assert reconciliation.source_total_amount == "300.00"
        assert reconciliation.preview_total_amount == "300.00"
        assert reconciliation.bill_count_match is True
        assert reconciliation.bill_identity_match is True
        assert reconciliation.total_amount_match is True
        assert reconciliation.reconciled is True
        assert reconciliation.accounting_complete is False
        assert "bill_line_gl_allocations" in reconciliation.unverified_controls
        assert db.query(Bill).count() == before_bills
        assert db.query(GLTransaction).count() == before_gl
        assert db.query(Charge).count() == before_charges

        stale = db.get(PlatformMigrationStagedRow, rows[0].id)
        stale.normalized_data = {
            **dict(stale.normalized_data or {}),
            "total_amount": "201.25",
        }
        db.commit()
        with pytest.raises(HTTPException) as exc:
            api.get_staged_appfolio_bills_reconciliation(
                run.id,
                upload.id,
                response=Response(),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert "exact latest Bills dry-run fingerprint" in exc.value.detail
        assert db.query(Bill).count() == before_bills
        assert db.query(GLTransaction).count() == before_gl
        assert db.query(Charge).count() == before_charges
    finally:
        db.close()
        engine.dispose()


def test_appfolio_bills_reconciliation_route_is_exposed_and_sales_is_denied():
    from app.main import app

    assert (
        "/api/platform/migrations/appfolio/runs/{run_id}/uploads/{upload_id}/bills/reconciliation"
        in set(app.openapi()["paths"])
    )

    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        sales = _platform_user(db, PlatformUserRole.PLATFORM_SALES)
        org = _org(db, name="Bills Reconciliation Auth Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="bills-reconciliation-auth",
            ),
            db=db,
            current_user=admin,
        )
        with pytest.raises(HTTPException) as exc:
            api.get_staged_appfolio_bills_reconciliation(
                run.id,
                999999,
                response=Response(),
                db=db,
                current_user=sales,
            )
        assert exc.value.status_code == 403
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Charges / Receivables CSV/XLSX ingestion + staging
# ---------------------------------------------------------------------------

def test_appfolio_charges_csv_stages_source_evidence_without_receivable_mutation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Charges Stage Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="charges-stage",
            ),
            db=db,
            current_user=admin,
        )
        _prop, _unit, account = _seed_general_ledger_source_mappings(
            db, run=run, org=org
        )

        before_charges = db.query(Charge).count()
        before_gl = db.query(GLTransaction).count()

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "charges.csv",
                    (
                        "Id,AmountDue,ChargedOn,Description,GlAccountId,OccupancyId\n"
                        "CHARGE-1,125.50,2026-09-10,Repair charge,GL-4100,OCC-100\n"
                    ).encode(),
                ),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert upload.detected_resource == "CHARGES"
        assert upload.validation_summary["charge_rows"] == 1
        assert upload.validation_summary["missing_charge_source_ids"] == 0
        assert upload.validation_summary["unresolved_charge_gl_accounts"] == 0
        assert upload.validation_summary["unresolved_charge_occupancies"] == 1
        assert upload.validation_summary["charge_mutation"] is False
        assert upload.validation_summary["payment_state_inferred"] is False
        assert upload.validation_summary["tenant_liability_inferred"] is False

        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        assert row.source_id == "CHARGE-1"
        assert row.disposition == "REVIEW"
        assert row.normalized_data == {
            "source_id": "CHARGE-1",
            "amount_due": "125.50",
            "charged_on": "2026-09-10",
            "description": "Repair charge",
            "source_gl_account_id": "GL-4100",
            "source_occupancy_id": "OCC-100",
        }
        assert any(
            "GL Account relationship is durably mapped" in warning
            for warning in row.warnings
        )
        assert any(
            "OccupancyId is preserved as source evidence only" in warning
            for warning in row.warnings
        )
        assert any(
            "AmountDue is preserved as the source current outstanding amount"
            in warning
            for warning in row.warnings
        )
        assert account.is_active is True
        assert db.query(Charge).count() == before_charges
        assert db.query(GLTransaction).count() == before_gl
    finally:
        db.close()
        engine.dispose()


def test_appfolio_charges_explicit_missing_identity_and_duplicate_ids_fail_safe():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Charges Guard Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="charges-guard",
            ),
            db=db,
            current_user=admin,
        )

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "charges-guard.csv",
                    (
                        "Id,AmountDue,ChargedOn,Description,GlAccountId,OccupancyId\n"
                        ",50.00,2026-09-11,Missing ID,GL-NO-MAP,OCC-1\n"
                        "CHARGE-DUP,60.00,2026-09-12,First duplicate,GL-NO-MAP,OCC-2\n"
                        "CHARGE-DUP,70.00,2026-09-13,Second duplicate,GL-NO-MAP,OCC-3\n"
                    ).encode(),
                ),
                resource="CHARGES",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()

        assert upload.detected_resource == "CHARGES"
        assert upload.validation_summary["charge_rows"] == 3
        assert upload.validation_summary["missing_charge_source_ids"] == 1
        assert upload.validation_summary["unresolved_charge_gl_accounts"] == 3
        assert upload.validation_summary["unresolved_charge_occupancies"] == 2
        assert upload.validation_summary["duplicates"] == 1

        assert rows[0].source_id is None
        assert rows[0].disposition == "REVIEW"
        assert any("Charge ID was not supplied" in warning for warning in rows[0].warnings)
        assert any("GL Account relationship is unresolved" in warning for warning in rows[0].warnings)
        assert rows[1].disposition == "REVIEW"
        assert rows[2].disposition == "INVALID"
        assert any("Duplicate AppFolio Charge ID" in error for error in rows[2].errors)
        assert db.query(Charge).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_charges_xlsx_multisheet_autodetects_and_replays_without_mutation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_DEV)
        org = _org(db, name="Charges XLSX Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="charges-xlsx",
            ),
            db=db,
            current_user=admin,
        )
        _seed_general_ledger_source_mappings(db, run=run, org=org)

        workbook = Workbook()
        ws = workbook.active
        ws.title = "Charges"
        ws.append([
            "Id", "AmountDue", "ChargedOn", "Description", "GlAccountId",
            "OccupancyId",
        ])
        ws.append([
            "CHARGE-XLSX-1", "88.25", "2026-09-14", "Source charge",
            "GL-4100", "OCC-XLSX-1",
        ])
        notes = workbook.create_sheet("Notes")
        notes.append(["Comment"])
        notes.append(["not a supported migration report"])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        before_charges = db.query(Charge).count()
        before_gl = db.query(GLTransaction).count()
        first = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("charges.xlsx", buffer.getvalue()),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert first.detected_resource == "CHARGES"
        assert first.sheet_name == "Charges"
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == first.id
        ).one()
        assert row.source_id == "CHARGE-XLSX-1"
        assert row.normalized_data["amount_due"] == "88.25"
        assert row.normalized_data["source_occupancy_id"] == "OCC-XLSX-1"

        replay = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("charges-copy.xlsx", buffer.getvalue()),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert replay.replayed is True
        assert replay.id == first.id
        assert db.query(Charge).count() == before_charges
        assert db.query(GLTransaction).count() == before_gl
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Charges staged relationship reconciliation
# ---------------------------------------------------------------------------

def test_appfolio_charge_resolution_accepts_only_durable_gl_mapping_and_is_idempotent():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Charges Resolve Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="charges-resolve",
            ),
            db=db,
            current_user=admin,
        )
        _prop, _unit, account = _seed_general_ledger_source_mappings(
            db, run=run, org=org
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "charges-resolve.csv",
                    (
                        "Id,AmountDue,ChargedOn,Description,GlAccountId,OccupancyId\n"
                        "CHARGE-R1,125.50,2026-09-10,Repair charge,GL-4100,OCC-100\n"
                    ).encode(),
                ),
                resource="CHARGES",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        run.last_dry_run_fingerprint = "a" * 64
        run.last_dry_run_summary = {"resource": "CHARGES"}
        db.commit()

        before_charges = db.query(Charge).count()
        before_gl = db.query(GLTransaction).count()
        before_audit = db.query(AuditLog).filter(
            AuditLog.action == "appfolio_charge_resolution_changed"
        ).count()

        resolved = api.resolve_staged_appfolio_charge(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedChargeResolutionIn(action="ACCEPT_RELATIONSHIP"),
            db=db,
            current_user=admin,
        )
        assert resolved.resolution_action == "ACCEPT_RELATIONSHIP"
        assert resolved.resolution_target_gl_account_id == account.id
        assert resolved.resolution_target_id is None
        assert resolved.resolution_target_unit_id is None
        assert resolved.resolution_target_tenant_user_id is None
        db.refresh(run)
        assert run.last_dry_run_fingerprint is None
        assert run.last_dry_run_summary is None
        assert run.status == "STAGED"

        replay = api.resolve_staged_appfolio_charge(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedChargeResolutionIn(action="ACCEPT_RELATIONSHIP"),
            db=db,
            current_user=admin,
        )
        assert replay.id == resolved.id
        assert db.query(AuditLog).filter(
            AuditLog.action == "appfolio_charge_resolution_changed"
        ).count() == before_audit + 1
        audit = db.query(AuditLog).filter(
            AuditLog.action == "appfolio_charge_resolution_changed"
        ).order_by(AuditLog.id.desc()).first()
        audit_value = json.loads(audit.new_value)
        assert audit_value["accepted_gl_relationship_only"] is True
        assert audit_value["occupancy_mapping_inferred"] is False
        assert audit_value["tenant_liability_inferred"] is False
        assert audit_value["payment_state_inferred"] is False
        assert audit_value["target_charge_mutation"] is False
        assert db.query(Charge).count() == before_charges
        assert db.query(GLTransaction).count() == before_gl
    finally:
        db.close()
        engine.dispose()


def test_appfolio_charge_resolution_rejects_unresolved_or_stale_gl_relationship():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Charges Resolve Guard Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="charges-resolve-guard",
            ),
            db=db,
            current_user=admin,
        )
        _prop, _unit, account = _seed_general_ledger_source_mappings(
            db, run=run, org=org
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "charges-resolve-guard.csv",
                    (
                        "Id,AmountDue,ChargedOn,Description,GlAccountId,OccupancyId\n"
                        "CHARGE-OK,10.00,2026-09-10,Known account,GL-4100,OCC-1\n"
                        "CHARGE-BAD,20.00,2026-09-11,Unknown account,GL-NO-MAP,OCC-2\n"
                    ).encode(),
                ),
                resource="CHARGES",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()

        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_charge(
                run.id,
                upload.id,
                rows[1].id,
                AppFolioStagedChargeResolutionIn(action="ACCEPT_RELATIONSHIP"),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409

        account.is_active = False
        db.commit()
        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_charge(
                run.id,
                upload.id,
                rows[0].id,
                AppFolioStagedChargeResolutionIn(action="ACCEPT_RELATIONSHIP"),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert db.query(Charge).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_charge_resolution_missing_identity_may_only_skip_and_invalid_cannot_resolve():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_DEV)
        org = _org(db, name="Charges Resolve Identity Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="charges-resolve-identity",
            ),
            db=db,
            current_user=admin,
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "charges-resolve-identity.csv",
                    (
                        "Id,AmountDue,ChargedOn,Description,GlAccountId,OccupancyId\n"
                        ",10.00,2026-09-10,Missing id,GL-1,OCC-1\n"
                        "CHARGE-DUP,20.00,2026-09-11,Duplicate one,GL-1,OCC-2\n"
                        "CHARGE-DUP,30.00,2026-09-12,Duplicate two,GL-1,OCC-3\n"
                    ).encode(),
                ),
                resource="CHARGES",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()

        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_charge(
                run.id,
                upload.id,
                rows[0].id,
                AppFolioStagedChargeResolutionIn(action="ACCEPT_RELATIONSHIP"),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409

        skipped = api.resolve_staged_appfolio_charge(
            run.id,
            upload.id,
            rows[0].id,
            AppFolioStagedChargeResolutionIn(action="SKIP"),
            db=db,
            current_user=admin,
        )
        assert skipped.resolution_action == "SKIP"
        assert skipped.resolution_target_gl_account_id is None

        assert rows[2].disposition == "INVALID"
        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_charge(
                run.id,
                upload.id,
                rows[2].id,
                AppFolioStagedChargeResolutionIn(action="SKIP"),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert db.query(Charge).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Charges staged dry run
# ---------------------------------------------------------------------------

def test_appfolio_charges_dry_run_binds_review_and_gl_mapping_and_replays():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Charges Dry Run Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="charges-dry-run",
            ),
            db=db,
            current_user=admin,
        )
        _prop, _unit, account = _seed_general_ledger_source_mappings(
            db, run=run, org=org
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "charges-dry-run.csv",
                    (
                        "Id,AmountDue,ChargedOn,Description,GlAccountId,OccupancyId\n"
                        "CHARGE-DRY-1,125.50,2026-09-10,Repair charge,GL-4100,OCC-100\n"
                    ).encode(),
                ),
                resource="CHARGES",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        api.resolve_staged_appfolio_charge(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedChargeResolutionIn(action="ACCEPT_RELATIONSHIP"),
            db=db,
            current_user=admin,
        )

        before_charges = db.query(Charge).count()
        before_gl = db.query(GLTransaction).count()
        before_audit = db.query(AuditLog).filter(
            AuditLog.action == "appfolio_staged_charges_dry_run"
        ).count()

        first = api.dry_run_staged_appfolio_charges(
            run.id, upload.id, db=db, current_user=admin
        )
        assert len(first.fingerprint) == 64
        assert first.replayed is False
        assert first.total == 1 and first.importable == 1 and first.invalid == 0
        assert first.rows[0].source_id == "CHARGE-DRY-1"
        assert first.rows[0].source_evidence == {
            "amount_due": "125.50",
            "charged_on": "2026-09-10",
            "description": "Repair charge",
            "source_gl_account_id": "GL-4100",
            "source_occupancy_id": "OCC-100",
        }
        assert first.rows[0].resolved_targets == {"gl_account_id": account.id}
        db.refresh(run)
        assert run.status == "DRY_RUN_READY"
        assert run.last_dry_run_fingerprint == first.fingerprint
        assert run.last_dry_run_summary["resource"] == "CHARGES"
        assert run.last_dry_run_summary["target_mutation"] is False
        assert run.last_dry_run_summary["occupancy_target_inferred"] is False
        assert run.last_dry_run_summary["tenant_liability_inferred"] is False
        assert run.last_dry_run_summary["payment_state_inferred"] is False

        replay = api.dry_run_staged_appfolio_charges(
            run.id, upload.id, db=db, current_user=admin
        )
        assert replay.replayed is True
        assert replay.fingerprint == first.fingerprint
        assert db.query(AuditLog).filter(
            AuditLog.action == "appfolio_staged_charges_dry_run"
        ).count() == before_audit + 1
        assert db.query(Charge).count() == before_charges
        assert db.query(GLTransaction).count() == before_gl
    finally:
        db.close()
        engine.dispose()


def test_appfolio_charges_dry_run_changes_when_mapping_fingerprint_changes_and_rejects_stale_target():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Charges Dry Run Mapping Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="charges-dry-map",
            ),
            db=db,
            current_user=admin,
        )
        _prop, _unit, account = _seed_general_ledger_source_mappings(
            db, run=run, org=org
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "charges-dry-map.csv",
                    (
                        "Id,AmountDue,ChargedOn,Description,GlAccountId,OccupancyId\n"
                        "CHARGE-MAP-1,80.00,2026-09-12,Source charge,GL-4100,OCC-1\n"
                    ).encode(),
                ),
                resource="CHARGES",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        api.resolve_staged_appfolio_charge(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedChargeResolutionIn(action="ACCEPT_RELATIONSHIP"),
            db=db,
            current_user=admin,
        )
        first = api.dry_run_staged_appfolio_charges(
            run.id, upload.id, db=db, current_user=admin
        )

        mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "GL_ACCOUNTS",
            PlatformMigrationItem.source_id == "GL-4100",
        ).one()
        mapping.source_fingerprint = "b" * 64
        db.commit()

        changed = api.dry_run_staged_appfolio_charges(
            run.id, upload.id, db=db, current_user=admin
        )
        assert changed.replayed is False
        assert changed.fingerprint != first.fingerprint

        account.is_active = False
        db.commit()
        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_charges(
                run.id, upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert db.query(Charge).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_charges_dry_run_blocks_unresolved_rows_and_empty_accepted_set():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_DEV)
        org = _org(db, name="Charges Dry Run Guard Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="charges-dry-guard",
            ),
            db=db,
            current_user=admin,
        )
        _seed_general_ledger_source_mappings(db, run=run, org=org)
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "charges-dry-guard.csv",
                    (
                        "Id,AmountDue,ChargedOn,Description,GlAccountId,OccupancyId\n"
                        "CHARGE-GUARD-1,50.00,2026-09-10,Guard charge,GL-4100,OCC-1\n"
                    ).encode(),
                ),
                resource="CHARGES",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()

        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_charges(
                run.id, upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409

        api.resolve_staged_appfolio_charge(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedChargeResolutionIn(action="SKIP"),
            db=db,
            current_user=admin,
        )
        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_charges(
                run.id, upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert db.query(Charge).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()

# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Work Orders ingestion + staging
# ---------------------------------------------------------------------------

def _seed_work_order_source_mappings(db, *, run, org):
    prop = Property(
        organization_id=org.id,
        name="Work Order Stage Property",
        address_line1="500 Service Ave",
        city="Cleveland",
        state="OH",
        zip_code="44113",
        is_active=True,
    )
    db.add(prop)
    db.flush()
    unit = Unit(
        property_id=prop.id,
        unit_number="WO-101",
        bedrooms=1,
        bathrooms=1,
        monthly_rent=1000,
        is_active=True,
    )
    vendor = Vendor(
        organization_id=org.id,
        company_name="Work Order Plumbing",
        is_active=True,
    )
    db.add_all([unit, vendor])
    db.flush()
    for resource, source_id, target_entity, target_id in (
        ("PROPERTIES", "PROP-WO-1", "PROPERTY", prop.id),
        ("UNITS", "UNIT-WO-1", "UNIT", unit.id),
        ("VENDORS", "VENDOR-WO-1", "VENDOR", vendor.id),
    ):
        db.add(
            PlatformMigrationItem(
                run_id=run.id,
                organization_id=org.id,
                provider="APPFOLIO",
                resource=resource,
                source_id=source_id,
                target_entity=target_entity,
                target_id=target_id,
                source_fingerprint=f"seed-{resource}-{source_id}",
            )
        )
    db.commit()
    return prop, unit, vendor


def test_appfolio_work_orders_csv_staging_preserves_source_evidence_and_replays_without_mutation():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Work Orders Stage Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="work-orders-stage",
            ),
            db=db,
            current_user=admin,
        )
        _seed_work_order_source_mappings(db, run=run, org=org)
        before_work_orders = db.query(WorkOrder).count()
        before_bills = db.query(Bill).count()
        before_charges = db.query(Charge).count()
        before_gl = db.query(GLTransaction).count()

        content = (
            "Id,PropertyId,UnitId,Status,JobDescription,AssignedUsers,CanceledOn,"
            "CompletedOn,PermissionToEnter,Priority,ScheduledStart,ScheduledEnd,"
            "VendorId,VendorTrade\n"
            "WO-100,PROP-WO-1,UNIT-WO-1,Open,Leaking faucet,USER-44,,,"
            "true,Urgent,2026-10-04T09:00:00,2026-10-04T10:00:00,"
            "VENDOR-WO-1,Plumbing\n"
        ).encode()

        first = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("work-orders.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert first.detected_resource == "WORK_ORDERS"
        assert first.status == "REVIEW_REQUIRED"
        assert first.validation_summary["work_order_rows"] == 1
        assert first.validation_summary["missing_work_order_source_ids"] == 0
        assert first.validation_summary["unresolved_work_order_properties"] == 0
        assert first.validation_summary["unresolved_work_order_units"] == 0
        assert first.validation_summary["unresolved_work_order_vendors"] == 0
        assert first.validation_summary["work_order_mutation"] is False
        assert first.validation_summary["accounting_history_mutation"] is False

        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == first.id
        ).one()
        assert row.resource == "WORK_ORDERS"
        assert row.source_id == "WO-100"
        assert row.disposition == "REVIEW"
        assert row.normalized_data == {
            "source_id": "WO-100",
            "source_property_id": "PROP-WO-1",
            "source_unit_id": "UNIT-WO-1",
            "status": "Open",
            "job_description": "Leaking faucet",
            "assigned_users": "USER-44",
            "canceled_on": None,
            "completed_on": None,
            "permission_to_enter": "true",
            "priority": "Urgent",
            "scheduled_start": "2026-10-04T09:00:00",
            "scheduled_end": "2026-10-04T10:00:00",
            "source_vendor_id": "VENDOR-WO-1",
            "vendor_trade": "Plumbing",
        }
        assert any("source evidence only" in warning for warning in row.warnings)
        assert any("creates or updates no WorkOrder" in warning for warning in row.warnings)

        replay = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("work-orders-copy.csv", content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert replay.replayed is True
        assert replay.id == first.id
        assert db.query(WorkOrder).count() == before_work_orders
        assert db.query(Bill).count() == before_bills
        assert db.query(Charge).count() == before_charges
        assert db.query(GLTransaction).count() == before_gl
    finally:
        db.close()
        engine.dispose()


def test_appfolio_work_orders_explicit_missing_id_and_duplicate_ids_fail_closed_as_required():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Work Orders Identity Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="work-orders-identity",
            ),
            db=db,
            current_user=admin,
        )

        content = (
            "Id,PropertyId,Status,JobDescription,Priority\n"
            ",PROP-MISSING,Open,Missing stable id,Normal\n"
            "WO-DUP,PROP-MISSING,Open,Duplicate one,High\n"
            "WO-DUP,PROP-MISSING,Closed,Duplicate two,Low\n"
        ).encode()
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("work-orders-identity.csv", content),
                resource="WORK_ORDERS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()

        assert rows[0].disposition == "REVIEW"
        assert rows[0].source_id is None
        assert any("Work Order ID was not supplied" in warning for warning in rows[0].warnings)
        assert rows[1].disposition == "REVIEW"
        assert rows[2].disposition == "INVALID"
        assert any("Duplicate AppFolio Work Order ID" in error for error in rows[2].errors)
        assert upload.validation_summary["duplicates"] == 1
        assert upload.validation_summary["invalid"] == 1
        assert upload.validation_summary["missing_work_order_source_ids"] == 1
        assert upload.validation_summary["unresolved_work_order_properties"] == 3
        assert db.query(WorkOrder).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_work_orders_xlsx_detects_verified_contract_and_rejects_contradictory_mappings():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_DEV)
        org = _org(db, name="Work Orders XLSX Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="work-orders-xlsx",
            ),
            db=db,
            current_user=admin,
        )
        prop, unit, vendor = _seed_work_order_source_mappings(db, run=run, org=org)

        second = Property(
            organization_id=org.id,
            name="Other Work Order Property",
            address_line1="700 Other Ave",
            city="Cleveland",
            state="OH",
            zip_code="44113",
            is_active=True,
        )
        db.add(second)
        db.flush()
        db.add(
            PlatformMigrationItem(
                run_id=run.id,
                organization_id=org.id,
                provider="APPFOLIO",
                resource="PROPERTIES",
                source_id="PROP-WO-2",
                target_entity="PROPERTY",
                target_id=second.id,
                source_fingerprint="seed-PROP-WO-2",
            )
        )
        db.commit()

        workbook = Workbook()
        ws = workbook.active
        ws.title = "Work Orders"
        ws.append([
            "Id", "PropertyId", "UnitId", "Status", "JobDescription",
            "AssignedUsers", "CanceledOn", "CompletedOn", "PermissionToEnter",
            "Priority", "ScheduledStart", "ScheduledEnd", "VendorId", "VendorTrade",
        ])
        ws.append([
            "WO-XLSX-1", "PROP-WO-1", "UNIT-WO-1", "Completed",
            "Replace disposal", "USER-9", None, "2026-10-02T14:00:00",
            True, "High", "2026-10-02T13:00:00", "2026-10-02T14:00:00",
            "VENDOR-WO-1", "Plumbing",
        ])
        buffer = BytesIO()
        workbook.save(buffer)
        workbook.close()

        first = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("work-orders.xlsx", buffer.getvalue()),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert first.detected_resource == "WORK_ORDERS"
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == first.id
        ).one()
        assert row.normalized_data["completed_on"] == "2026-10-02T14:00:00"
        assert row.normalized_data["permission_to_enter"] is True
        assert db.query(WorkOrder).count() == 0

        contradictory = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "work-orders-contradictory.csv",
                    (
                        "Id,PropertyId,UnitId,Status,JobDescription,VendorId\n"
                        "WO-BAD,PROP-WO-2,UNIT-WO-1,Open,Contradictory unit,"
                        "VENDOR-WO-1\n"
                    ).encode(),
                ),
                resource="WORK_ORDERS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        bad = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == contradictory.id
        ).one()
        assert bad.disposition == "INVALID"
        assert any(
            "resolve to different target Properties" in error
            for error in bad.errors
        )

        vendor.is_active = False
        db.commit()
        stale = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "work-orders-stale-vendor.csv",
                    (
                        "Id,PropertyId,UnitId,Status,JobDescription,VendorId\n"
                        "WO-STALE,PROP-WO-1,UNIT-WO-1,Open,Stale vendor,"
                        "VENDOR-WO-1\n"
                    ).encode(),
                ),
                resource="WORK_ORDERS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        stale_row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == stale.id
        ).one()
        assert stale_row.disposition == "INVALID"
        assert any(
            "active same-organization Vendor" in error
            for error in stale_row.errors
        )
        assert db.query(WorkOrder).count() == 0
    finally:
        db.close()
        engine.dispose()

# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Work Orders staged relationship reconciliation
# ---------------------------------------------------------------------------

def test_appfolio_work_order_resolution_accepts_mapped_relationships_and_is_idempotent():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Work Orders Resolve Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="work-orders-resolve",
            ),
            db=db,
            current_user=admin,
        )
        prop, unit, vendor = _seed_work_order_source_mappings(db, run=run, org=org)
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "work-orders-resolve.csv",
                    (
                        "Id,PropertyId,UnitId,Status,JobDescription,AssignedUsers,"
                        "Priority,VendorId,VendorTrade\n"
                        "WO-R1,PROP-WO-1,UNIT-WO-1,Open,Leaking sink,USER-9,"
                        "Urgent,VENDOR-WO-1,Plumbing\n"
                    ).encode(),
                ),
                resource="WORK_ORDERS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        run.last_dry_run_fingerprint = "a" * 64
        run.last_dry_run_summary = {"resource": "WORK_ORDERS"}
        db.commit()

        before_work_orders = db.query(WorkOrder).count()
        before_bills = db.query(Bill).count()
        before_charges = db.query(Charge).count()
        before_gl = db.query(GLTransaction).count()
        before_audit = db.query(AuditLog).filter(
            AuditLog.action == "appfolio_work_order_resolution_changed"
        ).count()

        resolved = api.resolve_staged_appfolio_work_order(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedWorkOrderResolutionIn(action="ACCEPT_RELATIONSHIP"),
            db=db,
            current_user=admin,
        )
        assert resolved.resolution_action == "ACCEPT_RELATIONSHIP"
        assert resolved.resolution_target_id == prop.id
        assert resolved.resolution_target_unit_id == unit.id
        assert resolved.resolution_target_vendor_id == vendor.id
        assert resolved.resolution_target_tenant_user_id is None
        assert resolved.resolution_target_gl_account_id is None
        db.refresh(run)
        assert run.last_dry_run_fingerprint is None
        assert run.last_dry_run_summary is None
        assert run.status == "STAGED"

        replay = api.resolve_staged_appfolio_work_order(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedWorkOrderResolutionIn(action="ACCEPT_RELATIONSHIP"),
            db=db,
            current_user=admin,
        )
        assert replay.id == resolved.id
        assert db.query(AuditLog).filter(
            AuditLog.action == "appfolio_work_order_resolution_changed"
        ).count() == before_audit + 1
        audit = db.query(AuditLog).filter(
            AuditLog.action == "appfolio_work_order_resolution_changed"
        ).order_by(AuditLog.id.desc()).first()
        audit_value = json.loads(audit.new_value)
        assert audit_value["accepted_for_later_commit"] is True
        assert audit_value["assigned_users_inferred"] is False
        assert audit_value["workflow_state_inferred"] is False
        assert audit_value["target_work_order_mutation"] is False
        assert audit_value["accounting_mutation"] is False
        assert db.query(WorkOrder).count() == before_work_orders
        assert db.query(Bill).count() == before_bills
        assert db.query(Charge).count() == before_charges
        assert db.query(GLTransaction).count() == before_gl
    finally:
        db.close()
        engine.dispose()


def test_appfolio_work_order_resolution_missing_identity_may_only_skip_and_invalid_cannot_resolve():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Work Orders Resolve Identity Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="work-orders-resolve-identity",
            ),
            db=db,
            current_user=admin,
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "work-orders-resolve-identity.csv",
                    (
                        "Id,PropertyId,Status,JobDescription\n"
                        ",PROP-X,Open,Missing source id\n"
                        "WO-DUP,PROP-X,Open,Duplicate one\n"
                        "WO-DUP,PROP-X,Closed,Duplicate two\n"
                    ).encode(),
                ),
                resource="WORK_ORDERS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).order_by(PlatformMigrationStagedRow.row_number).all()

        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_work_order(
                run.id,
                upload.id,
                rows[0].id,
                AppFolioStagedWorkOrderResolutionIn(action="ACCEPT_RELATIONSHIP"),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409

        skipped = api.resolve_staged_appfolio_work_order(
            run.id,
            upload.id,
            rows[0].id,
            AppFolioStagedWorkOrderResolutionIn(action="SKIP"),
            db=db,
            current_user=admin,
        )
        assert skipped.resolution_action == "SKIP"
        assert skipped.resolution_target_id is None
        assert skipped.resolution_target_unit_id is None
        assert skipped.resolution_target_vendor_id is None

        assert rows[2].disposition == "INVALID"
        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_work_order(
                run.id,
                upload.id,
                rows[2].id,
                AppFolioStagedWorkOrderResolutionIn(action="SKIP"),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        assert db.query(WorkOrder).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_work_order_resolution_revalidates_required_supplied_relationships():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_DEV)
        org = _org(db, name="Work Orders Resolve Guard Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="work-orders-resolve-guard",
            ),
            db=db,
            current_user=admin,
        )
        prop, unit, vendor = _seed_work_order_source_mappings(db, run=run, org=org)

        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "work-orders-unmapped-vendor.csv",
                    (
                        "Id,PropertyId,UnitId,Status,JobDescription,VendorId\n"
                        "WO-UNMAPPED,PROP-WO-1,UNIT-WO-1,Open,Unknown vendor,VENDOR-NO-MAP\n"
                    ).encode(),
                ),
                resource="WORK_ORDERS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        with pytest.raises(HTTPException) as exc:
            api.resolve_staged_appfolio_work_order(
                run.id,
                upload.id,
                row.id,
                AppFolioStagedWorkOrderResolutionIn(action="ACCEPT_RELATIONSHIP"),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409

        vendor.is_active = False
        db.commit()
        stale = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "work-orders-stale-resolve.csv",
                    (
                        "Id,PropertyId,UnitId,Status,JobDescription,VendorId\n"
                        "WO-STALE-RESOLVE,PROP-WO-1,UNIT-WO-1,Open,Stale vendor,VENDOR-WO-1\n"
                    ).encode(),
                ),
                resource="WORK_ORDERS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        stale_row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == stale.id
        ).one()
        assert stale_row.disposition == "INVALID"

        vendor.is_active = True
        second = Property(
            organization_id=org.id,
            name="Resolve Other Property",
            address_line1="910 Other Ave",
            city="Cleveland",
            state="OH",
            zip_code="44113",
            is_active=True,
        )
        db.add(second)
        db.flush()
        mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "PROPERTIES",
            PlatformMigrationItem.source_id == "PROP-WO-1",
        ).one()
        mapping.target_id = second.id
        db.commit()

        mismatch = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "work-orders-mismatch-resolve.csv",
                    (
                        "Id,PropertyId,UnitId,Status,JobDescription\n"
                        "WO-MISMATCH,PROP-WO-1,UNIT-WO-1,Open,Unit property mismatch\n"
                    ).encode(),
                ),
                resource="WORK_ORDERS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        mismatch_row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == mismatch.id
        ).one()
        assert mismatch_row.disposition == "INVALID"
        assert db.query(WorkOrder).count() == 0
    finally:
        db.close()
        engine.dispose()



# ---------------------------------------------------------------------------
# Phase 4.13 AppFolio Work Orders staged dry run
# ---------------------------------------------------------------------------

def test_appfolio_work_orders_dry_run_is_exact_replay_safe_and_non_mutating():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Work Orders Dry Run Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="work-orders-dry-run",
            ),
            db=db,
            current_user=admin,
        )
        prop, unit, vendor = _seed_work_order_source_mappings(db, run=run, org=org)
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "work-orders-dry-run.csv",
                    (
                        "Id,PropertyId,UnitId,Status,JobDescription,AssignedUsers,CanceledOn,"
                        "CompletedOn,PermissionToEnter,Priority,ScheduledStart,ScheduledEnd,"
                        "VendorId,VendorTrade\n"
                        "WO-DRY-1,PROP-WO-1,UNIT-WO-1,Open,Leaking faucet,USER-44,,,"
                        "true,Urgent,2026-10-04T09:00:00,2026-10-04T10:00:00,"
                        "VENDOR-WO-1,Plumbing\n"
                    ).encode(),
                ),
                resource="WORK_ORDERS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        api.resolve_staged_appfolio_work_order(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedWorkOrderResolutionIn(action="ACCEPT_RELATIONSHIP"),
            db=db,
            current_user=admin,
        )

        before_work_orders = db.query(WorkOrder).count()
        before_bills = db.query(Bill).count()
        before_charges = db.query(Charge).count()
        before_gl = db.query(GLTransaction).count()
        before_vendors = db.query(Vendor).count()

        first = api.dry_run_staged_appfolio_work_orders(
            run.id, upload.id, db=db, current_user=admin
        )
        assert first.replayed is False
        assert first.total == 1
        assert first.importable == 1
        assert first.invalid == 0
        assert len(first.rows) == 1
        preview = first.rows[0]
        assert preview.source_id == "WO-DRY-1"
        assert preview.source_evidence == {
            "source_property_id": "PROP-WO-1",
            "source_unit_id": "UNIT-WO-1",
            "source_vendor_id": "VENDOR-WO-1",
            "assigned_users": "USER-44",
            "status": "Open",
            "job_description": "Leaking faucet",
            "canceled_on": None,
            "completed_on": None,
            "permission_to_enter": "true",
            "priority": "Urgent",
            "scheduled_start": "2026-10-04T09:00:00",
            "scheduled_end": "2026-10-04T10:00:00",
            "vendor_trade": "Plumbing",
        }
        assert preview.resolved_targets == {
            "property_id": prop.id,
            "unit_id": unit.id,
            "vendor_id": vendor.id,
        }
        assert any("review-only" in warning for warning in preview.warnings)

        db.refresh(run)
        assert run.last_dry_run_fingerprint == first.fingerprint
        assert run.status == "DRY_RUN_READY"
        assert run.last_dry_run_summary["resource"] == "WORK_ORDERS"
        assert run.last_dry_run_summary["review_only"] is True
        assert run.last_dry_run_summary["work_order_mutation"] is False
        assert run.last_dry_run_summary["staff_assignment_inferred"] is False
        assert run.last_dry_run_summary["workflow_state_translated"] is False

        audit_count = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_staged_work_orders_dry_run",
        ).count()
        assert audit_count == 1

        replay = api.dry_run_staged_appfolio_work_orders(
            run.id, upload.id, db=db, current_user=admin
        )
        assert replay.replayed is True
        assert replay.fingerprint == first.fingerprint
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_run",
            AuditLog.entity_id == run.id,
            AuditLog.action == "appfolio_staged_work_orders_dry_run",
        ).count() == audit_count

        assert db.query(WorkOrder).count() == before_work_orders
        assert db.query(Bill).count() == before_bills
        assert db.query(Charge).count() == before_charges
        assert db.query(GLTransaction).count() == before_gl
        assert db.query(Vendor).count() == before_vendors
    finally:
        db.close()
        engine.dispose()


def test_appfolio_work_orders_dry_run_revalidates_mapping_fingerprints_and_scope():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Work Orders Dry Run Stale Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="work-orders-dry-run-stale",
            ),
            db=db,
            current_user=admin,
        )
        prop, unit, vendor = _seed_work_order_source_mappings(db, run=run, org=org)
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "work-orders-dry-run-stale.csv",
                    (
                        "Id,PropertyId,UnitId,Status,JobDescription,VendorId\n"
                        "WO-DRY-STALE,PROP-WO-1,UNIT-WO-1,Open,Stale checks,VENDOR-WO-1\n"
                    ).encode(),
                ),
                resource="WORK_ORDERS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()
        api.resolve_staged_appfolio_work_order(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedWorkOrderResolutionIn(action="ACCEPT_RELATIONSHIP"),
            db=db,
            current_user=admin,
        )
        first = api.dry_run_staged_appfolio_work_orders(
            run.id, upload.id, db=db, current_user=admin
        )

        vendor_mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "VENDORS",
            PlatformMigrationItem.source_id == "VENDOR-WO-1",
        ).one()
        vendor_mapping.source_fingerprint = "changed-work-order-vendor-fingerprint"
        db.commit()

        changed = api.dry_run_staged_appfolio_work_orders(
            run.id, upload.id, db=db, current_user=admin
        )
        assert changed.fingerprint != first.fingerprint
        assert changed.replayed is False
        assert changed.rows[0].resolved_targets["vendor_id"] == vendor.id

        property_mapping = db.query(PlatformMigrationItem).filter(
            PlatformMigrationItem.run_id == run.id,
            PlatformMigrationItem.resource == "PROPERTIES",
            PlatformMigrationItem.source_id == "PROP-WO-1",
        ).one()
        property_mapping.target_entity = "UNIT"
        db.commit()
        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_work_orders(
                run.id, upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "Property relationship is stale" in exc.value.detail

        property_mapping.target_entity = "PROPERTY"
        second_property = Property(
            organization_id=org.id,
            name="Work Order Dry Run Other Property",
            address_line1="777 Other",
            city="Cleveland",
            state="OH",
            zip_code="44113",
            is_active=True,
        )
        db.add(second_property)
        db.flush()
        property_mapping.target_id = second_property.id
        row.resolution_target_id = second_property.id
        db.commit()
        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_work_orders(
                run.id, upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "Unit no longer belongs" in exc.value.detail

        assert db.get(Property, prop.id).is_active is True
        assert db.get(Unit, unit.id).is_active is True
        assert db.query(WorkOrder).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


def test_appfolio_work_orders_dry_run_blocks_unresolved_and_empty_accepted_set():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_DEV)
        org = _org(db, name="Work Orders Dry Run Guard Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="work-orders-dry-run-guard",
            ),
            db=db,
            current_user=admin,
        )
        _seed_work_order_source_mappings(db, run=run, org=org)
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "work-orders-dry-run-guard.csv",
                    (
                        "Id,PropertyId,Status,JobDescription\n"
                        "WO-DRY-GUARD,PROP-WO-1,Open,Needs explicit review\n"
                    ).encode(),
                ),
                resource="WORK_ORDERS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == upload.id
        ).one()

        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_work_orders(
                run.id, upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "unresolved staged rows" in exc.value.detail

        api.resolve_staged_appfolio_work_order(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedWorkOrderResolutionIn(action="SKIP"),
            db=db,
            current_user=admin,
        )
        with pytest.raises(HTTPException) as exc:
            api.dry_run_staged_appfolio_work_orders(
                run.id, upload.id, db=db, current_user=admin
            )
        assert exc.value.status_code == 409
        assert "at least one accepted staged row" in exc.value.detail
        assert db.query(WorkOrder).count() == 0
        assert db.query(Bill).count() == 0
        assert db.query(Charge).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 migration source coverage/completeness checklist
# ---------------------------------------------------------------------------

def test_appfolio_migration_coverage_reports_durable_sources_partial_and_blocked_states():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Coverage Checklist Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="coverage-checklist",
            ),
            db=db,
            current_user=admin,
        )
        _prop, _unit, _vendor = _seed_work_order_source_mappings(
            db, run=run, org=org
        )
        upload = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file(
                    "coverage-work-orders.csv",
                    (
                        "Id,PropertyId,UnitId,Status,JobDescription,VendorId\n"
                        "WO-COVERAGE-1,PROP-WO-1,UNIT-WO-1,Open,Coverage test,VENDOR-WO-1\n"
                    ).encode(),
                ),
                resource="WORK_ORDERS",
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert upload.detected_resource == "WORK_ORDERS"

        response = Response()
        result = api.get_appfolio_migration_coverage(
            run.id,
            response=response,
            db=db,
            current_user=admin,
        )
        by_resource = {item.resource: item for item in result.items}

        assert response.headers["cache-control"] == "no-store"
        assert by_resource["PROPERTIES"].state == "SUPPLIED"
        assert by_resource["PROPERTIES"].upload_count == 0
        assert by_resource["PROPERTIES"].mapping_count == 1
        assert by_resource["UNITS"].state == "SUPPLIED"
        assert by_resource["VENDORS"].state == "SUPPLIED"
        assert by_resource["WORK_ORDERS"].state == "PARTIAL"
        assert by_resource["WORK_ORDERS"].upload_count == 1
        assert by_resource["WORK_ORDERS"].row_count == 1
        assert any(
            "target WorkOrder requires" in blocker
            for blocker in by_resource["WORK_ORDERS"].blockers
        )
        assert by_resource["TENANTS"].state == "MISSING"
        assert by_resource["SECURITY_DEPOSITS"].state == "BLOCKED"
        assert by_resource["DOCUMENTS"].state == "BLOCKED"
        assert result.accounting_complete is False
        assert result.partial_operational_migration_may_be_possible is True
        assert result.blocked_count == 2
        assert any(
            "Accounting migration is not complete" in blocker
            for blocker in result.blockers
        )
    finally:
        db.close()
        engine.dispose()


def test_appfolio_migration_coverage_is_repeatable_read_only_and_preserves_run_state():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_TECH)
        org = _org(db, name="Coverage Read Only Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="coverage-read-only",
            ),
            db=db,
            current_user=admin,
        )
        _seed_work_order_source_mappings(db, run=run, org=org)
        run.status = "DRY_RUN_READY"
        run.last_dry_run_fingerprint = "f" * 64
        run.last_dry_run_summary = {"resource": "WORK_ORDERS", "review_only": True}
        db.commit()

        before_uploads = db.query(PlatformMigrationUpload).count()
        before_mappings = db.query(PlatformMigrationItem).count()
        before_work_orders = db.query(WorkOrder).count()
        before_bills = db.query(Bill).count()
        before_charges = db.query(Charge).count()
        before_gl = db.query(GLTransaction).count()
        before_audit = db.query(AuditLog).count()

        first = api.get_appfolio_migration_coverage(
            run.id, response=Response(), db=db, current_user=admin
        )
        second = api.get_appfolio_migration_coverage(
            run.id, response=Response(), db=db, current_user=admin
        )
        assert first.model_dump() == second.model_dump()

        db.refresh(run)
        assert run.status == "DRY_RUN_READY"
        assert run.last_dry_run_fingerprint == "f" * 64
        assert run.last_dry_run_summary == {
            "resource": "WORK_ORDERS",
            "review_only": True,
        }
        assert db.query(PlatformMigrationUpload).count() == before_uploads
        assert db.query(PlatformMigrationItem).count() == before_mappings
        assert db.query(WorkOrder).count() == before_work_orders
        assert db.query(Bill).count() == before_bills
        assert db.query(Charge).count() == before_charges
        assert db.query(GLTransaction).count() == before_gl
        assert db.query(AuditLog).count() == before_audit
    finally:
        db.close()
        engine.dispose()


def test_appfolio_migration_coverage_enforces_platform_view_role_and_never_claims_accounting_complete():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        support = _platform_user(db, PlatformUserRole.PLATFORM_SUPPORT)
        sales = _platform_user(db, PlatformUserRole.PLATFORM_SALES)
        org = _org(db, name="Coverage Auth Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="coverage-auth",
            ),
            db=db,
            current_user=admin,
        )

        support_result = api.get_appfolio_migration_coverage(
            run.id, response=Response(), db=db, current_user=support
        )
        assert support_result.accounting_complete is False
        assert all(
            item.state in {"MISSING", "BLOCKED"}
            for item in support_result.items
        )
        assert support_result.partial_operational_migration_may_be_possible is False

        with pytest.raises(HTTPException) as exc:
            api.get_appfolio_migration_coverage(
                run.id, response=Response(), db=db, current_user=sales
            )
        assert exc.value.status_code == 403
    finally:
        db.close()
        engine.dispose()



def test_appfolio_staged_correction_preserves_source_and_invalidates_preview():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Correction Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="correction-run",
            ),
            db=db,
            current_user=admin,
        )
        upload = PlatformMigrationUpload(
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            filename="properties.csv",
            file_format="CSV",
            file_sha256="1" * 64,
            normalized_fingerprint="2" * 64,
            detected_resource="PROPERTIES",
            sheet_name="CSV",
            headers=["Id", "Name", "City"],
            column_mapping={"Id": "source_id", "Name": "name", "City": "city"},
            validation_summary={"valid": 1},
            status="STAGED",
            row_count=1,
            created_by_platform_user_id=admin.id,
        )
        db.add(upload)
        db.flush()
        row = PlatformMigrationStagedRow(
            upload_id=upload.id,
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            resource="PROPERTIES",
            row_number=2,
            source_id="AF-CORR-1",
            disposition="NEW",
            row_fingerprint="3" * 64,
            normalized_data={
                "source_id": "AF-CORR-1",
                "name": "Apt Bldg",
                "city": "Cleveland",
            },
            warnings=[],
            errors=[],
        )
        db.add(row)
        run.last_dry_run_fingerprint = "4" * 64
        run.last_dry_run_summary = {"resource": "PROPERTIES"}
        run.status = "DRY_RUN_READY"
        db.commit()
        db.refresh(row)
        db.refresh(upload)
        db.refresh(run)

        source_upload_fingerprint = upload.normalized_fingerprint
        source_row_fingerprint = row.row_fingerprint
        before_review = api._staged_review_fingerprint(upload, [row])
        before_properties = db.query(Property).count()
        before_charges = db.query(Charge).count()
        before_gl = db.query(GLTransaction).count()
        audit_before = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == row.id,
            AuditLog.action == "appfolio_staged_value_corrected",
        ).count()

        corrected = api.correct_appfolio_staged_row(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedRowCorrectionIn(
                field_name="name",
                corrected_value="Apartment Building",
            ),
            db=db,
            current_user=admin,
        )
        assert corrected.normalized_data["name"] == "Apartment Building"
        assert corrected.source_id == "AF-CORR-1"
        assert corrected.row_fingerprint == source_row_fingerprint
        assert upload.normalized_fingerprint == source_upload_fingerprint

        db.refresh(run)
        assert run.last_dry_run_fingerprint is None
        assert run.last_dry_run_summary is None
        assert run.status == "STAGED"
        after_review = api._staged_review_fingerprint(upload, [corrected])
        assert after_review != before_review

        audit_rows = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == row.id,
            AuditLog.action == "appfolio_staged_value_corrected",
        ).all()
        assert len(audit_rows) == audit_before + 1
        audit = audit_rows[-1]
        audit_old = json.loads(audit.old_value)
        audit_new = json.loads(audit.new_value)
        assert audit_old["field_name"] == "name"
        assert audit_old["value"] == "Apt Bldg"
        assert audit_new["value"] == "Apartment Building"
        assert audit_new["source_file_rewritten"] is False
        assert audit_new["customer_target_mutation"] is False
        assert audit_new["accounting_mutation"] is False

        replay = api.correct_appfolio_staged_row(
            run.id,
            upload.id,
            row.id,
            AppFolioStagedRowCorrectionIn(
                field_name="name",
                corrected_value="Apartment Building",
            ),
            db=db,
            current_user=admin,
        )
        assert replay.normalized_data["name"] == "Apartment Building"
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == row.id,
            AuditLog.action == "appfolio_staged_value_corrected",
        ).count() == audit_before + 1

        assert db.query(Property).count() == before_properties
        assert db.query(Charge).count() == before_charges
        assert db.query(GLTransaction).count() == before_gl
    finally:
        db.close()
        engine.dispose()


def test_appfolio_staged_correction_blocks_identity_invalid_rows_and_nonwrite_roles():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        support = _platform_user(db, PlatformUserRole.PLATFORM_SUPPORT)
        org = _org(db, name="Correction Guard Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="correction-guards",
            ),
            db=db,
            current_user=admin,
        )
        upload = PlatformMigrationUpload(
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            filename="properties.csv",
            file_format="CSV",
            file_sha256="5" * 64,
            normalized_fingerprint="6" * 64,
            detected_resource="PROPERTIES",
            sheet_name="CSV",
            headers=["Id", "Name"],
            column_mapping={"Id": "source_id", "Name": "name"},
            validation_summary={"valid": 1, "invalid": 1},
            status="STAGED",
            row_count=2,
            created_by_platform_user_id=admin.id,
        )
        db.add(upload)
        db.flush()
        safe = PlatformMigrationStagedRow(
            upload_id=upload.id,
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            resource="PROPERTIES",
            row_number=2,
            source_id="AF-SAFE",
            disposition="NEW",
            row_fingerprint="7" * 64,
            normalized_data={"source_id": "AF-SAFE", "name": "Safe Name"},
            warnings=[],
            errors=[],
        )
        invalid = PlatformMigrationStagedRow(
            upload_id=upload.id,
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            resource="PROPERTIES",
            row_number=3,
            source_id=None,
            disposition="INVALID",
            row_fingerprint="8" * 64,
            normalized_data={"source_id": None, "name": "Invalid Name"},
            warnings=[],
            errors=["Source property identity is missing."],
        )
        db.add_all([safe, invalid])
        db.commit()

        with pytest.raises(HTTPException) as exc:
            api.correct_appfolio_staged_row(
                run.id,
                upload.id,
                safe.id,
                AppFolioStagedRowCorrectionIn(
                    field_name="source_id",
                    corrected_value="SYNTHETIC-ID",
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409

        with pytest.raises(HTTPException) as exc:
            api.correct_appfolio_staged_row(
                run.id,
                upload.id,
                invalid.id,
                AppFolioStagedRowCorrectionIn(
                    field_name="name",
                    corrected_value="Cannot Rescue",
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409

        with pytest.raises(HTTPException) as exc:
            api.correct_appfolio_staged_row(
                run.id,
                upload.id,
                safe.id,
                AppFolioStagedRowCorrectionIn(
                    field_name="name",
                    corrected_value="Support Change",
                ),
                db=db,
                current_user=support,
            )
        assert exc.value.status_code == 403
        assert db.get(PlatformMigrationStagedRow, safe.id).normalized_data["name"] == "Safe Name"
    finally:
        db.close()
        engine.dispose()



def test_appfolio_run_scoped_correction_rule_applies_exact_value_and_replays():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Correction Rule Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="rule-run",
            ),
            db=db,
            current_user=admin,
        )
        upload = PlatformMigrationUpload(
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            filename="properties.csv",
            file_format="CSV",
            file_sha256="9" * 64,
            normalized_fingerprint="a" * 64,
            detected_resource="PROPERTIES",
            sheet_name="CSV",
            headers=["Id", "Name", "City"],
            column_mapping={"Id": "source_id", "Name": "name", "City": "city"},
            validation_summary={"valid": 3},
            status="STAGED",
            row_count=3,
            created_by_platform_user_id=admin.id,
        )
        db.add(upload)
        db.flush()
        rows = [
            PlatformMigrationStagedRow(
                upload_id=upload.id,
                run_id=run.id,
                organization_id=org.id,
                provider="APPFOLIO",
                resource="PROPERTIES",
                row_number=2,
                source_id="RULE-1",
                disposition="NEW",
                row_fingerprint="b" * 64,
                normalized_data={"source_id": "RULE-1", "name": "Apt Bldg", "city": "Cleveland"},
                correction_evidence=[],
                warnings=[],
                errors=[],
            ),
            PlatformMigrationStagedRow(
                upload_id=upload.id,
                run_id=run.id,
                organization_id=org.id,
                provider="APPFOLIO",
                resource="PROPERTIES",
                row_number=3,
                source_id="RULE-2",
                disposition="NEW",
                row_fingerprint="c" * 64,
                normalized_data={"source_id": "RULE-2", "name": "Apt Bldg", "city": "Lakewood"},
                correction_evidence=[],
                warnings=[],
                errors=[],
            ),
            PlatformMigrationStagedRow(
                upload_id=upload.id,
                run_id=run.id,
                organization_id=org.id,
                provider="APPFOLIO",
                resource="PROPERTIES",
                row_number=4,
                source_id="RULE-3",
                disposition="NEW",
                row_fingerprint="d" * 64,
                normalized_data={"source_id": "RULE-3", "name": "Apartment", "city": "Parma"},
                correction_evidence=[],
                warnings=[],
                errors=[],
            ),
        ]
        db.add_all(rows)
        run.last_dry_run_fingerprint = "e" * 64
        run.last_dry_run_summary = {"resource": "PROPERTIES"}
        run.status = "DRY_RUN_READY"
        db.commit()
        for row in rows:
            db.refresh(row)
        db.refresh(upload)

        before_review = api._staged_review_fingerprint(upload, rows)
        before_properties = db.query(Property).count()
        before_charges = db.query(Charge).count()
        before_gl = db.query(GLTransaction).count()

        result = api.create_appfolio_correction_rule(
            run.id,
            upload.id,
            rows[0].id,
            AppFolioMigrationCorrectionRuleCreateIn(
                field_name="name",
                corrected_value="Apartment",
            ),
            db=db,
            current_user=admin,
        )
        assert result.replayed is False
        assert result.applied_row_count == 2
        assert result.resource == "PROPERTIES"
        assert result.source_value == "Apt Bldg"
        assert result.corrected_value == "Apartment"

        db.refresh(run)
        for row in rows:
            db.refresh(row)
        assert rows[0].normalized_data["name"] == "Apartment"
        assert rows[1].normalized_data["name"] == "Apartment"
        assert rows[2].normalized_data["name"] == "Apartment"
        assert rows[0].source_id == "RULE-1"
        assert rows[0].row_fingerprint == "b" * 64
        assert upload.normalized_fingerprint == "a" * 64
        assert rows[0].correction_evidence[-1]["source_value"] == "Apt Bldg"
        assert rows[0].correction_evidence[-1]["corrected_value"] == "Apartment"
        assert rows[0].correction_evidence[-1]["kind"] == "RUN_CORRECTION_RULE"
        assert rows[2].correction_evidence == []
        assert run.last_dry_run_fingerprint is None
        assert run.last_dry_run_summary is None
        assert run.status == "STAGED"
        assert api._staged_review_fingerprint(upload, rows) != before_review

        rule = db.query(PlatformMigrationCorrectionRule).one()
        assert rule.run_id == run.id
        assert rule.organization_id == org.id
        assert rule.source_value == "Apt Bldg"

        application_audits = db.query(AuditLog).filter(
            AuditLog.action == "appfolio_correction_rule_applied"
        ).all()
        assert len(application_audits) == 2
        assert all(json.loads(a.new_value)["customer_target_mutation"] is False for a in application_audits)
        assert all(json.loads(a.new_value)["accounting_mutation"] is False for a in application_audits)

        # Replay the same rule from a fresh staged row carrying the exact source value.
        upload2 = PlatformMigrationUpload(
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            filename="properties-2.csv",
            file_format="CSV",
            file_sha256="f" * 64,
            normalized_fingerprint="0" * 64,
            detected_resource="PROPERTIES",
            sheet_name="CSV",
            headers=["Id", "Name"],
            column_mapping={"Id": "source_id", "Name": "name"},
            validation_summary={"valid": 1},
            status="STAGED",
            row_count=1,
            created_by_platform_user_id=admin.id,
        )
        db.add(upload2)
        db.flush()
        row4 = PlatformMigrationStagedRow(
            upload_id=upload2.id,
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            resource="PROPERTIES",
            row_number=2,
            source_id="RULE-4",
            disposition="NEW",
            row_fingerprint="1" * 64,
            normalized_data={"source_id": "RULE-4", "name": "Apt Bldg"},
            correction_evidence=[],
            warnings=[],
            errors=[],
        )
        db.add(row4)
        db.commit()
        replay = api.create_appfolio_correction_rule(
            run.id,
            upload2.id,
            row4.id,
            AppFolioMigrationCorrectionRuleCreateIn(
                field_name="name",
                corrected_value="Apartment",
            ),
            db=db,
            current_user=admin,
        )
        assert replay.replayed is True
        assert replay.id == rule.id
        assert replay.applied_row_count == 1
        db.refresh(row4)
        assert row4.normalized_data["name"] == "Apartment"
        assert db.query(PlatformMigrationCorrectionRule).count() == 1

        assert db.query(Property).count() == before_properties
        assert db.query(Charge).count() == before_charges
        assert db.query(GLTransaction).count() == before_gl
    finally:
        db.close()
        engine.dispose()


def test_appfolio_correction_rule_conflict_scope_and_authorization_fail_closed():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        support = _platform_user(db, PlatformUserRole.PLATFORM_SUPPORT)
        org = _org(db, name="Correction Conflict Org")
        other_org = _org(db, name="Correction Other Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="conflict-run",
            ),
            db=db,
            current_user=admin,
        )
        upload = PlatformMigrationUpload(
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            filename="vendors.csv",
            file_format="CSV",
            file_sha256="2" * 64,
            normalized_fingerprint="3" * 64,
            detected_resource="VENDORS",
            sheet_name="CSV",
            headers=["VendorId", "Name"],
            column_mapping={"VendorId": "source_id", "Name": "name"},
            validation_summary={"valid": 1},
            status="STAGED",
            row_count=1,
            created_by_platform_user_id=admin.id,
        )
        db.add(upload)
        db.flush()
        row = PlatformMigrationStagedRow(
            upload_id=upload.id,
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            resource="VENDORS",
            row_number=2,
            source_id="VEN-1",
            disposition="NEW",
            row_fingerprint="4" * 64,
            normalized_data={"source_id": "VEN-1", "name": "ABC Maint"},
            correction_evidence=[],
            warnings=[],
            errors=[],
        )
        db.add(row)
        db.add(
            PlatformMigrationCorrectionRule(
                run_id=run.id,
                organization_id=org.id,
                provider="APPFOLIO",
                resource="VENDORS",
                field_name="name",
                source_value="ABC Maint",
                corrected_value="ABC Maintenance",
                created_by_platform_user_id=admin.id,
            )
        )
        db.commit()

        with pytest.raises(HTTPException) as exc:
            api.create_appfolio_correction_rule(
                run.id,
                upload.id,
                row.id,
                AppFolioMigrationCorrectionRuleCreateIn(
                    field_name="name",
                    corrected_value="ABC Maintenance LLC",
                ),
                db=db,
                current_user=admin,
            )
        assert exc.value.status_code == 409
        db.refresh(row)
        assert row.normalized_data["name"] == "ABC Maint"

        with pytest.raises(HTTPException) as exc:
            api.create_appfolio_correction_rule(
                run.id,
                upload.id,
                row.id,
                AppFolioMigrationCorrectionRuleCreateIn(
                    field_name="name",
                    corrected_value="ABC Maintenance",
                ),
                db=db,
                current_user=support,
            )
        assert exc.value.status_code == 403

        # A rule from another organization/run never participates in this run.
        other_run = PlatformMigrationRun(
            organization_id=other_org.id,
            provider="APPFOLIO",
            source_account_ref="other",
            status="DRAFT",
            created_by_platform_user_id=admin.id,
        )
        db.add(other_run)
        db.flush()
        db.add(
            PlatformMigrationCorrectionRule(
                run_id=other_run.id,
                organization_id=other_org.id,
                provider="APPFOLIO",
                resource="VENDORS",
                field_name="name",
                source_value="ABC Maint",
                corrected_value="Foreign Rule",
                created_by_platform_user_id=admin.id,
            )
        )
        db.commit()
        response = Response()
        listed = api.list_appfolio_correction_rules(
            run.id,
            response=response,
            db=db,
            current_user=admin,
        )
        assert response.headers["cache-control"] == "no-store"
        assert len(listed) == 1
        assert listed[0].corrected_value == "ABC Maintenance"
    finally:
        db.close()
        engine.dispose()



def test_appfolio_correction_rule_update_is_explicit_audited_and_changes_fingerprint():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Correction Update Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="correction-update",
            ),
            db=db,
            current_user=admin,
        )
        upload = PlatformMigrationUpload(
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            filename="properties.csv",
            file_format="CSV",
            file_sha256="2" * 64,
            normalized_fingerprint="3" * 64,
            detected_resource="PROPERTIES",
            sheet_name="CSV",
            headers=["Id", "Name"],
            column_mapping={"Id": "source_id", "Name": "name"},
            validation_summary={"valid": 1},
            status="STAGED",
            row_count=1,
            created_by_platform_user_id=admin.id,
        )
        db.add(upload)
        db.flush()
        row = PlatformMigrationStagedRow(
            upload_id=upload.id,
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            resource="PROPERTIES",
            row_number=2,
            source_id="UPD-1",
            disposition="NEW",
            row_fingerprint="4" * 64,
            normalized_data={"source_id": "UPD-1", "name": "Apt Bldg"},
            correction_evidence=[],
            warnings=[],
            errors=[],
        )
        db.add(row)
        db.commit()

        created = api.create_appfolio_correction_rule(
            run.id,
            upload.id,
            row.id,
            AppFolioMigrationCorrectionRuleCreateIn(
                field_name="name",
                corrected_value="Apartment",
            ),
            db=db,
            current_user=admin,
        )
        db.refresh(row)
        before_fingerprint = api._staged_review_fingerprint(upload, [row])
        run.last_dry_run_fingerprint = "5" * 64
        run.last_dry_run_summary = {"staged_review_fingerprint": before_fingerprint}
        run.status = "DRY_RUN_READY"
        db.commit()

        updated = api.update_appfolio_correction_rule(
            run.id,
            created.id,
            AppFolioMigrationCorrectionRuleUpdateIn(
                corrected_value="Apartment Building",
            ),
            db=db,
            current_user=admin,
        )
        db.refresh(row)
        db.refresh(run)
        assert updated.replayed is False
        assert updated.applied_row_count == 1
        assert updated.corrected_value == "Apartment Building"
        assert row.normalized_data["name"] == "Apartment Building"
        assert row.correction_evidence[-1]["source_value"] == "Apt Bldg"
        assert row.correction_evidence[-1]["corrected_value"] == "Apartment Building"
        assert run.last_dry_run_fingerprint is None
        assert run.last_dry_run_summary is None
        assert run.status == "STAGED"
        assert api._staged_review_fingerprint(upload, [row]) != before_fingerprint

        audit = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_correction_rule",
            AuditLog.entity_id == created.id,
            AuditLog.action == "appfolio_correction_rule_updated",
        ).one()
        old_value = json.loads(audit.old_value)
        new_value = json.loads(audit.new_value)
        assert old_value["corrected_value"] == "Apartment"
        assert new_value["corrected_value"] == "Apartment Building"
        assert new_value["dry_run_invalidated"] is True

        replay = api.update_appfolio_correction_rule(
            run.id,
            created.id,
            AppFolioMigrationCorrectionRuleUpdateIn(
                corrected_value="Apartment Building",
            ),
            db=db,
            current_user=admin,
        )
        assert replay.replayed is True
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_correction_rule",
            AuditLog.entity_id == created.id,
            AuditLog.action == "appfolio_correction_rule_updated",
        ).count() == 1
    finally:
        db.close()
        engine.dispose()


def test_appfolio_new_upload_applies_current_exact_run_correction_rule_only():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Correction Replay Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="correction-replay",
            ),
            db=db,
            current_user=admin,
        )
        first_content = (
            "Property Id,Property Name,Address 1,City,State,Zip\n"
            "RULE-SEED,Apt Bldg,1 Rule St,Cleveland,OH,44113\n"
        ).encode()
        first = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("seed.csv", first_content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        seed_row = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == first.id
        ).one()
        rule = api.create_appfolio_correction_rule(
            run.id,
            first.id,
            seed_row.id,
            AppFolioMigrationCorrectionRuleCreateIn(
                field_name="name",
                corrected_value="Apartment",
            ),
            db=db,
            current_user=admin,
        )
        assert rule.applied_row_count == 1

        second_content = (
            "Property Id,Property Name,Address 1,City,State,Zip\n"
            "RULE-LATER,Apt Bldg,2 Rule St,Cleveland,OH,44113\n"
            "RULE-OTHER,Apt Bldg East,3 Rule St,Cleveland,OH,44113\n"
        ).encode()
        second = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("later.csv", second_content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        later_rows = db.query(PlatformMigrationStagedRow).filter(
            PlatformMigrationStagedRow.upload_id == second.id
        ).order_by(PlatformMigrationStagedRow.row_number.asc()).all()
        assert later_rows[0].normalized_data["name"] == "Apartment"
        assert later_rows[0].correction_evidence[-1]["rule_id"] == rule.id
        assert later_rows[1].normalized_data["name"] == "Apt Bldg East"
        assert later_rows[1].correction_evidence == []

        applied_audits = db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == later_rows[0].id,
            AuditLog.action == "appfolio_correction_rule_applied",
        ).count()
        assert applied_audits == 1

        replay = asyncio.run(
            api.stage_appfolio_upload(
                run.id,
                file=_upload_file("later-replay.csv", second_content),
                resource=None,
                sheet_name=None,
                column_mapping_json=None,
                db=db,
                current_user=admin,
            )
        )
        assert replay.replayed is True
        assert db.query(AuditLog).filter(
            AuditLog.entity_type == "platform_migration_staged_row",
            AuditLog.entity_id == later_rows[0].id,
            AuditLog.action == "appfolio_correction_rule_applied",
        ).count() == 1
        assert db.query(Property).count() == 0
        assert db.query(Charge).count() == 0
        assert db.query(GLTransaction).count() == 0
    finally:
        db.close()
        engine.dispose()


# ---------------------------------------------------------------------------
# Phase 4.13 run-level migration review/readiness summary
# ---------------------------------------------------------------------------

def _review_summary_upload(
    db,
    *,
    run: PlatformMigrationRun,
    organization_id: int,
    resource: str,
    suffix: str,
) -> PlatformMigrationUpload:
    upload = PlatformMigrationUpload(
        run_id=run.id,
        organization_id=organization_id,
        provider="APPFOLIO",
        filename=f"{resource.lower()}-{suffix}.csv",
        file_format="CSV",
        file_sha256=(suffix[:1] or "a") * 64,
        normalized_fingerprint=(suffix[-1:] or "b") * 64,
        detected_resource=resource,
        sheet_name="Sheet1",
        headers=["Id"],
        column_mapping={"Id": "source_id"},
        validation_summary={"total": 0},
        status="STAGED",
        row_count=0,
        created_by_platform_user_id=run.created_by_platform_user_id,
    )
    db.add(upload)
    db.flush()
    return upload


def _review_summary_row(
    db,
    *,
    run: PlatformMigrationRun,
    upload: PlatformMigrationUpload,
    row_number: int,
    source_id: str,
    disposition: str,
    resolution_action: str | None = None,
    warnings: list[str] | None = None,
    errors: list[str] | None = None,
    organization_id: int | None = None,
) -> PlatformMigrationStagedRow:
    row = PlatformMigrationStagedRow(
        upload_id=upload.id,
        run_id=run.id,
        organization_id=organization_id or run.organization_id,
        provider="APPFOLIO",
        resource=upload.detected_resource,
        row_number=row_number,
        source_id=source_id,
        disposition=disposition,
        row_fingerprint=(f"{row_number:x}" * 64)[:64],
        normalized_data={"source_id": source_id},
        correction_evidence=[],
        warnings=warnings or [],
        errors=errors or [],
        resolution_action=resolution_action,
    )
    db.add(row)
    db.flush()
    return row


def test_appfolio_run_review_summary_mixed_resources_is_read_only_and_no_store():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        support = _platform_user(db, PlatformUserRole.PLATFORM_SUPPORT)
        org = _org(db, name="Review Summary Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="review-summary",
            ),
            db=db,
            current_user=admin,
        )

        target = Property(
            organization_id=org.id,
            name="Existing Review Target",
            address_line1="1 Review Way",
            city="Cleveland",
            state="OH",
            zip_code="44113",
            is_active=True,
        )
        db.add(target)
        db.flush()

        properties = _review_summary_upload(
            db, run=run, organization_id=org.id, resource="PROPERTIES", suffix="a"
        )
        general_ledger = _review_summary_upload(
            db, run=run, organization_id=org.id, resource="GENERAL_LEDGER", suffix="b"
        )
        tenants = _review_summary_upload(
            db, run=run, organization_id=org.id, resource="TENANTS", suffix="c"
        )

        _review_summary_row(
            db, run=run, upload=properties, row_number=1,
            source_id="P-NEW", disposition="NEW",
            warnings=["descriptive warning"],
        )
        matched = _review_summary_row(
            db, run=run, upload=properties, row_number=2,
            source_id="P-MATCH", disposition="POSSIBLE_MATCH",
            resolution_action="MATCH_EXISTING",
        )
        matched.resolution_target_id = target.id
        _review_summary_row(
            db, run=run, upload=properties, row_number=3,
            source_id="P-SKIP", disposition="REVIEW",
            resolution_action="SKIP",
        )

        _review_summary_row(
            db, run=run, upload=general_ledger, row_number=1,
            source_id="GL-ACCEPT", disposition="REVIEW",
            resolution_action="ACCEPT_RELATIONSHIP",
        )
        _review_summary_row(
            db, run=run, upload=general_ledger, row_number=2,
            source_id="GL-BAD", disposition="INVALID",
            errors=["source row invalid"],
        )
        _review_summary_row(
            db, run=run, upload=general_ledger, row_number=3,
            source_id="GL-POSSIBLE", disposition="POSSIBLE_MATCH",
            warnings=["needs review"],
        )
        _review_summary_row(
            db, run=run, upload=tenants, row_number=1,
            source_id="T-MAPPED", disposition="ALREADY_MAPPED",
        )
        db.add(
            PlatformMigrationItem(
                run_id=run.id,
                organization_id=org.id,
                provider="APPFOLIO",
                resource="TENANTS",
                source_id="T-MAPPED",
                target_entity="TENANT_USER",
                target_id=999999,
                source_fingerprint="d" * 64,
                created_by_platform_user_id=admin.id,
            )
        )

        properties.row_count = 3
        general_ledger.row_count = 3
        tenants.row_count = 1
        stored_run = db.get(PlatformMigrationRun, run.id)
        stored_run.last_dry_run_fingerprint = "e" * 64
        stored_run.last_dry_run_summary = {"resource": "PROPERTIES", "total": 2}
        stored_run.status = "DRY_RUN_READY"
        db.commit()

        before = {
            "uploads": db.query(PlatformMigrationUpload).count(),
            "rows": db.query(PlatformMigrationStagedRow).count(),
            "items": db.query(PlatformMigrationItem).count(),
            "properties": db.query(Property).count(),
            "units": db.query(Unit).count(),
            "users": db.query(User).count(),
            "vendors": db.query(Vendor).count(),
            "leases": db.query(Lease).count(),
            "charges": db.query(Charge).count(),
            "bills": db.query(Bill).count(),
            "gl_transactions": db.query(GLTransaction).count(),
            "work_orders": db.query(WorkOrder).count(),
        }

        response = Response()
        summary = api.get_appfolio_migration_review_summary(
            run.id,
            response=response,
            db=db,
            current_user=support,
        )

        assert response.headers["cache-control"] == "no-store"
        assert summary.upload_count == 3
        assert summary.row_count == 7
        assert summary.new == 1
        assert summary.possible_matches == 1
        assert summary.matched_existing == 1
        assert summary.create_new == 0
        assert summary.accepted_relationships == 1
        assert summary.already_imported == 1
        assert summary.skipped == 1
        assert summary.invalid == 1
        assert summary.unresolved_review == 0
        assert summary.warning_count == 2
        assert summary.blocking_count == 2
        assert summary.mapping_count == 1
        assert summary.dry_run_state == "CURRENT"
        assert summary.last_dry_run_resource == "PROPERTIES"
        assert summary.last_dry_run_fingerprint == "e" * 64

        by_resource = {item.resource: item for item in summary.resources}
        assert by_resource["PROPERTIES"].row_count == 3
        assert by_resource["PROPERTIES"].new == 1
        assert by_resource["PROPERTIES"].matched_existing == 1
        assert by_resource["PROPERTIES"].skipped == 1
        assert by_resource["GENERAL_LEDGER"].accepted_relationships == 1
        assert by_resource["GENERAL_LEDGER"].possible_matches == 1
        assert by_resource["GENERAL_LEDGER"].invalid == 1
        assert by_resource["GENERAL_LEDGER"].blocking_count == 2
        assert by_resource["TENANTS"].already_imported == 1
        assert by_resource["TENANTS"].mapping_count == 1

        after = {
            "uploads": db.query(PlatformMigrationUpload).count(),
            "rows": db.query(PlatformMigrationStagedRow).count(),
            "items": db.query(PlatformMigrationItem).count(),
            "properties": db.query(Property).count(),
            "units": db.query(Unit).count(),
            "users": db.query(User).count(),
            "vendors": db.query(Vendor).count(),
            "leases": db.query(Lease).count(),
            "charges": db.query(Charge).count(),
            "bills": db.query(Bill).count(),
            "gl_transactions": db.query(GLTransaction).count(),
            "work_orders": db.query(WorkOrder).count(),
        }
        assert after == before
        assert db.get(PlatformMigrationRun, run.id).last_dry_run_fingerprint == "e" * 64
    finally:
        db.close()
        engine.dispose()


def test_appfolio_run_review_summary_filters_inconsistent_org_rows_and_fails_closed_for_role():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        sales = _platform_user(db, PlatformUserRole.PLATFORM_SALES)
        org = _org(db, name="Review Scope Org")
        foreign_org = _org(db, name="Review Scope Foreign")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="review-scope",
            ),
            db=db,
            current_user=admin,
        )

        own_upload = _review_summary_upload(
            db, run=run, organization_id=org.id, resource="PROPERTIES", suffix="f"
        )
        foreign_upload = _review_summary_upload(
            db, run=run, organization_id=foreign_org.id, resource="UNITS", suffix="9"
        )
        _review_summary_row(
            db, run=run, upload=own_upload, row_number=1,
            source_id="OWN", disposition="REVIEW",
        )
        _review_summary_row(
            db, run=run, upload=foreign_upload, row_number=1,
            source_id="FOREIGN", disposition="INVALID",
            errors=["must not leak"],
            organization_id=foreign_org.id,
        )
        own_upload.row_count = 1
        foreign_upload.row_count = 1
        db.commit()

        summary = api.get_appfolio_migration_review_summary(
            run.id,
            response=Response(),
            db=db,
            current_user=admin,
        )
        assert summary.upload_count == 1
        assert summary.row_count == 1
        assert summary.unresolved_review == 1
        assert summary.blocking_count == 1
        assert [item.resource for item in summary.resources] == ["PROPERTIES"]
        assert summary.dry_run_state == "NONE_OR_STALE"

        with pytest.raises(HTTPException) as exc:
            api.get_appfolio_migration_review_summary(
                run.id,
                response=Response(),
                db=db,
                current_user=sales,
            )
        assert exc.value.status_code == 403
    finally:
        db.close()
        engine.dispose()


def test_appfolio_run_review_summary_route_is_exposed():
    from app.main import app

    assert (
        "/api/platform/migrations/appfolio/runs/{run_id}/review-summary"
        in set(app.openapi()["paths"])
    )


def test_appfolio_run_review_fingerprint_is_deterministic_and_binds_all_review_evidence():
    db, engine = _session()
    try:
        admin = _platform_user(db, PlatformUserRole.PLATFORM_ADMIN)
        org = _org(db, name="Review Fingerprint Org")
        run = api.create_run(
            AppFolioMigrationRunCreateIn(
                organization_id=org.id,
                source_account_ref="review-fingerprint",
            ),
            db=db,
            current_user=admin,
        )
        upload = _review_summary_upload(
            db,
            run=run,
            organization_id=org.id,
            resource="PROPERTIES",
            suffix="a",
        )
        row = _review_summary_row(
            db,
            run=run,
            upload=upload,
            row_number=1,
            source_id="P-FP",
            disposition="NEW",
        )
        row.normalized_data = {
            "source_id": "P-FP",
            "name": "Original Name",
            "city": "Cleveland",
        }
        upload.row_count = 1

        rule = PlatformMigrationCorrectionRule(
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            resource="PROPERTIES",
            field_name="name",
            source_value="Original Name",
            corrected_value="Reviewed Name",
            created_by_platform_user_id=admin.id,
        )
        mapping = PlatformMigrationItem(
            run_id=run.id,
            organization_id=org.id,
            provider="APPFOLIO",
            resource="PROPERTIES",
            source_id="P-FP",
            target_entity="PROPERTY",
            target_id=12345,
            source_fingerprint="1" * 64,
            created_by_platform_user_id=admin.id,
        )
        db.add_all([rule, mapping])
        db.commit()

        def fingerprint() -> str:
            return api.get_appfolio_migration_review_summary(
                run.id,
                response=Response(),
                db=db,
                current_user=admin,
            ).review_fingerprint

        first = fingerprint()
        assert len(first) == 64
        assert fingerprint() == first

        upload.normalized_fingerprint = "2" * 64
        db.commit()
        after_upload = fingerprint()
        assert after_upload != first

        row.normalized_data = {
            **dict(row.normalized_data),
            "city": "Lakewood",
        }
        db.commit()
        after_normalized = fingerprint()
        assert after_normalized != after_upload

        row.row_fingerprint = "3" * 64
        db.commit()
        after_source_row = fingerprint()
        assert after_source_row != after_normalized

        row.correction_evidence = [
            {
                "kind": "ROW_CORRECTION",
                "field_name": "city",
                "source_value": "Cleveland",
                "corrected_value": "Lakewood",
            }
        ]
        db.commit()
        after_evidence = fingerprint()
        assert after_evidence != after_source_row

        row.resolution_action = "SKIP"
        db.commit()
        after_resolution = fingerprint()
        assert after_resolution != after_evidence

        rule.corrected_value = "Final Reviewed Name"
        db.commit()
        after_rule = fingerprint()
        assert after_rule != after_resolution

        mapping.source_fingerprint = "4" * 64
        db.commit()
        after_mapping_source = fingerprint()
        assert after_mapping_source != after_rule

        mapping.target_id = 54321
        db.commit()
        after_mapping_target = fingerprint()
        assert after_mapping_target != after_mapping_source

        before_counts = (
            db.query(PlatformMigrationUpload).count(),
            db.query(PlatformMigrationStagedRow).count(),
            db.query(PlatformMigrationCorrectionRule).count(),
            db.query(PlatformMigrationItem).count(),
            db.query(Property).count(),
            db.query(Unit).count(),
            db.query(User).count(),
            db.query(Vendor).count(),
            db.query(Lease).count(),
            db.query(Charge).count(),
            db.query(Bill).count(),
            db.query(GLTransaction).count(),
            db.query(WorkOrder).count(),
        )
        assert fingerprint() == after_mapping_target
        after_counts = (
            db.query(PlatformMigrationUpload).count(),
            db.query(PlatformMigrationStagedRow).count(),
            db.query(PlatformMigrationCorrectionRule).count(),
            db.query(PlatformMigrationItem).count(),
            db.query(Property).count(),
            db.query(Unit).count(),
            db.query(User).count(),
            db.query(Vendor).count(),
            db.query(Lease).count(),
            db.query(Charge).count(),
            db.query(Bill).count(),
            db.query(GLTransaction).count(),
            db.query(WorkOrder).count(),
        )
        assert after_counts == before_counts
    finally:
        db.close()
        engine.dispose()
